from __future__ import annotations

from typing import TYPE_CHECKING
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from expo_jbm329.services.analysis.statistics import analyze_descriptive_statistics
from expo_jbm329.services.analysis.statistics_charts import StatisticsChartLabels, render_statistics_charts
from expo_jbm329.services.data_io_service import DataIOService
from expo_jbm329.services.file_loader import FileLoader
from expo_jbm329.services.file_writer import ExportCancelledError, FileWriter

if TYPE_CHECKING:
    from pathlib import Path


@pytest.fixture
def tmp_csv(tmp_path: Path) -> Path:
    path = tmp_path / "test.csv"
    pd.DataFrame({"A": [1, 2], "B": ["x", "y"]}).to_csv(path, index=False)
    return path


@pytest.fixture
def loader() -> FileLoader:
    return FileLoader()


@pytest.fixture
def writer() -> FileWriter:
    return FileWriter()


@pytest.fixture
def io_service(loader: FileLoader, writer: FileWriter) -> DataIOService:
    return DataIOService(loader, writer)


def test_file_loader_load_csv(loader: FileLoader, tmp_csv: Path) -> None:
    df = loader.load_df_auto(str(tmp_csv))

    assert len(df) == 2
    assert list(df.columns) == ["A", "B"]


def test_file_writer_save_csv(writer: FileWriter, tmp_path: Path) -> None:
    dest = tmp_path / "out.csv"
    df = pd.DataFrame({"X": [10, 20]})

    writer.save_csv(df, dest)

    assert dest.exists()
    assert pd.read_csv(dest)["X"].tolist() == [10, 20]


def test_data_io_service_load(io_service: DataIOService, tmp_csv: Path) -> None:
    res = io_service.resolve_and_load_df(str(tmp_csv))

    assert res.ok is True
    assert isinstance(res.data, pd.DataFrame)
    assert list(res.data.columns) == ["A", "B"]


def test_data_io_service_load_missing_file_returns_failure(io_service: DataIOService) -> None:
    res = io_service.resolve_and_load_df("missing.csv")

    assert res.ok is False
    assert res.cancelled is False
    assert res.data is None
    assert res.error is not None
    assert "missing.csv" in res.error


def test_data_io_service_export_csv(io_service: DataIOService, tmp_path: Path) -> None:
    dest = tmp_path / "export.csv"
    df = pd.DataFrame({"Z": [100]})

    res = io_service.export_df_csv(df, dest)

    assert res.ok is True
    assert dest.exists()


def test_data_io_service_export_profile(io_service: DataIOService, tmp_path: Path) -> None:
    dest = tmp_path / "profile.html"
    df = pd.DataFrame({"A": [1]})
    mock_profile = MagicMock()

    with patch("expo_jbm329.services.data_io_service.generate_profile_report", return_value=mock_profile) as mock_gen:
        io_service._writer = MagicMock()
        res = io_service.export_df_profile(df, dest, "My Title")

    assert res.ok is True
    mock_gen.assert_called_once()
    io_service._writer.save_profile.assert_called_once_with(mock_profile, str(dest), corr_id=None)


def test_data_io_service_export_comparison_profile(io_service: DataIOService, tmp_path: Path) -> None:
    dest = tmp_path / "compare.html"
    data = [(pd.DataFrame({"A": [1]}), "One"), (pd.DataFrame({"A": [2]}), "Two")]
    mock_report = MagicMock()

    with patch(
        "expo_jbm329.services.data_io_service.generate_comparison_profile_report",
        return_value=mock_report,
    ) as mock_gen:
        io_service._writer = MagicMock()
        res = io_service.export_dfs_profile(data, dest)

    assert res.ok is True
    mock_gen.assert_called_once_with(data, corr_id=None)
    io_service._writer.save_profile.assert_called_once_with(mock_report, str(dest), corr_id=None)


@pytest.mark.parametrize("streaming", [True, False])
def test_named_excel_service_roundtrip(
    io_service: DataIOService, writer: FileWriter, tmp_path: Path, streaming: bool
) -> None:
    writer.reload_settings({"excel": {"streaming": streaming}})
    path = tmp_path / "named.xlsx"
    progress: list[int] = []
    result = io_service.export_dfs_excel(
        {"Summary": pd.DataFrame({"count": [1]}), "Empty": pd.DataFrame(columns=["value"])},
        path,
        progress_cb=progress.append,
        corr_id="named-service",
        job_id="job",
        job_scope="export:excel",
    )
    assert result.ok
    assert not result.cancelled
    assert result.path == str(path)
    assert result.corr_id == "named-service"
    assert result.elapsed is not None and result.elapsed >= 0
    assert progress[-1] == 100
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Summary", "Empty"]
        assert list(workbook["Summary"].values) == [("count",), (1,)]
        assert list(workbook["Empty"].values) == [("value",)]
    finally:
        workbook.close()


@pytest.mark.parametrize("named", [False, True])
@pytest.mark.parametrize("outcome", ["success", "cancel", "failure"])
def test_excel_service_shared_results_and_forwarding(
    io_service: DataIOService, tmp_path: Path, monkeypatch, named: bool, outcome: str
) -> None:
    path = tmp_path / "out.xlsx"
    frame = pd.DataFrame({"a": [1]})
    source = {"One": frame} if named else frame
    calls = []
    progress: list[int] = []

    def cancel() -> bool:
        return False

    def save(self, actual, dest, **kwargs):
        calls.append((actual, dest, kwargs))
        if outcome == "cancel":
            message = "cancelled"
            raise ExportCancelledError(message, path=dest, rows_written=1, sheets_written=1, corr_id="corr")
        if outcome == "failure":
            message = "writer failure"
            raise ValueError(message)
        return path

    monkeypatch.setattr(FileWriter, "save_excel_sheets" if named else "save_excel", save)
    export = io_service.export_dfs_excel if named else io_service.export_df_excel
    result = export(source, path, progress_cb=progress.append, cancel_cb=cancel, corr_id="corr")
    assert len(calls) == 1
    actual, dest, kwargs = calls[0]
    assert actual["One"] is frame if named else actual is frame
    assert dest == str(path)
    assert kwargs == {"progress_cb": progress.append, "cancel_cb": cancel, "corr_id": "corr", "na_rep": ""}
    assert result.ok is (outcome == "success")
    assert result.cancelled is (outcome == "cancel")
    assert result.path == str(path)
    assert result.corr_id == "corr"
    assert result.error == ("writer failure" if outcome == "failure" else None)
    assert (result.elapsed is not None) is (outcome != "failure")


def test_named_service_early_cancellation_skips_writer(io_service: DataIOService, tmp_path: Path, monkeypatch) -> None:
    def unexpected(*args, **kwargs):
        pytest.fail("Early cancellation must not write")

    monkeypatch.setattr(FileWriter, "save_excel_sheets", unexpected)
    path = tmp_path / "out.xlsx"
    result = io_service.export_dfs_excel(
        {"Data": pd.DataFrame({"a": [1]})}, path, cancel_cb=lambda: True, corr_id="early"
    )
    assert result.cancelled and not result.ok
    assert result.elapsed is None
    assert result.corr_id == "early"
    assert not path.exists()


@pytest.mark.parametrize(
    "sheets",
    [
        {},
        {"bad:name": pd.DataFrame({"a": [1]})},
        {"Empty": pd.DataFrame()},
        {"Bad": object()},
        [],
        pd.DataFrame({"a": [1]}),
    ],
)
def test_named_service_validation_returns_failed_result(io_service: DataIOService, tmp_path: Path, sheets) -> None:
    path = tmp_path / "out.xlsx"
    result = io_service.export_dfs_excel(sheets, path, corr_id="invalid")
    assert not result.ok and not result.cancelled
    assert result.error
    assert result.path == str(path)
    assert result.corr_id == "invalid"
    assert not path.exists()


def test_named_service_real_cancellation_preserves_existing(io_service: DataIOService, tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    path.write_bytes(b"original")
    progress: list[int] = []
    result = io_service.export_dfs_excel(
        {"Data": pd.DataFrame({"a": range(4)})},
        path,
        progress_cb=progress.append,
        cancel_cb=lambda: bool(progress and progress[-1] >= 25),
        corr_id="during",
    )
    assert result.cancelled and not result.ok
    assert result.elapsed is not None
    assert result.error is None
    assert path.read_bytes() == b"original"
    assert 100 not in progress


def test_named_service_validation_facade_preserves_frames(io_service: DataIOService) -> None:
    frame = pd.DataFrame({"a": [1]})
    sheets = {"One": frame}
    captured = io_service.validate_excel_sheets(sheets)
    assert captured is not sheets
    assert captured["One"] is frame


def test_chart_only_excel_export_prepares_images_and_commits_progress(
    io_service: DataIOService,
    tmp_path: Path,
) -> None:
    result = analyze_descriptive_statistics(pd.DataFrame({"first": [1, 2, 3], "second": [3, 4, 5]}))
    path = tmp_path / "charts-only.xlsx"
    progress: list[int] = []

    def make_charts(chart_progress, cancel_cb):
        return render_statistics_charts(
            result.columns,
            StatisticsChartLabels("Histogram", "Boxplot", "No data"),
            progress_cb=chart_progress,
            cancel_cb=cancel_cb,
        )

    exported = io_service.export_dfs_excel(
        {},
        path,
        chart_factory=make_charts,
        chart_sheet_name="Charts",
        progress_cb=progress.append,
    )

    assert exported.ok is True
    assert progress == sorted(set(progress))
    assert progress[0] == 0
    assert progress[-1] == 100
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Charts"]
        assert len(workbook["Charts"]._images) == 2
    finally:
        workbook.close()


def test_chart_export_cancellation_after_preparation_preserves_destination(
    io_service: DataIOService,
    tmp_path: Path,
) -> None:
    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"original")
    result = analyze_descriptive_statistics(pd.DataFrame({"number": [1, 2, 3]}))
    calls = 0

    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls == 2

    def make_charts(_progress_cb, _cancel_cb):
        return render_statistics_charts(
            result.columns,
            StatisticsChartLabels("Histogram", "Boxplot", "No data"),
        )

    exported = io_service.export_dfs_excel(
        {},
        path,
        chart_factory=make_charts,
        chart_sheet_name="Charts",
        cancel_cb=cancel,
    )

    assert exported.ok is False
    assert exported.cancelled is True
    assert path.read_bytes() == b"original"


def test_chart_export_rejects_duplicate_chart_sheet_before_writing(
    io_service: DataIOService,
    tmp_path: Path,
) -> None:
    path = tmp_path / "collision.xlsx"
    exported = io_service.export_dfs_excel(
        {"Charts": pd.DataFrame({"value": [1]})},
        path,
        chart_factory=lambda _progress, _cancel: (),
        chart_sheet_name="charts",
    )

    assert exported.ok is False
    assert exported.error is not None
    assert "Duplicate Excel sheet name" in exported.error
    assert not path.exists()
