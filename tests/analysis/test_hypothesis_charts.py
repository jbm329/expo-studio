from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from typing import TYPE_CHECKING
from unittest.mock import patch

import numpy as np
import pytest
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PyQt6.QtCore import QCoreApplication, QTranslator

from expo_jbm329.services.analysis.categories import HypothesisTest
from expo_jbm329.services.analysis.hypothesis_charts import (
    ChiSquareChartPresentation,
    GroupComparisonChartPresentation,
    HypothesisChartPresentation,
    PairedComparisonChartPresentation,
    build_hypothesis_chart_presentation,
    draw_group_distribution,
    draw_hypothesis_charts,
    heatmap_color_limit,
    render_hypothesis_charts,
)
from expo_jbm329.services.analysis.hypothesis_export import (
    HypothesisExportComponent,
    hypothesis_export_tables,
)
from expo_jbm329.services.analysis.paired_comparison import MAX_PAIRED_TRAJECTORIES
from expo_jbm329.services.excel_chart import ExcelChartCancelledError
from expo_jbm329.services.file_writer import FileWriter
from expo_jbm329.utils.format_utils import fmt_num
from tests.analysis.test_hypothesis_export import chi_snapshot, group_snapshot, paired_snapshot

if TYPE_CHECKING:
    from collections.abc import Callable

    from expo_jbm329.services.analysis.hypothesis_export import HypothesisExportSnapshot, HypothesisResult


def _presentation(snapshot: HypothesisExportSnapshot) -> HypothesisChartPresentation:
    return build_hypothesis_chart_presentation(snapshot.result, lambda _context, text: text)


@pytest.mark.parametrize("factory", [group_snapshot, chi_snapshot, paired_snapshot])
def test_render_exact_displayed_images_is_deterministic_and_reports_progress(
    factory: Callable[[], HypothesisExportSnapshot],
) -> None:
    snapshot = factory()
    presentation = _presentation(snapshot)
    progress = []
    images = render_hypothesis_charts(snapshot, presentation, progress_cb=progress.append)
    assert len(images) == 1
    assert images[0].heading == ""
    assert images[0].accessible_name == " / ".join(snapshot.columns)
    assert images[0].captions == ()
    assert images[0].image_data.startswith(b"\x89PNG")
    assert (images[0].width_px, images[0].height_px) == (1080, 420)
    assert progress == [100]
    assert images == render_hypothesis_charts(snapshot, presentation)


@pytest.mark.parametrize("cancel_at", [1, 2])
def test_renderer_cancellation_never_reports_completed_progress(cancel_at: int) -> None:
    calls = 0
    progress = []

    def cancel():
        nonlocal calls
        calls += 1
        return calls == cancel_at

    with pytest.raises(ExcelChartCancelledError):
        snapshot = group_snapshot()
        render_hypothesis_charts(
            snapshot,
            _presentation(snapshot),
            cancel_cb=cancel,
            progress_cb=progress.append,
        )
    assert progress == []


def test_group_boxes_use_full_summary_min_max_and_literal_labels() -> None:
    snapshot = group_snapshot()
    figure = Figure()
    ax = figure.add_subplot(111)
    draw_group_distribution(ax, snapshot.result.groups, "Distribution")
    assert [label.get_text() for label in ax.get_xticklabels()] == [group.label for group in snapshot.result.groups]
    assert ax.lines[1].get_ydata()[1] == snapshot.result.groups[0].minimum
    assert ax.lines[2].get_ydata()[1] == snapshot.result.groups[0].maximum


def test_residual_heatmap_orientation_color_scale_and_undefined_cells() -> None:
    snapshot = chi_snapshot(3)
    figure = Figure()
    presentation = _presentation(snapshot)
    draw_hypothesis_charts(figure, snapshot.result, presentation)
    ax = figure.axes[0]
    np.testing.assert_allclose(ax.images[0].get_array(), snapshot.result.adjusted_residuals)
    assert ax.get_xlabel() == snapshot.columns[1]
    assert ax.get_ylabel() == snapshot.columns[0]
    assert [label.get_text() for label in ax.get_xticklabels()] == list(snapshot.result.column_labels)
    assert len(ax.texts) == 6
    assert heatmap_color_limit(np.array([[np.nan]])) == 1.96
    assert heatmap_color_limit(np.array([[-4.0, 2.0]])) == 4
    malformed = replace(snapshot.result, adjusted_residuals=((float("nan"), 2, 3), (4, 5, 6)))
    figure = Figure()
    presentation = build_hypothesis_chart_presentation(malformed, lambda _context, text: text)
    draw_hypothesis_charts(figure, malformed, presentation)
    ax = figure.axes[0]
    assert len(ax.texts) == 5
    figure = Figure()
    annotations = tuple(tuple("translated" for _ in row) for row in malformed.adjusted_residuals)
    presentation = build_hypothesis_chart_presentation(
        malformed,
        lambda _context, text: text,
        number_format=lambda _value: "translated",
    )
    assert isinstance(presentation, ChiSquareChartPresentation)
    assert presentation.annotations == annotations
    draw_hypothesis_charts(figure, malformed, presentation)
    ax = figure.axes[0]
    assert all(text.get_text() == "translated" for text in ax.texts)


def test_large_heatmap_omits_annotations() -> None:
    result = chi_snapshot().result
    result = replace(
        result,
        adjusted_residuals=tuple(tuple(0.1 for _ in range(11)) for _ in range(10)),
        row_labels=tuple(map(str, range(10))),
        column_labels=tuple(map(str, range(11))),
    )
    figure = Figure()
    presentation = build_hypothesis_chart_presentation(result, lambda _context, text: text)
    draw_hypothesis_charts(figure, result, presentation)
    ax = figure.axes[0]
    assert not ax.texts


def test_paired_charts_preserve_selected_order_full_boxes_and_deterministic_sample() -> None:
    snapshot = paired_snapshot(3, MAX_PAIRED_TRAJECTORIES + 50)
    result = snapshot.result
    assert result.plot_data.sampled
    figure = Figure()
    presentation = build_hypothesis_chart_presentation(result, lambda _context, text: text)
    draw_hypothesis_charts(figure, result, presentation)
    boxes, trajectories = figure.axes
    assert [label.get_text() for label in boxes.get_xticklabels()] == list(snapshot.columns)
    assert [label.get_text() for label in trajectories.get_xticklabels()] == list(snapshot.columns)
    assert len(trajectories.lines) == MAX_PAIRED_TRAJECTORIES
    assert list(trajectories.lines[0].get_ydata()) == list(result.plot_data.trajectories[0])
    assert boxes.lines[1].get_ydata()[1] == result.summaries[0].minimum
    assert boxes.lines[2].get_ydata()[1] == result.summaries[0].maximum
    figure = Figure()
    result_without_plot_data = replace(result, plot_data=None)
    presentation = build_hypothesis_chart_presentation(result_without_plot_data, lambda _context, text: text)
    draw_hypothesis_charts(figure, result_without_plot_data, presentation)
    assert not figure.axes[1].lines


@pytest.mark.parametrize(
    "factory",
    [
        lambda: group_snapshot(2),
        lambda: group_snapshot(3),
        lambda: chi_snapshot(2),
        lambda: chi_snapshot(3),
        lambda: paired_snapshot(2),
        lambda: paired_snapshot(3),
    ],
)
@pytest.mark.parametrize("locale", ["en", "sv"])
def test_real_localized_app_and_worker_axes_share_semantic_presentation(
    factory: Callable[[], HypothesisExportSnapshot],
    locale: str,
) -> None:
    snapshot = factory()
    translator = QTranslator()
    locale_path = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales" / f"app_{locale}.qm"
    assert translator.load(str(locale_path))
    application = QCoreApplication.instance()
    assert application is not None
    application.installTranslator(translator)
    try:
        presentation = build_hypothesis_chart_presentation(snapshot.result, QCoreApplication.translate)
        app_figure = Figure()
        draw_hypothesis_charts(app_figure, snapshot.result, presentation)
        worker_figures: list[Figure] = []

        def capture_worker_figure(
            figure: Figure,
            result: HypothesisResult,
            captured_presentation: HypothesisChartPresentation,
        ) -> None:
            draw_hypothesis_charts(figure, result, captured_presentation)
            worker_figures.append(figure)

        with patch(
            "expo_jbm329.services.analysis.hypothesis_charts.draw_hypothesis_charts",
            side_effect=capture_worker_figure,
        ):
            images = render_hypothesis_charts(snapshot, presentation)

        def semantic_axes(figure: Figure) -> tuple[tuple[str, str, str, tuple[str, ...], tuple[str, ...]], ...]:
            axes = [axis for axis in figure.axes if axis.images or axis.lines or axis.patches]
            return tuple(
                (
                    axis.get_title(),
                    axis.get_xlabel(),
                    axis.get_ylabel(),
                    tuple(label.get_text() for label in axis.get_xticklabels()),
                    tuple(text.get_text() for text in axis.texts),
                )
                for axis in axes
            )

        assert len(worker_figures) == 1
        axes = semantic_axes(app_figure)
        assert axes == semantic_axes(worker_figures[0])
        if snapshot.test is HypothesisTest.GROUP_COMPARISON:
            assert isinstance(presentation, GroupComparisonChartPresentation)
            expected_axes = (
                (
                    QCoreApplication.translate("GroupComparisonView", "Distribution by group"),
                    presentation.x_axis_label,
                    presentation.y_axis_label,
                ),
            )
            assert tuple(axis[:3] for axis in axes) == expected_axes
            assert presentation.x_axis_label == snapshot.result.grouping_column
            assert presentation.y_axis_label == snapshot.result.numeric_column
            assert app_figure.axes[0].get_xlabel() == presentation.x_axis_label
            assert app_figure.axes[0].get_ylabel() == presentation.y_axis_label
        elif snapshot.test is HypothesisTest.CHI_SQUARE:
            assert isinstance(presentation, ChiSquareChartPresentation)
            expected_axes = (
                (
                    QCoreApplication.translate("ChiSquareView", "Adjusted standardized residuals"),
                    snapshot.result.column_column,
                    snapshot.result.row_column,
                ),
            )
            assert tuple(axis[:3] for axis in axes) == expected_axes
            assert tuple(text.get_text() for text in app_figure.axes[0].texts) == tuple(
                fmt_num(value, sig=3)
                for row in snapshot.result.adjusted_residuals
                for value in row
                if np.isfinite(value)
            )
        else:
            assert isinstance(presentation, PairedComparisonChartPresentation)
            expected_axes = (
                (
                    QCoreApplication.translate("PairedComparisonView", "Distribution by occasion"),
                    QCoreApplication.translate("PairedComparisonView", "Measurement occasion"),
                    QCoreApplication.translate("PairedComparisonView", "Value"),
                ),
                (
                    QCoreApplication.translate("PairedComparisonView", "Subject trajectories"),
                    QCoreApplication.translate("PairedComparisonView", "Measurement occasion"),
                    QCoreApplication.translate("PairedComparisonView", "Value"),
                ),
            )
            assert tuple(axis[:3] for axis in axes) == expected_axes
            assert tuple(label.get_text() for label in app_figure.axes[1].get_xticklabels()) == snapshot.columns
        assert images[0].heading == ""
        assert images[0].accessible_name == " / ".join(snapshot.columns)
        assert images[0].captions == ()
    finally:
        application.removeTranslator(translator)


@pytest.mark.parametrize("factory", [group_snapshot, chi_snapshot, paired_snapshot])
@pytest.mark.parametrize("charts_only", [False, True])
@pytest.mark.parametrize("streaming", [False, True])
def test_real_workbooks_preserve_typed_tables_and_one_chart_sheet(
    tmp_path,
    factory: Callable[[], HypothesisExportSnapshot],
    charts_only: bool,
    streaming: bool,
) -> None:
    snapshot = factory()
    frames = (
        hypothesis_export_tables(
            snapshot,
            (HypothesisExportComponent.SUMMARY, HypothesisExportComponent.TEST_RESULTS),
        )
        if not charts_only
        else {}
    )
    sheets = {key.value: frame for key, frame in frames.items()}
    images = render_hypothesis_charts(snapshot, _presentation(snapshot))
    path = tmp_path / "hypothesis.xlsx"
    FileWriter().save_excel_sheets(sheets, path, chart_images=images, chart_sheet_name="Charts", streaming=streaming)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == [*sheets, "Charts"]
        assert len(workbook["Charts"]._images) == 1
        assert workbook["Charts"]["A1"].value is None
        for name, frame in sheets.items():
            worksheet = workbook[name]
            assert [cell.value for cell in worksheet[1]] == list(frame.columns)
            assert worksheet.max_row == len(frame) + 1
        if not charts_only:
            assert isinstance(workbook["Test results"]["B2"].value, (float, int))
    finally:
        workbook.close()
