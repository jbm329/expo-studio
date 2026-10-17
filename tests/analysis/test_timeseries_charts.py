from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PyQt6.QtCore import QCoreApplication, QTranslator
from PyQt6.QtWidgets import QApplication

from expo_jbm329.gui.dialogs.analysis.timeseries_view import TimeSeriesView
from expo_jbm329.services.analysis import timeseries_charts as chart_module
from expo_jbm329.services.analysis.categories import AnalysisCategory
from expo_jbm329.services.analysis.report import (
    ReportDocument,
    capture_report_labels,
    create_report_entry,
    prepare_report_content,
    render_prepared_report_html,
)
from expo_jbm329.services.analysis.timeseries import analyze_time_series
from expo_jbm329.services.analysis.timeseries_charts import (
    draw_acf_plot,
    draw_decomposition,
    draw_series_plot,
    render_time_series_charts,
)
from expo_jbm329.services.analysis.timeseries_export import (
    TimeSeriesExportComponent,
    TimeSeriesExportSnapshot,
    series_table,
    time_series_export_tables,
)
from expo_jbm329.services.excel_chart import ExcelChartCancelledError
from expo_jbm329.services.file_writer import FileWriter
from expo_jbm329.utils.dataset_ref import DatasetRef
from tests.analysis.test_timeseries_export import CHART_LABELS, irregular_snapshot, time_series_snapshot

_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


def _snapshot_with_plot_count(count: int) -> TimeSeriesExportSnapshot:
    if count == 1:
        return irregular_snapshot()
    if count == 2:
        frame = pd.DataFrame({
            "day": pd.date_range("2024-01-01", periods=4, freq="D"),
            "sales": [1.0, 2.0, 4.0, 3.0],
        })
        result = analyze_time_series(frame, "day", "sales", seasonal_period=7)
        assert result.error is None
        assert result.acf_values and not result.trend
        return replace(time_series_snapshot(), result=result, applied_configuration=None)
    return time_series_snapshot()


@pytest.mark.parametrize("locale", ["en", "sv"])
@pytest.mark.parametrize("plot_count", [1, 2, 3])
def test_live_worker_and_html_share_captured_localized_semantics(locale, plot_count, monkeypatch):
    result = _snapshot_with_plot_count(plot_count).result
    translator = QTranslator()
    catalog = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales" / f"app_{locale}.qm"
    assert translator.load(str(catalog))
    app = QApplication.instance()
    assert app is not None
    app.installTranslator(translator)
    try:
        view = TimeSeriesView(result)
        labels = view.chart_labels()
        snapshot = TimeSeriesExportSnapshot(result, labels, view.export_notes(), view.autocorrelation_note())
        report_labels = capture_report_labels(QCoreApplication.translate)
        expected_titles = (
            ("Series", "Autocorrelation", "Seasonal decomposition")
            if locale == "en"
            else ("Serie", "Autokorrelation", "Säsongsdekomponering")
        )
        assert (
            labels.series_metadata.accessible_name,
            labels.acf_metadata.accessible_name,
            labels.decomposition_metadata.accessible_name,
        ) == expected_titles
        assert labels.value_axis == result.value_column
        assert labels.lag_axis == ("Lag" if locale == "en" else "Fördröjning")
        assert labels.acf_axis == expected_titles[1]
        assert (labels.observed, labels.trend, labels.seasonal, labels.residual) == (
            ("Observed", "Trend", "Seasonal", "Residual")
            if locale == "en"
            else ("Observerad", "Trend", "Säsong", "Residual")
        )
        unavailable_notes = tuple(
            view.note_text(note) for note in (result.diagnostics_note, result.decomposition_note) if note
        )
    finally:
        app.removeTranslator(translator)

    assert view.chart_labels() is labels
    live_canvases = view.findChildren(FigureCanvasQTAgg)
    live = {canvas.figure.axes[0].get_title(): canvas.figure for canvas in live_canvases}
    assert tuple(live) == expected_titles
    for index in range(plot_count, 3):
        axis = live[expected_titles[index]].axes[0]
        assert not axis.axison
        assert axis.texts[0].get_text() in unavailable_notes
    assert all(note in snapshot.notes for note in unavailable_notes)
    if locale == "sv":
        assert all(
            note != source
            for note, source in zip(
                unavailable_notes,
                (note for note in (result.diagnostics_note, result.decomposition_note) if note),
                strict=True,
            )
        )

    def semantics(figure):
        return tuple(
            (
                axis.get_title(),
                axis.get_xlabel(),
                axis.get_ylabel(),
                tuple(
                    tuple(None if pd.isna(value) else float(value) for value in line.get_ydata())
                    for line in axis.get_lines()
                ),
                len(axis.collections),
            )
            for axis in figure.axes
        )

    expected = [semantics(live[title]) for title in expected_titles[:plot_count]]
    assert expected[0][0][:3] == (expected_titles[0], "", result.value_column)
    if plot_count >= 2:
        assert expected[1][0][:3] == (expected_titles[1], labels.lag_axis, labels.acf_axis)
    if plot_count == 3:
        assert [panel[:3] for panel in expected[2]] == [
            (expected_titles[2], "", labels.observed),
            ("", "", labels.trend),
            ("", "", labels.seasonal),
            ("", "", labels.residual),
        ]
    rendered = []
    original_render = chart_module.render_chart_png

    def inspect_png(figure, *args, **kwargs):
        rendered.append(semantics(figure))
        return original_render(figure, *args, **kwargs)

    monkeypatch.setattr(chart_module, "render_chart_png", inspect_png)
    document = ReportDocument()
    entry = create_report_entry(
        AnalysisCategory.TIME_SERIES, DatasetRef("series", "Sales dataset", result.source_rows, 2), snapshot
    )
    document.upsert(entry.with_selection(("charts", "model_summary"), ""))
    prepared = prepare_report_content(document, labels=report_labels)
    assert rendered == expected
    charts = prepared.entries[0].components[0].charts
    assert len(charts) == plot_count
    assert all(chart.heading == "" and chart.captions == () for chart in charts)
    assert tuple(chart.accessible_name for chart in charts) == expected_titles[:plot_count]
    html = render_prepared_report_html(prepared)
    assert "<figcaption>" not in html
    assert "<strong></strong>" not in html
    assert 'alt=""' not in html
    positions = [html.index(f'alt="{title}"') for title in expected_titles[:plot_count]]
    assert positions == sorted(positions)
    assert all(note in html for note in unavailable_notes)
    view.close()


def test_series_plot_draws_the_prepared_values():
    result = time_series_snapshot().result
    ax = Figure().add_subplot(111)

    draw_series_plot(ax, result, CHART_LABELS)

    (line,) = ax.get_lines()
    assert list(line.get_ydata()) == pytest.approx(list(result.values))
    assert ax.get_title() == "Series"
    assert ax.get_ylabel() == "sales"


def test_acf_plot_draws_one_stem_per_lag():
    result = time_series_snapshot().result
    ax = Figure().add_subplot(111)

    draw_acf_plot(ax, result, CHART_LABELS)

    stems, points = ax.collections
    assert len(stems.get_segments()) == len(result.acf_values)
    assert points.get_offsets()[:, 1].tolist() == pytest.approx(list(result.acf_values))
    assert ax.get_title() == "Autocorrelation"
    assert (ax.get_xlabel(), ax.get_ylabel()) == ("Lag", "Autocorrelation")


def test_decomposition_draws_four_labeled_panels():
    result = time_series_snapshot().result
    axes = list(Figure().subplots(4, 1, sharex=True))

    draw_decomposition(axes, result, CHART_LABELS)

    assert [ax.get_ylabel() for ax in axes] == ["Observed", "Trend", "Seasonal", "Residual"]
    assert axes[0].get_title() == "Seasonal decomposition"
    assert list(axes[2].get_lines()[0].get_ydata()) == pytest.approx(list(result.seasonal))


def test_decomposition_requires_four_axes():
    with pytest.raises(ValueError, match="four axes"):
        draw_decomposition([Figure().add_subplot(111)], time_series_snapshot().result, CHART_LABELS)


def test_render_returns_every_available_plot_with_progress():
    progress: list[int] = []

    images = render_time_series_charts(time_series_snapshot(), progress_cb=progress.append, cancel_cb=lambda: False)

    assert [image.heading for image in images] == ["", "", ""]
    assert [image.accessible_name for image in images] == ["Series", "Autocorrelation", "Seasonal decomposition"]
    assert [(image.width_px, image.height_px) for image in images] == [(960, 480), (960, 480), (960, 960)]
    assert all(image.image_data.startswith(_PNG_SIGNATURE) and image.captions == () for image in images)
    assert progress == [33, 66, 100]


def test_render_skips_unavailable_plots():
    progress: list[int] = []

    (image,) = render_time_series_charts(irregular_snapshot(), progress_cb=progress.append)

    assert image.heading == ""
    assert image.accessible_name == "Series"
    assert progress == [100]


@pytest.mark.parametrize("plot_count", [1, 2, 3])
def test_optional_plot_order_progress_and_cancellation(plot_count):
    snapshot = _snapshot_with_plot_count(plot_count)
    progress: list[int] = []
    images = render_time_series_charts(snapshot, progress_cb=progress.append)
    assert [image.accessible_name for image in images] == ["Series", "Autocorrelation", "Seasonal decomposition"][
        :plot_count
    ]
    assert progress == [(index + 1) * 100 // plot_count for index in range(plot_count)]
    for cancel_after in range(plot_count + 1):
        cancelled_progress: list[int] = []
        with pytest.raises(ExcelChartCancelledError):
            render_time_series_charts(
                snapshot,
                progress_cb=cancelled_progress.append,
                cancel_cb=lambda progress=cancelled_progress, limit=cancel_after: len(progress) >= limit,
            )
        assert cancelled_progress == progress[:cancel_after]


@pytest.mark.parametrize("cancel_after", [0, 1, 2, 3])
def test_render_honors_cancellation_before_each_image_and_at_the_end(cancel_after):
    calls = 0

    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls > cancel_after

    with pytest.raises(ExcelChartCancelledError):
        render_time_series_charts(time_series_snapshot(), cancel_cb=cancel)


@pytest.mark.parametrize("streaming", [False, True])
def test_workbook_keeps_typed_tables_and_every_plot(tmp_path, streaming):
    snapshot = time_series_snapshot()
    frames = time_series_export_tables(
        snapshot,
        (
            TimeSeriesExportComponent.SERIES,
            TimeSeriesExportComponent.AUTOCORRELATION,
            TimeSeriesExportComponent.MODEL_SUMMARY,
        ),
    )
    sheets = {key.value: frame for key, frame in frames.items()}
    path = tmp_path / "time_series.xlsx"

    FileWriter().save_excel_sheets(
        sheets,
        path,
        chart_images=render_time_series_charts(snapshot),
        chart_sheet_name="Charts",
        streaming=streaming,
    )

    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["series", "autocorrelation", "model_summary", "Charts"]
        assert len(workbook["Charts"]._images) == 3
        chart_sheet = workbook["Charts"]
        assert chart_sheet["A1"].value is None
        assert all(cell.value is None for row in chart_sheet.iter_rows() for cell in row)
        anchors = [image.anchor._from.row for image in chart_sheet._images]
        assert anchors == sorted(set(anchors))
        series = workbook["series"]
        assert [cell.value for cell in series[1]] == ["timestamp", "value", "trend", "seasonal", "residual"]
        assert series["A2"].is_date
        assert series["B2"].value == pytest.approx(10.0)
        acf = workbook["autocorrelation"]
        assert [cell.value for cell in acf[1]] == ["lag", "autocorrelation"]
        assert acf["A2"].value == 1
        assert [row[1].value for row in list(series.rows)[1:]] == pytest.approx(snapshot.result.values)
        assert [row[1].value for row in list(acf.rows)[1:]] == pytest.approx(snapshot.result.acf_values)
        summary = workbook["model_summary"]
        assert summary["A2"].value == "source_rows"
        assert summary["B2"].value == pytest.approx(23.0)
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [False, True])
@pytest.mark.parametrize("plot_count", [1, 2, 3])
def test_workbook_optional_charts_have_no_redundant_titles(tmp_path, streaming, plot_count):
    snapshot = _snapshot_with_plot_count(plot_count)
    path = tmp_path / "optional_charts.xlsx"
    FileWriter().save_excel_sheets(
        {"series": series_table(snapshot)},
        path,
        chart_images=render_time_series_charts(snapshot),
        chart_sheet_name="Charts",
        streaming=streaming,
    )
    workbook = load_workbook(path)
    try:
        chart_sheet = workbook["Charts"]
        assert len(chart_sheet._images) == plot_count
        assert all(cell.value is None for row in chart_sheet.iter_rows() for cell in row)
        anchors = [image.anchor._from.row for image in chart_sheet._images]
        assert anchors == sorted(set(anchors))
        assert [row[1].value for row in list(workbook["series"].rows)[1:]] == pytest.approx(snapshot.result.values)
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [False, True])
def test_workbook_accepts_timezone_aware_series(tmp_path, streaming):
    frame = pd.DataFrame({
        "day": pd.date_range("2024-01-01", periods=4, freq="D", tz="Europe/Stockholm"),
        "sales": [1.0, 2.0, 4.0, 3.0],
    })
    snapshot = time_series_snapshot(frame)
    path = tmp_path / "time_series_tz.xlsx"

    FileWriter().save_excel_sheets({"series": series_table(snapshot)}, path, streaming=streaming)

    workbook = load_workbook(path)
    try:
        assert workbook["series"]["A2"].is_date
        assert workbook["series"]["A2"].value.isoformat().startswith("2024-01-01")
    finally:
        workbook.close()
