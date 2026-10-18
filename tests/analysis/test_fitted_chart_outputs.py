from __future__ import annotations

from io import BytesIO
from math import ceil

import pandas as pd
import pytest
from openpyxl import load_workbook
from PIL import Image

from expo_jbm329.services.analysis.chart_config import AxisPolicy, ChartConfig, RotationMode
from expo_jbm329.services.analysis.chart_export import render_chart_png
from expo_jbm329.services.analysis.chart_layout import ChartFigure
from expo_jbm329.services.analysis.report import (
    PreparedReportComponent,
    PreparedReportDocument,
    PreparedReportEntry,
    render_prepared_report_html,
)
from expo_jbm329.services.excel_chart import ExcelChartImage
from expo_jbm329.services.file_writer import FileWriter


def _fitted_image() -> ExcelChartImage:
    figure = ChartFigure(
        figsize=(6, 3),
        dpi=100,
        constrained_layout=True,
        config=ChartConfig(default_axis_policy=AxisPolicy(rotation_mode=RotationMode.VERTICAL)),
    )
    labels = [f"Group {index}: " + "long measurement name " * 3 for index in range(6)]
    axis = figure.add_subplot()
    axis.bar(range(6), range(1, 7))
    axis.set_xticks(range(6), labels)
    raster = render_chart_png(figure)
    return ExcelChartImage(
        "",
        raster.image_data,
        raster.width_px,
        raster.height_px,
        ("Sampling explanation retained.",),
        "Complete category labels",
    )


@pytest.mark.parametrize("streaming", [False, True])
def test_fitted_workbook_images_and_captions_use_actual_dimensions(tmp_path, streaming):
    image = _fitted_image()
    assert image.height_px > 300
    with Image.open(BytesIO(image.image_data)) as png:
        assert png.size == (image.width_px, image.height_px)
    path = tmp_path / "fitted.xlsx"
    data = pd.DataFrame({"value": [1.25, 2.5]})
    FileWriter().save_excel_sheets(
        {"Data": data},
        path,
        chart_images=(image, image),
        chart_sheet_name="Charts",
        streaming=streaming,
    )
    workbook = load_workbook(path)
    try:
        assert list(workbook["Data"].values) == [("value",), (1.25,), (2.5,)]
        sheet = workbook["Charts"]
        first, second = sheet._images
        assert first.width == image.width_px
        assert first.height == image.height_px
        assert first.anchor.ext.cx == image.width_px * 9525
        assert first.anchor.ext.cy == image.height_px * 9525
        assert second.anchor._from.row > first.anchor._from.row + ceil(image.height_px / 20)
        captions = [cell for row in sheet for cell in row if cell.value == image.captions[0]]
        assert len(captions) == 2
        assert captions[0].row > first.anchor._from.row + ceil(image.height_px / 20)
        assert sheet["A1"].value is None
    finally:
        workbook.close()


def test_fitted_html_keeps_image_dimensions_accessibility_and_explanation():
    image = _fitted_image()
    document = PreparedReportDocument(
        "Fitted charts",
        "",
        (
            PreparedReportEntry(
                "result",
                "Dataset",
                6,
                1,
                "Hypothesis Tests",
                "2026-10-17",
                "1",
                "",
                (PreparedReportComponent("chart", "Distribution", "chart", charts=(image,)),),
            ),
        ),
        "N/A",
    )
    html = render_prepared_report_html(document)
    assert f'width="{image.width_px}" height="{image.height_px}"' in html
    assert 'alt="Complete category labels"' in html
    assert "<p>Sampling explanation retained.</p>" in html
    assert "<strong>Complete category labels</strong>" not in html
    assert "break-inside: avoid; overflow-x: auto;" in html
    assert "height: auto; max-width: none;" in html
    assert ".report-chart img { max-width: 100%; }" in html
