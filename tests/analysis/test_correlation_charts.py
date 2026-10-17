from __future__ import annotations

from dataclasses import replace
from itertools import combinations
from unittest.mock import Mock

import numpy as np
import pandas as pd
import pytest
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PyQt6.QtCore import QLocale

from expo_jbm329.services.analysis import correlation_charts as charts
from expo_jbm329.services.analysis.chart_config import AxisPolicy, ChartConfig, RotationMode
from expo_jbm329.services.analysis.chart_export import ChartRaster
from expo_jbm329.services.analysis.chart_layout import ChartFigure
from expo_jbm329.services.analysis.chart_presentation import ChartNumericLocale
from expo_jbm329.services.analysis.correlation import (
    SCATTER_SAMPLE_SIZE,
    CorrelationError,
    CorrelationMethod,
    analyze_correlation_matrix,
    analyze_correlation_pair,
)
from expo_jbm329.services.analysis.correlation_charts import (
    CorrelationChartLabels,
    build_correlation_chart_labels,
    build_correlation_matrix_presentation,
    build_correlation_pair_presentation,
    draw_correlation_matrix,
    draw_correlation_scatter,
    render_correlation_charts,
    scatter_captions,
)
from expo_jbm329.services.analysis.correlation_export import (
    CorrelationExportComponent as Component,
)
from expo_jbm329.services.analysis.correlation_export import CorrelationExportRequest, CorrelationExportSnapshot
from expo_jbm329.services.analysis.statistics import StatisticsExportFormat
from expo_jbm329.services.excel_chart import ExcelChartCancelledError
from expo_jbm329.services.file_writer import FileWriter
from expo_jbm329.utils.format_utils import capture_chart_numeric_locale, fmt_int, fmt_num

LABELS = CorrelationChartLabels(
    "Matrix",
    "{x} / {y}",
    "Least-squares line: y = ({slope}) x + ({intercept})",
    "Descriptive linear fit, not the rank correlation estimator.",
    "Showing {shown} of {total} points. Statistics use all points.",
    "Insufficient observations.",
    "Constant input.",
)


def snapshot(method=CorrelationMethod.PEARSON, df=None):
    if df is None:
        df = pd.DataFrame({"z": [1.0, 2, 3, 4], "a": [8.0, 3, 6, 5], "b": [3.0, 5, 8, 9]})
    matrix = analyze_correlation_matrix(df, tuple(df.columns), method)
    assert matrix is not None and matrix.error is None
    return CorrelationExportSnapshot(matrix, df)


def test_matrix_keeps_vertical_labels_under_global_horizontal_policy() -> None:
    result = snapshot().matrix
    figure = ChartFigure(
        figsize=(9, 6),
        config=ChartConfig(default_axis_policy=AxisPolicy(rotation_mode=RotationMode.HORIZONTAL)),
    )
    canvas = FigureCanvasAgg(figure)
    ax = figure.add_subplot()
    draw_correlation_matrix(figure, ax, result, build_correlation_matrix_presentation(LABELS))

    canvas.draw()

    assert all(label.get_rotation() == 90 for label in ax.get_xticklabels())


@pytest.mark.parametrize("method", list(CorrelationMethod))
def test_chart_order_and_rank_method_captions_without_retesting_correlations(method, monkeypatch):
    result = snapshot(method)
    progress = []
    monkeypatch.setattr(
        "expo_jbm329.services.analysis.correlation._pair_statistics",
        Mock(side_effect=AssertionError("Correlation tests must not be repeated.")),
    )
    images = render_correlation_charts(
        result, (Component.SCATTERPLOTS, Component.MATRIX_PLOT), LABELS, progress_cb=progress.append
    )
    assert [image.accessible_name for image in images] == ["Matrix", "z / a", "z / b", "a / b"]
    assert all(image.heading == "" for image in images)
    assert all(image.image_data.startswith(b"\x89PNG") for image in images)
    assert images[0].captions == ()
    assert all("y = (" in image.captions[0] for image in images[1:])
    assert all(
        (LABELS.descriptive_line in image.captions) is (method is not CorrelationMethod.PEARSON) for image in images[1:]
    )
    assert progress == [25, 50, 75, 100]


def test_scatter_uses_deterministic_sample_but_full_cohort_fit():
    x = np.arange(SCATTER_SAMPLE_SIZE + 100, dtype=float)
    y = x * 1.5 + 2
    y[-1] += 10000
    df = pd.DataFrame({"x": x, "y": y})
    detail = analyze_correlation_pair(df, "x", "y")
    assert detail.sampled
    assert len(detail.sample_x) == SCATTER_SAMPLE_SIZE
    repeated = analyze_correlation_pair(df, "x", "y")
    assert detail.sample_x == repeated.sample_x and detail.sample_y == repeated.sample_y
    np.testing.assert_allclose((detail.slope, detail.intercept), np.polyfit(x, y, 1))
    figure = Figure()
    ax = figure.add_subplot()
    draw_correlation_scatter(ax, detail, build_correlation_pair_presentation(detail, LABELS))
    assert (ax.get_title(), ax.get_xlabel(), ax.get_ylabel()) == ("x / y", "x", "y")
    assert len(ax.collections[0].get_offsets()) == SCATTER_SAMPLE_SIZE
    np.testing.assert_allclose(
        ax.lines[0].get_ydata(),
        detail.slope * np.asarray(ax.lines[0].get_xdata()) + detail.intercept,
    )
    assert scatter_captions(detail, LABELS)[-1] == "Showing 5,000 of 5,100 points. Statistics use all points."


def test_undefined_pairs_have_explicit_placeholders_in_sequence():
    result = snapshot(df=pd.DataFrame({"x": [1.0, 2, 3, 4], "c": [1.0] * 4, "s": [1.0, 2, np.nan, np.nan]}))
    images = render_correlation_charts(result, (Component.SCATTERPLOTS,), LABELS)
    assert [image.accessible_name for image in images] == ["x / c", "x / s", "c / s"]
    assert all(image.heading == "" for image in images)
    assert [image.captions for image in images] == [
        (LABELS.constant_input,),
        (LABELS.insufficient_observations,),
        (LABELS.insufficient_observations,),
    ]
    assert all(image.image_data.startswith(b"\x89PNG") for image in images)


def test_matrix_scale_orientation_annotation_threshold_and_nan():
    result = snapshot()
    figure = Figure()
    ax = figure.add_subplot()
    draw_correlation_matrix(
        figure, ax, result.matrix, build_correlation_matrix_presentation(replace(LABELS, matrix_title="Applied"))
    )
    assert ax.get_title() == "Applied"
    assert ax.images[0].get_clim() == (-1, 1)
    np.testing.assert_allclose(ax.images[0].get_array(), result.matrix.coefficients)
    assert [tick.get_text() for tick in ax.get_xticklabels()] == list(result.matrix.columns)
    assert len(ax.texts) == 9
    assert {text.get_color() for text in ax.texts} == {"white", "black"}
    wide = snapshot(df=pd.DataFrame({f"c{i}": np.arange(4, dtype=float) for i in range(13)}))
    figure = Figure()
    ax = figure.add_subplot()
    draw_correlation_matrix(
        figure, ax, wide.matrix, build_correlation_matrix_presentation(replace(LABELS, matrix_title="Wide"))
    )
    assert len(ax.texts) == 0
    undefined = snapshot(df=pd.DataFrame({"c": [1.0] * 4, "d": [2.0] * 4}))
    ax = Figure().add_subplot()
    draw_correlation_matrix(
        ax.figure,
        ax,
        undefined.matrix,
        build_correlation_matrix_presentation(replace(LABELS, matrix_title="Undefined")),
    )
    assert len(ax.texts) == 0
    assert np.ma.getmaskarray(ax.images[0].get_array()).all()


def test_matrix_uses_captured_localized_annotations():
    result = snapshot()
    figure = Figure()
    ax = figure.add_subplot()
    annotations = tuple(tuple(f"{row},{column}" for column in range(3)) for row in range(3))
    draw_correlation_matrix(
        figure,
        ax,
        result.matrix,
        build_correlation_matrix_presentation(
            replace(LABELS, matrix_title="Localized", matrix_annotations=annotations)
        ),
    )
    assert [text.get_text() for text in ax.texts] == [value for row in annotations for value in row]


@pytest.mark.parametrize("fail", [False, True])
def test_render_releases_worker_figure_on_success_or_render_failure(monkeypatch, fail):
    cleared = []
    original = charts.Figure.clear

    def clear(figure, *args, **kwargs):
        cleared.append(figure)
        return original(figure, *args, **kwargs)

    monkeypatch.setattr(charts.Figure, "clear", clear)
    if fail:
        monkeypatch.setattr(charts, "render_chart_png", Mock(side_effect=OSError("PNG error")))
        with pytest.raises(OSError, match="PNG error"):
            render_correlation_charts(snapshot(), (Component.MATRIX_PLOT,), LABELS)
    else:
        render_correlation_charts(snapshot(), (Component.MATRIX_PLOT,), LABELS)
    assert cleared
    assert not cleared[-1].axes


@pytest.mark.parametrize("cancel_at", [1, 2, 3, 4])
def test_cancel_before_during_or_between_images_never_returns_success(cancel_at):
    count = 0
    progress = []

    def cancel():
        nonlocal count
        count += 1
        return count == cancel_at

    with pytest.raises(ExcelChartCancelledError):
        render_correlation_charts(
            snapshot(),
            (Component.MATRIX_PLOT, Component.SCATTERPLOTS),
            LABELS,
            cancel_cb=cancel,
            progress_cb=progress.append,
        )
    assert 100 not in progress


@pytest.mark.parametrize("components", [(), (Component.MATRIX_PLOT,) * 2, (Component.STRONGEST_CORRELATIONS,)])
def test_renderer_rejects_invalid_chart_selection(components):
    with pytest.raises(ValueError, match="distinct correlation chart"):
        render_correlation_charts(snapshot(), components, LABELS)


def test_unexpected_chart_error_and_invalid_scatter_propagate():
    detail = analyze_correlation_pair(pd.DataFrame({"x": [1.0] * 3, "y": [2.0] * 3}), "x", "y")
    presentation = build_correlation_pair_presentation(detail, LABELS)
    with pytest.raises(ValueError, match="defined pair"):
        draw_correlation_scatter(Figure().add_subplot(), detail, replace(presentation, unavailable_message=None))
    with pytest.raises(ValueError, match="Unexpected"):
        scatter_captions(replace(detail, error=CorrelationError.INVALID_COLUMN), LABELS)
    df = pd.DataFrame({"x": [1.0, 2, 3], "y": [2.0, 3, 5]})
    detail = analyze_correlation_pair(df, "x", "y")
    with pytest.raises(ValueError, match="match the plotted"):
        analyze_correlation_pair(df, "x", "y", statistics=replace(detail.pair, n=99))


def test_maximum_matrix_has_all_435_unique_ordered_pairs_without_rendering(monkeypatch):
    result = snapshot(df=pd.DataFrame({f"c{i}": np.arange(4, dtype=float) for i in range(30)}))
    seen = []
    original = charts.analyze_correlation_pair

    def record(df, x, y, method, *, statistics):
        seen.append((x, y))
        return original(df, x, y, method, statistics=statistics)

    monkeypatch.setattr(charts, "analyze_correlation_pair", record)
    monkeypatch.setattr(charts, "render_chart_png", lambda *_args, **_kwargs: ChartRaster(b"PNG", 1080, 720))
    render_correlation_charts(result, (Component.SCATTERPLOTS,), LABELS)
    assert seen == list(combinations(result.matrix.columns, 2))
    assert len(seen) == len(set(map(frozenset, seen))) == 435


@pytest.mark.parametrize("streaming", [True, False])
def test_real_correlation_workbook_contains_matrix_pairs_and_below_image_captions(tmp_path, streaming):
    result = snapshot(CorrelationMethod.SPEARMAN)
    images = render_correlation_charts(result, (Component.MATRIX_PLOT, Component.SCATTERPLOTS), LABELS)
    path = tmp_path / "correlations.xlsx"
    FileWriter().save_excel_sheets({}, path, chart_images=images, chart_sheet_name="Charts", streaming=streaming)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Charts"]
        assert len(sheet._images) == 4
        assert sheet["A1"].value is None
        values = [row[0].value for row in sheet.iter_rows() if row[0].value]
        assert values == [
            *[caption for image in images[1:] for caption in image.captions],
        ]
        starts = [image.anchor._from.row for image in sheet._images]
        assert all(second > first + 36 for first, second in zip(starts, starts[1:], strict=False))
    finally:
        workbook.close()


@pytest.mark.parametrize("components", [(Component.MATRIX_PLOT,), (Component.SCATTERPLOTS,), tuple(Component)])
def test_request_accepts_independent_and_combined_excel_charts(components):
    assert CorrelationExportRequest(snapshot(), components, StatisticsExportFormat.EXCEL).components == components
    for format_choice in (StatisticsExportFormat.CSV, StatisticsExportFormat.BINARY):
        with pytest.raises(ValueError, match="charts require Excel"):
            CorrelationExportRequest(snapshot(), components, format_choice)


@pytest.mark.parametrize("language", ["en_US", "sv_SE"])
def test_captured_chart_formatting_matches_app_without_worker_locale_access(language, monkeypatch):
    monkeypatch.setattr(QLocale, "system", lambda: QLocale(language))
    numeric_locale = capture_chart_numeric_locale()
    values = (float("nan"), -0.123456, 1.0, 1234.0, 0.987654, float("inf"))
    assert [numeric_locale.number(value, precision=2) for value in values] == [
        fmt_num(value, sig=2) for value in values
    ]
    assert [numeric_locale.number(value) for value in values] == [fmt_num(value) for value in values]
    assert numeric_locale.count(5100) == fmt_int(5100)
    result = snapshot()
    labels = build_correlation_chart_labels(result.matrix, lambda _context, source: source, numeric_locale)
    expected = tuple(tuple(fmt_num(value, sig=2) for value in row) for row in result.matrix.coefficients)
    assert labels.matrix_annotations == expected
    monkeypatch.setattr(QLocale, "system", Mock(side_effect=AssertionError("Worker consulted Qt locale.")))
    images = render_correlation_charts(result, (Component.MATRIX_PLOT, Component.SCATTERPLOTS), labels)
    assert all(image.accessible_name and not image.heading for image in images)
    assert len(images[1].captions) == 1


@pytest.mark.parametrize(
    ("frame", "expected_error"),
    [
        (pd.DataFrame({"x": [1.0, 2, 3], "y": [1.0] * 3}), CorrelationError.CONSTANT_INPUT),
        (pd.DataFrame({"x": [1.0, 2], "y": [1.0, 2]}), CorrelationError.NOT_ENOUGH_OBSERVATIONS),
    ],
)
@pytest.mark.parametrize("method", list(CorrelationMethod))
def test_unavailable_pair_drawer_preserves_identity_and_explicit_error(frame, expected_error, method):
    detail = analyze_correlation_pair(frame, "x", "y", method)
    assert detail.error is expected_error
    presentation = build_correlation_pair_presentation(detail, LABELS)
    ax = Figure().add_subplot()
    draw_correlation_scatter(ax, detail, presentation)
    assert (ax.get_title(), ax.get_xlabel(), ax.get_ylabel()) == ("x / y", "x", "y")
    assert [text.get_text() for text in ax.texts] == list(presentation.metadata.captions)
    assert len(ax.collections) == len(ax.lines) == 0


def test_matrix_formatting_uses_canonical_precision_not_report_table_formatter():
    result = snapshot()
    labels = build_correlation_chart_labels(
        result.matrix, lambda _context, source: source, ChartNumericLocale(decimal_separator=",", group_separator=" ")
    )
    matrix = build_correlation_matrix_presentation(labels)
    assert matrix.metadata.worksheet_heading == ""
    assert matrix.metadata.accessible_name == matrix.title
    assert labels.matrix_annotations[0][1] == "-0,37"
