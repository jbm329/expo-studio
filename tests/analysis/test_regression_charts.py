from __future__ import annotations

import dataclasses
import math
from pathlib import Path
from unittest.mock import patch

import pytest
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PyQt6.QtCore import QCoreApplication, QTranslator

from expo_jbm329.gui.dialogs.analysis.regression_glm_view import GeneralizedRegressionView
from expo_jbm329.gui.dialogs.analysis.regression_view import RegressionView
from expo_jbm329.gui.dialogs.analysis.survival_view import SurvivalRegressionView
from expo_jbm329.services.analysis.categories import AnalysisCategory
from expo_jbm329.services.analysis.regression import RegressionResult
from expo_jbm329.services.analysis.regression_charts import (
    COUNT_POINT_SIZE,
    CoxChartLabels,
    draw_actual_vs_predicted,
    draw_count_charts,
    draw_cox_charts,
    draw_cox_forest,
    draw_cox_risk_counts,
    draw_cox_survival,
    draw_linear_charts,
    draw_normal_qq,
    draw_residuals_vs_fitted,
    render_regression_charts,
)
from expo_jbm329.services.analysis.regression_export import (
    RegressionExportComponent,
    RegressionExportSnapshot,
    regression_export_tables,
)
from expo_jbm329.services.analysis.regression_glm import RegressionModel
from expo_jbm329.services.analysis.report import (
    ReportDocument,
    create_report_entry,
    render_report_html,
)
from expo_jbm329.services.analysis.survival import SurvivalResult, survival_term_name
from expo_jbm329.services.excel_chart import ExcelChartCancelledError
from expo_jbm329.services.file_writer import FileWriter
from expo_jbm329.utils.dataset_ref import DatasetRef
from tests.analysis.test_regression_export import (
    COUNT_LABELS,
    COX_LABELS,
    LINEAR_LABELS,
    cox_result,
    cox_snapshot,
    generalized_result,
    linear_result,
    linear_snapshot,
    logistic_snapshot,
    negative_binomial_snapshot,
    poisson_snapshot,
    regression_chart_labels,
)

ALL_SNAPSHOTS = [linear_snapshot, logistic_snapshot, poisson_snapshot, negative_binomial_snapshot, cox_snapshot]


@pytest.mark.parametrize("factory", ALL_SNAPSHOTS)
@pytest.mark.parametrize("language", ["en", "sv"])
def test_regression_views_and_exports_share_axes_and_hide_redundant_headings(factory, language, tmp_path) -> None:
    translator = QTranslator()
    locales = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales"
    assert translator.load(str(locales / f"app_{language}.qm"))
    assert QCoreApplication.installTranslator(translator)
    result = factory().result
    if isinstance(result, RegressionResult):
        view = RegressionView(result)
    elif isinstance(result, SurvivalResult):
        view = SurvivalRegressionView(result)
    else:
        view = GeneralizedRegressionView(result)
    try:
        labels = view.chart_labels()
        snapshot = RegressionExportSnapshot(result, labels)
        app_axes = [
            (axis.get_title(), axis.get_xlabel(), axis.get_ylabel())
            for canvas in view.findChildren(FigureCanvasQTAgg)
            for axis in canvas.figure.axes
        ]
        figure = Figure()
        with patch("expo_jbm329.services.analysis.regression_charts.Figure", return_value=figure):
            image = render_regression_charts(snapshot)[0]
        assert sorted(app_axes) == sorted(
            (axis.get_title(), axis.get_xlabel(), axis.get_ylabel()) for axis in figure.axes
        )
        assert image.heading == ""
        assert image.accessible_name == labels.metadata.accessible_name
        assert image.accessible_name
        for streaming in (False, True):
            path = tmp_path / f"diagnostics-{streaming}.xlsx"
            FileWriter().save_excel_sheets(
                {}, path, chart_images=(image,), chart_sheet_name="Charts", streaming=streaming
            )
            workbook = load_workbook(path)
            try:
                assert workbook["Charts"]["A1"].value is None
                assert len(workbook["Charts"]._images) == 1
            finally:
                workbook.close()
        document = ReportDocument()
        document.upsert(
            create_report_entry(
                AnalysisCategory.REGRESSION,
                DatasetRef("source", "Dataset", 100, 3),
                snapshot,
            ).with_selection(("charts",), "")
        )
        html = render_report_html(document)
        assert "<figcaption>" not in html
        assert 'alt=""' not in html
    finally:
        view.close()
        QCoreApplication.removeTranslator(translator)


# ----------------------------------------------------------------------
# Rendering
# ----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("factory", "size"),
    [
        (linear_snapshot, (1620, 540)),
        (poisson_snapshot, (1320, 570)),
        (negative_binomial_snapshot, (1320, 570)),
        (logistic_snapshot, (1200, 900)),
        (cox_snapshot, (1440, 720)),
    ],
)
def test_render_is_one_deterministic_png_with_progress(factory, size):
    snapshot = factory()
    progress: list[int] = []
    images = render_regression_charts(snapshot, progress_cb=progress.append)
    assert len(images) == 1
    image = images[0]
    assert image.heading == snapshot.chart_labels.metadata.worksheet_heading
    assert image.accessible_name == snapshot.chart_labels.metadata.accessible_name
    assert image.image_data.startswith(b"\x89PNG")
    assert (image.width_px, image.height_px) == size
    assert progress == [100]
    assert images == render_regression_charts(snapshot)


def test_render_keeps_captions_below_the_image():
    result = linear_result()
    snapshot = RegressionExportSnapshot(
        result, regression_chart_labels("Linear: y", LINEAR_LABELS, ("Sampled 2,000 of 9,000 rows",))
    )
    assert render_regression_charts(snapshot)[0].captions == ("Sampled 2,000 of 9,000 rows",)


def test_logistic_export_grows_with_predictor_count():
    snapshot = logistic_snapshot()
    result = snapshot.result
    many_terms = dataclasses.replace(result, terms=(result.terms[0], *(result.terms[1:] * 20)))
    labels = dataclasses.replace(
        snapshot.chart_labels.model,  # type: ignore[type-var]
        effect_annotations=tuple("effect" for _ in range(20)),
    )
    grown = RegressionExportSnapshot(many_terms, regression_chart_labels("h", labels))
    image = render_regression_charts(grown)[0]
    assert image.height_px == round((5.0 + 20 * 0.24) * 120)


@pytest.mark.parametrize("cancel_at", [1, 2])
def test_cancellation_never_reports_completed_progress(cancel_at):
    calls = 0
    progress: list[int] = []

    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls == cancel_at

    with pytest.raises(ExcelChartCancelledError):
        render_regression_charts(linear_snapshot(), cancel_cb=cancel, progress_cb=progress.append)
    assert progress == []


def test_render_rejects_models_without_chart_data():
    snapshot = RegressionExportSnapshot(
        dataclasses.replace(linear_result(), plot=None), regression_chart_labels("h", LINEAR_LABELS)
    )
    with pytest.raises(ValueError, match="no chart data"):
        render_regression_charts(snapshot)


# ----------------------------------------------------------------------
# Linear drawing
# ----------------------------------------------------------------------


def test_linear_charts_use_displayed_order_and_labels():
    plot = linear_result().plot
    assert plot is not None
    figure = Figure()
    draw_linear_charts(figure, plot, LINEAR_LABELS)
    residuals, actual, qq = figure.axes
    assert residuals.get_title() == LINEAR_LABELS.residuals_title
    assert residuals.get_xlabel() == LINEAR_LABELS.fitted_axis
    assert actual.get_title() == LINEAR_LABELS.actual_title
    assert actual.get_ylabel() == LINEAR_LABELS.actual_axis
    assert qq.get_title() == LINEAR_LABELS.qq_title
    assert qq.get_ylabel() == LINEAR_LABELS.standardized_axis
    assert len(residuals.collections[0].get_offsets()) == len(plot.fitted)


def test_linear_reference_lines_span_the_data():
    plot = linear_result().plot
    assert plot is not None
    figure = Figure()
    ax = figure.add_subplot(111)
    draw_actual_vs_predicted(ax, plot, LINEAR_LABELS)
    low, high = min((*plot.fitted, *plot.actual)), max((*plot.fitted, *plot.actual))
    assert list(ax.lines[0].get_xdata()) == [low, high]
    ax = figure.add_subplot(111)
    draw_normal_qq(ax, plot, LINEAR_LABELS)
    assert ax.lines[0].get_xdata()[0] == min(plot.qq_theoretical[0], plot.qq_sample[0])
    ax = figure.add_subplot(111)
    draw_residuals_vs_fitted(ax, plot, LINEAR_LABELS)
    assert list(ax.lines[0].get_ydata()) == [0.0, 0.0]


# ----------------------------------------------------------------------
# Count drawing
# ----------------------------------------------------------------------


def test_count_charts_compare_observed_and_residuals():
    data = generalized_result(RegressionModel.POISSON).plot_data
    assert data is not None
    figure = Figure()
    draw_count_charts(figure, data, COUNT_LABELS)
    comparison, residuals = figure.axes
    assert comparison.get_title() == COUNT_LABELS.comparison_title
    assert comparison.get_ylabel() == COUNT_LABELS.observed_axis
    assert residuals.get_ylabel() == COUNT_LABELS.pearson_axis
    assert residuals.get_xlabel() == COUNT_LABELS.fitted_axis
    assert comparison.collections[0].get_sizes()[0] == COUNT_POINT_SIZE
    assert comparison.lines[0].get_xdata()[1] == max((*data.fitted, *data.observed))


# ----------------------------------------------------------------------
# Cox drawing
# ----------------------------------------------------------------------


def test_cox_forest_lists_terms_with_locale_ticks():
    result = cox_result()
    labels = dataclasses.replace(COX_LABELS, tick_format=lambda value: f"<{value:.1f}>")
    figure = Figure()
    ax = figure.add_subplot(111)
    draw_cox_forest(ax, result.terms, labels)
    assert ax.get_xscale() == "log"
    assert [label.get_text() for label in ax.get_yticklabels()] == [survival_term_name(t) for t in result.terms]
    tick_texts = [label.get_text() for label in ax.xaxis.get_majorticklabels()]
    assert tick_texts
    assert all(text.startswith("<") and text.endswith(">") for text in tick_texts)
    assert ax.get_xlabel() == labels.hazard_axis


@pytest.mark.parametrize("bad_low", [0.0, math.nan, math.inf])
def test_cox_forest_explains_unavailable_log_scale(bad_low):
    result = cox_result()
    broken = dataclasses.replace(result.terms[0], hazard_ratio_ci_low=bad_low)
    figure = Figure()
    ax = figure.add_subplot(111)
    draw_cox_forest(ax, (broken, *result.terms[1:]), COX_LABELS)
    assert [text.get_text() for text in ax.texts] == [COX_LABELS.forest_unavailable]
    assert not ax.axison


def test_cox_forest_without_terms_is_unavailable():
    figure = Figure()
    ax = figure.add_subplot(111)
    draw_cox_forest(ax, (), COX_LABELS)
    assert [text.get_text() for text in ax.texts] == [COX_LABELS.forest_unavailable]


def test_cox_survival_and_risk_table_use_formatters():
    data = cox_result().plot_data
    assert data is not None
    labels = dataclasses.replace(COX_LABELS, time_format=lambda value: f"t{value:.0f}", count_format=lambda n: f"n{n}")
    figure = Figure()
    survival = figure.add_subplot(211)
    risk = figure.add_subplot(212)
    draw_cox_survival(survival, data, labels)
    draw_cox_risk_counts(risk, data, labels)
    assert survival.get_title() == labels.survival_title
    assert survival.get_ylim() == (0, 1.05)
    assert list(survival.get_xticks()) == list(data.risk_times)
    cells = risk.tables[0].get_celld()
    assert cells[0, 0].get_text().get_text() == labels.time_format(data.risk_times[0])
    assert cells[1, 0].get_text().get_text() == labels.count_format(data.risk_counts[0])
    assert risk.get_title() == labels.risk_title


def test_cox_survival_without_data_is_explained():
    figure = Figure()
    survival = figure.add_subplot(211)
    risk = figure.add_subplot(212)
    draw_cox_survival(survival, None, COX_LABELS)
    draw_cox_risk_counts(risk, None, COX_LABELS)
    assert [text.get_text() for text in survival.texts] == [COX_LABELS.survival_unavailable]
    assert not risk.tables


def test_cox_default_formatters_are_used_without_locale():
    labels = CoxChartLabels("f", "u", "h", "s", "n", "d", "p", "r")
    assert labels.tick_format(1.23456) == "1.23"
    assert labels.count_format(12) == "12"


def test_cox_charts_layout_has_three_panels():
    result = cox_result()
    figure = Figure()
    draw_cox_charts(figure, result, COX_LABELS)
    assert len(figure.axes) == 3
    figure = Figure()
    draw_cox_charts(figure, dataclasses.replace(result, plot_data=None), COX_LABELS)
    assert figure.axes[1].texts[0].get_text() == COX_LABELS.survival_unavailable


def test_cox_render_without_survival_data_still_exports():
    snapshot = RegressionExportSnapshot(
        dataclasses.replace(cox_result(), plot_data=None), regression_chart_labels("Cox", COX_LABELS)
    )
    assert isinstance(snapshot.result, SurvivalResult)
    assert render_regression_charts(snapshot)[0].image_data.startswith(b"\x89PNG")


# ----------------------------------------------------------------------
# Workbooks
# ----------------------------------------------------------------------


@pytest.mark.parametrize("factory", ALL_SNAPSHOTS)
@pytest.mark.parametrize("streaming", [False, True])
def test_real_workbooks_keep_typed_tables_and_one_chart_sheet(tmp_path, factory, streaming):
    snapshot = factory()
    frames = regression_export_tables(
        snapshot,
        (RegressionExportComponent.COEFFICIENTS, RegressionExportComponent.MODEL_METRICS),
    )
    sheets = {key.value: frame for key, frame in frames.items()}
    images = render_regression_charts(snapshot)
    path = tmp_path / "regression.xlsx"
    FileWriter().save_excel_sheets(sheets, path, chart_images=images, chart_sheet_name="Charts", streaming=streaming)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == [*sheets, "Charts"]
        assert len(workbook["Charts"]._images) == 1
        assert workbook["Charts"]["A1"].value == snapshot.chart_labels.metadata.worksheet_heading
        coefficients = workbook["coefficients"]
        assert [cell.value for cell in coefficients[1]] == list(frames[RegressionExportComponent.COEFFICIENTS].columns)
        assert isinstance(coefficients["B2"].value, float)
        metrics = workbook["model_metrics"]
        assert metrics["A2"].value == "rows_used"
        assert isinstance(metrics["B2"].value, (float, int))
    finally:
        workbook.close()
