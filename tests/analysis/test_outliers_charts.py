from __future__ import annotations

import dataclasses
import math
from html import escape
from pathlib import Path
from typing import TYPE_CHECKING, Any
from unittest.mock import patch

import pytest
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
from openpyxl import load_workbook
from PyQt6.QtCore import QCoreApplication, QEvent, QTranslator
from PyQt6.QtWidgets import QLabel

from expo_jbm329.gui.dialogs.analysis.outliers_multivariate_view import MultivariateOutliersView
from expo_jbm329.gui.dialogs.analysis.outliers_view import OutliersView
from expo_jbm329.services.analysis.categories import AnalysisCategory
from expo_jbm329.services.analysis.chart_layout import ChartFigure as Figure
from expo_jbm329.services.analysis.multivariate_outliers import MultivariateOutlierMethod
from expo_jbm329.services.analysis.outliers import OutlierMethod, analyze_outlier_column, analyze_outlier_summary
from expo_jbm329.services.analysis.outliers_charts import (
    build_outlier_histogram_presentation,
    draw_outlier_histogram,
    draw_outlier_projection,
    render_multivariate_outlier_charts,
    render_univariate_outlier_charts,
)
from expo_jbm329.services.analysis.outliers_export import (
    MultivariateOutlierExportSnapshot,
    OutlierExportComponent,
    UnivariateOutlierExportSnapshot,
    outlier_export_tables,
)
from expo_jbm329.services.analysis.report import (
    ReportDocument,
    capture_report_labels,
    create_report_entry,
    prepare_report_content,
    render_prepared_report_html,
)
from expo_jbm329.services.excel_chart import ExcelChartCancelledError
from expo_jbm329.services.file_writer import FileWriter
from expo_jbm329.utils.dataset_ref import DatasetRef
from tests.analysis.test_outliers_export import (
    HISTOGRAM_LABELS,
    PROJECTION_LABELS,
    multivariate_snapshot,
    univariate_detail,
    univariate_snapshot,
)

if TYPE_CHECKING:
    from collections.abc import Iterator

    from matplotlib.axes import Axes

_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


def _axes():
    return Figure().add_subplot(111)


def _semantics(axis: Axes) -> tuple[str, str, str, tuple[str, ...], tuple[str, ...]]:
    legend = axis.get_legend()
    return (
        axis.get_title(),
        axis.get_xlabel(),
        axis.get_ylabel(),
        () if legend is None else tuple(text.get_text() for text in legend.get_texts()),
        tuple(text.get_text() for text in axis.texts),
    )


@pytest.fixture(params=["en", "sv"])
def chart_language(request: pytest.FixtureRequest) -> Iterator[str]:
    language = request.param
    translator = QTranslator()
    locales = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales"
    assert translator.load(str(locales / f"app_{language}.qm"))
    assert QCoreApplication.installTranslator(translator)
    try:
        yield language
    finally:
        QCoreApplication.removeTranslator(translator)


@pytest.mark.parametrize("method", list(OutlierMethod))
def test_univariate_app_worker_and_html_share_captured_semantics_for_every_column(
    method: OutlierMethod, chart_language: str
) -> None:
    source = univariate_snapshot(method, column="b")
    view = OutliersView(source.summary, source.detail)
    snapshot = dataclasses.replace(source, chart_labels=view.histogram_labels())
    app_semantics = []
    for column in snapshot.chart_columns:
        view.set_column_detail(analyze_outlier_column(snapshot.data, column, method, snapshot.summary.threshold))
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete.value)
        canvas = view.findChildren(FigureCanvasQTAgg)[0]
        app_semantics.append(_semantics(canvas.figure.axes[0]))
    figures: list[Figure] = []

    def figure_factory(*args: Any, **kwargs: Any) -> Figure:
        figure = Figure(*args, **kwargs)
        figures.append(figure)
        return figure

    try:
        with (
            patch("expo_jbm329.services.analysis.outliers_charts.Figure", side_effect=figure_factory),
            patch.object(QCoreApplication, "translate", side_effect=AssertionError("Worker must not translate")),
        ):
            images = render_univariate_outlier_charts(snapshot)
        assert [_semantics(figure.axes[0]) for figure in figures] == app_semantics
        assert [image.accessible_name for image in images] == list(snapshot.chart_columns)
        assert all(image.heading == "" and image.captions == () for image in images)
        assert [(entry[0], entry[1]) for entry in app_semantics] == [(name, name) for name in snapshot.chart_columns]
        assert all(entry[2] == ("Antal" if chart_language == "sv" else "Count") for entry in app_semantics)
        assert app_semantics[0][3] == (
            ("Inom gränserna", "Potentiella avvikare")
            if chart_language == "sv"
            else ("Inside the fences", "Potential outliers")
        )
        assert not any(label.text() == "Column distribution" for label in view.findChildren(QLabel))
        document = ReportDocument()
        document.upsert(create_report_entry(AnalysisCategory.OUTLIERS, DatasetRef("outliers", "Data", 10, 4), snapshot))
        with patch("expo_jbm329.services.analysis.report.render_univariate_outlier_charts", return_value=images):
            prepared = prepare_report_content(document, labels=capture_report_labels())
        html = render_prepared_report_html(prepared)
        assert "<figcaption>" not in html
        for name in snapshot.chart_columns:
            assert f'alt="{escape(name)}"' in html
    finally:
        view.close()


@pytest.mark.parametrize("method", list(MultivariateOutlierMethod))
@pytest.mark.parametrize("sampled", [False, True])
@pytest.mark.parametrize("streaming", [False, True])
def test_multivariate_app_worker_and_html_share_captured_semantics(
    method: MultivariateOutlierMethod, sampled: bool, chart_language: str, streaming: bool, tmp_path: Path
) -> None:
    result = dataclasses.replace(multivariate_snapshot(method).result, sampled=sampled)
    view = MultivariateOutliersView(result)
    labels = view.chart_labels()
    snapshot = MultivariateOutlierExportSnapshot(result, labels, view.export_notes())
    app_axis = view.findChildren(FigureCanvasQTAgg)[0].figure.axes[0]
    figure = Figure()
    try:
        with (
            patch("expo_jbm329.services.analysis.outliers_charts.Figure", return_value=figure),
            patch.object(QCoreApplication, "translate", side_effect=AssertionError("Worker must not translate")),
        ):
            (image,) = render_multivariate_outlier_charts(snapshot)
        assert _semantics(figure.axes[0]) == _semantics(app_axis)
        assert (app_axis.get_xlabel(), app_axis.get_ylabel()) == ("PC1", "PC2")
        assert image.accessible_name == (
            "PCA-projektion (endast visualisering)" if chart_language == "sv" else "PCA projection (visualization only)"
        )
        assert image.heading == ""
        expected_caption = (
            "Projektionen visar ett deterministiskt urval av 7 rader."
            if chart_language == "sv"
            else "The projection shows a deterministic sample of 7 rows."
        )
        assert image.captions == ((expected_caption,) if sampled else ())
        if sampled:
            assert any(label.text() == expected_caption for label in view.findChildren(QLabel))
            assert expected_caption in "\n".join(snapshot.notes)
        document = ReportDocument()
        document.upsert(create_report_entry(AnalysisCategory.OUTLIERS, DatasetRef("outliers", "Data", 7, 3), snapshot))
        with patch("expo_jbm329.services.analysis.report.render_multivariate_outlier_charts", return_value=(image,)):
            prepared = prepare_report_content(document, labels=capture_report_labels())
        html = render_prepared_report_html(prepared)
        assert f'alt="{escape(image.accessible_name)}"' in html
        assert f"<figcaption>{escape(image.accessible_name)}" not in html
        assert ("<figcaption>" in html) is sampled
        if sampled:
            assert f"<figcaption>\n<p>{expected_caption}</p>\n</figcaption>" in html
        path = tmp_path / "projection.xlsx"
        FileWriter().save_excel_sheets(
            {
                "numeric": outlier_export_tables(snapshot, (OutlierExportComponent.MODEL_SUMMARY,))[
                    OutlierExportComponent.MODEL_SUMMARY
                ]
            },
            path,
            chart_images=(image,),
            chart_sheet_name="Projection",
            streaming=streaming,
        )
        workbook = load_workbook(path)
        try:
            assert len(workbook["Projection"]._images) == 1
            values = [cell.value for row in workbook["Projection"] for cell in row if cell.value is not None]
            assert values == ([expected_caption] if sampled else [])
            assert workbook["numeric"]["A2"].value == "features"
            assert workbook["numeric"]["B2"].value == 2
        finally:
            workbook.close()
    finally:
        view.close()


@pytest.mark.parametrize("column", ["", "<sensor & reading>"])
def test_histogram_preserves_actual_column_axes_and_accessible_identity(column: str, chart_language: str) -> None:
    source = univariate_snapshot()
    data = source.data.rename(columns={"a": column})
    summary = analyze_outlier_summary(data)
    detail = analyze_outlier_column(data, column)
    view = OutliersView(summary, detail)
    snapshot = UnivariateOutlierExportSnapshot(summary, data, view.histogram_labels(), detail=detail)
    try:
        axis = view.findChildren(FigureCanvasQTAgg)[0].figure.axes[0]
        assert (axis.get_title(), axis.get_xlabel()) == (column, column)
        images = render_univariate_outlier_charts(snapshot)
        expected_name = column or ("Kolumnfördelning" if chart_language == "sv" else "Column distribution")
        assert images[0].accessible_name == expected_name
        document = ReportDocument()
        document.upsert(create_report_entry(AnalysisCategory.OUTLIERS, DatasetRef("outliers", "Data", 10, 4), snapshot))
        with patch("expo_jbm329.services.analysis.report.render_univariate_outlier_charts", return_value=images):
            prepared = prepare_report_content(document, labels=capture_report_labels())
        html = render_prepared_report_html(prepared)
        assert f'alt="{escape(expected_name)}"' in html
        assert "<figcaption>" not in html
    finally:
        view.close()


def test_captured_outlier_labels_survive_an_application_language_change(chart_language: str) -> None:
    univariate = univariate_snapshot()
    univariate_view = OutliersView(univariate.summary, univariate.detail)
    univariate_labels = univariate_view.histogram_labels()
    univariate = dataclasses.replace(univariate, chart_labels=univariate_labels)
    multivariate = multivariate_snapshot()
    multivariate_view = MultivariateOutliersView(dataclasses.replace(multivariate.result, sampled=True))
    projection_labels = multivariate_view.chart_labels()
    multivariate = dataclasses.replace(multivariate, chart_labels=projection_labels)
    translator = QTranslator()
    locales = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales"
    other_language = "en" if chart_language == "sv" else "sv"
    assert translator.load(str(locales / f"app_{other_language}.qm"))
    assert QCoreApplication.installTranslator(translator)
    figure = Figure()
    try:
        assert univariate_view.histogram_labels() is univariate_labels
        assert multivariate_view.chart_labels() is projection_labels
        presentation = build_outlier_histogram_presentation("a", univariate.chart_labels)
        assert univariate.detail is not None
        draw_outlier_histogram(figure.add_subplot(111), univariate.detail, presentation)
        assert _semantics(figure.axes[0]) == _semantics(
            univariate_view.findChildren(FigureCanvasQTAgg)[0].figure.axes[0]
        )
        figure.clear()
        with patch("expo_jbm329.services.analysis.outliers_charts.Figure", return_value=figure):
            (image,) = render_multivariate_outlier_charts(multivariate)
        assert _semantics(figure.axes[0]) == _semantics(
            multivariate_view.findChildren(FigureCanvasQTAgg)[0].figure.axes[0]
        )
        assert image.captions == projection_labels.metadata.captions
    finally:
        QCoreApplication.removeTranslator(translator)
        univariate_view.close()
        multivariate_view.close()


# ----------------------------------------------------------------------
# Drawing
# ----------------------------------------------------------------------


def test_histogram_stacks_flagged_bars_and_marks_both_fences():
    detail = univariate_detail()
    ax = _axes()

    draw_outlier_histogram(ax, detail, build_outlier_histogram_presentation(detail.summary.column, HISTOGRAM_LABELS))

    bins = len(detail.histogram_edges) - 1
    assert len(ax.patches) == 2 * bins
    assert [patch.get_height() for patch in ax.patches[bins:]] == list(detail.outlier_counts)
    assert [patch.get_y() for patch in ax.patches[bins:]] == list(detail.inlier_counts)
    fences = [line.get_xdata()[0] for line in ax.get_lines()]
    assert fences == pytest.approx([detail.summary.lower_fence, detail.summary.upper_fence])
    assert ax.get_title() == "a"
    assert ax.get_xlabel() == "a"
    assert ax.get_ylabel() == "Count"
    assert [text.get_text() for text in ax.get_legend().get_texts()] == ["Inside", "Flagged"]


def test_histogram_of_a_column_without_values_shows_only_its_title():
    detail = univariate_detail("few")
    empty = dataclasses.replace(detail, histogram_edges=())
    ax = _axes()

    draw_outlier_histogram(ax, empty, build_outlier_histogram_presentation(empty.summary.column, HISTOGRAM_LABELS))

    assert ax.get_title() == "few"
    assert len(ax.patches) == 0
    assert ax.get_legend() is None


def test_histogram_skips_undefined_fences():
    detail = univariate_detail("few")
    ax = _axes()

    draw_outlier_histogram(ax, detail, build_outlier_histogram_presentation(detail.summary.column, HISTOGRAM_LABELS))

    assert math.isnan(detail.summary.lower_fence)
    assert len(ax.get_lines()) == 0


def test_projection_colors_flagged_rows_and_states_it_is_a_visualization():
    snapshot = multivariate_snapshot()
    result = snapshot.result
    ax = _axes()

    draw_outlier_projection(ax, result, PROJECTION_LABELS)

    inliers, outliers = ax.collections
    assert len(outliers.get_offsets()) == sum(result.sample_outliers)
    assert len(inliers.get_offsets()) == len(result.sample_outliers) - sum(result.sample_outliers)
    assert ax.get_title() == "PCA projection (visualization only)"
    assert (ax.get_xlabel(), ax.get_ylabel()) == ("PC1", "PC2")
    assert [text.get_text() for text in ax.get_legend().get_texts()] == ["Inliers", "Potential outliers"]


# ----------------------------------------------------------------------
# Univariate rendering
# ----------------------------------------------------------------------


def test_univariate_render_returns_one_histogram_per_screened_column_with_progress():
    snapshot = univariate_snapshot()
    progress: list[int] = []

    images = render_univariate_outlier_charts(snapshot, progress_cb=progress.append, cancel_cb=lambda: False)

    assert [image.accessible_name for image in images] == list(snapshot.chart_columns)
    assert all(image.heading == "" for image in images)
    assert all(image.image_data.startswith(_PNG_SIGNATURE) for image in images)
    assert {(image.width_px, image.height_px) for image in images} == {(960, 450)}
    assert progress == [34, 67, 100]


@pytest.mark.parametrize("cancel_after", [0, 1, 3])
def test_univariate_render_honors_cancellation_before_each_column_and_at_the_end(cancel_after):
    calls = 0

    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls > cancel_after

    with pytest.raises(ExcelChartCancelledError):
        render_univariate_outlier_charts(univariate_snapshot(), cancel_cb=cancel)


def test_univariate_render_requires_a_screened_column(monkeypatch):
    snapshot = univariate_snapshot()
    monkeypatch.setattr(type(snapshot), "chart_columns", property(lambda _self: ()))

    with pytest.raises(ValueError, match="no screened column"):
        render_univariate_outlier_charts(snapshot)


def test_univariate_render_rejects_data_that_can_no_longer_be_screened():
    snapshot = univariate_snapshot()
    snapshot.data["a"] = "text"

    with pytest.raises(ValueError, match="could not be screened"):
        render_univariate_outlier_charts(snapshot)


# ----------------------------------------------------------------------
# Multivariate rendering
# ----------------------------------------------------------------------


def test_multivariate_render_returns_the_projection_with_its_caption():
    progress: list[int] = []

    (image,) = render_multivariate_outlier_charts(
        multivariate_snapshot(), progress_cb=progress.append, cancel_cb=lambda: False
    )

    assert image.heading == ""
    assert image.accessible_name == "PCA projection (visualization only)"
    assert image.captions == ("Sampled",)
    assert image.image_data.startswith(_PNG_SIGNATURE)
    assert (image.width_px, image.height_px) == (960, 720)
    assert progress == [100]


@pytest.mark.parametrize("cancel_after", [0, 1])
def test_multivariate_render_honors_cancellation_before_and_after_drawing(cancel_after):
    calls = 0

    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls > cancel_after

    with pytest.raises(ExcelChartCancelledError):
        render_multivariate_outlier_charts(multivariate_snapshot(), cancel_cb=cancel)


# ----------------------------------------------------------------------
# Workbooks
# ----------------------------------------------------------------------


@pytest.mark.parametrize("streaming", [False, True])
def test_univariate_workbook_keeps_typed_tables_and_one_chart_per_column(tmp_path, streaming):
    snapshot = univariate_snapshot()
    frames = outlier_export_tables(
        snapshot,
        (
            OutlierExportComponent.COLUMN_SUMMARY,
            OutlierExportComponent.COLUMN_DETAILS,
            OutlierExportComponent.POTENTIAL_OUTLIERS,
        ),
    )
    sheets = {key.value: frame for key, frame in frames.items()}
    images = render_univariate_outlier_charts(snapshot)
    path = tmp_path / "outliers.xlsx"

    FileWriter().save_excel_sheets(
        sheets,
        path,
        chart_images=images,
        chart_sheet_name="Column distributions",
        streaming=streaming,
    )

    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == [*sheets, "Column distributions"]
        assert len(workbook["Column distributions"]._images) == len(snapshot.chart_columns)
        assert workbook["Column distributions"]["A1"].value is None
        assert not [cell.value for row in workbook["Column distributions"] for cell in row if cell.value is not None]
        summary = workbook["column_summary"]
        assert [cell.value for cell in summary[1]] == list(frames[OutlierExportComponent.COLUMN_SUMMARY].columns)
        assert summary["A2"].value == "a"
        assert isinstance(summary["E2"].value, float)
        details = workbook["column_details"]
        assert details["A2"].value == "threshold"
        assert details["B2"].value == pytest.approx(1.5)
        rows = workbook["potential_outliers"]
        assert [cell.value for cell in rows[1]][:3] == ["Row (2)", "a", "Score"]
        assert rows["A2"].value == 10
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [False, True])
def test_multivariate_workbook_keeps_model_summary_and_projection(tmp_path, streaming):
    snapshot = multivariate_snapshot()
    frames = outlier_export_tables(snapshot, (OutlierExportComponent.MODEL_SUMMARY,))
    sheets = {key.value: frame for key, frame in frames.items()}
    path = tmp_path / "multivariate.xlsx"

    FileWriter().save_excel_sheets(
        sheets,
        path,
        chart_images=render_multivariate_outlier_charts(snapshot),
        chart_sheet_name="Outlier projection",
        streaming=streaming,
    )

    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["model_summary", "Outlier projection"]
        assert len(workbook["Outlier projection"]._images) == 1
        assert workbook["Outlier projection"]["A1"].value is None
        assert [cell.value for row in workbook["Outlier projection"] for cell in row if cell.value is not None] == [
            "Sampled"
        ]
        assert workbook["model_summary"]["A2"].value == "features"
    finally:
        workbook.close()
