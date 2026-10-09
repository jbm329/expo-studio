from __future__ import annotations

import dataclasses
from pathlib import Path
from unittest.mock import Mock

import numpy as np
import pandas as pd
import pytest
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
from openpyxl import load_workbook
from PyQt6.QtCore import Qt, QTranslator
from PyQt6.QtWidgets import QApplication, QComboBox, QDialog, QLabel, QTableWidget, QWidget

from expo_jbm329.gui.dialogs.analysis.analysis_dialog import AnalysisDialog
from expo_jbm329.gui.dialogs.analysis.analysis_export_dialog import AnalysisExportDialog
from expo_jbm329.gui.dialogs.analysis.chi_square_view import ChiSquareView
from expo_jbm329.gui.dialogs.analysis.clustering_config import ClusteringConfigWidget
from expo_jbm329.gui.dialogs.analysis.clustering_view import ClusteringView
from expo_jbm329.gui.dialogs.analysis.correlation_config import CorrelationConfigWidget
from expo_jbm329.gui.dialogs.analysis.correlation_view import CorrelationView
from expo_jbm329.gui.dialogs.analysis.group_comparison_view import GroupComparisonView
from expo_jbm329.gui.dialogs.analysis.hypothesis_tests_config import HypothesisTestsConfigWidget
from expo_jbm329.gui.dialogs.analysis.outliers_config import OutliersConfigWidget
from expo_jbm329.gui.dialogs.analysis.outliers_multivariate_config import MultivariateOutliersConfigWidget
from expo_jbm329.gui.dialogs.analysis.outliers_multivariate_view import MultivariateOutliersView
from expo_jbm329.gui.dialogs.analysis.outliers_view import OutliersView
from expo_jbm329.gui.dialogs.analysis.overview_view import OverviewView
from expo_jbm329.gui.dialogs.analysis.paired_comparison_view import PairedComparisonView
from expo_jbm329.gui.dialogs.analysis.pca_config import PCAConfigWidget
from expo_jbm329.gui.dialogs.analysis.pca_view import PCAView
from expo_jbm329.gui.dialogs.analysis.regression_config import RegressionConfigWidget
from expo_jbm329.gui.dialogs.analysis.regression_glm_view import GeneralizedRegressionView
from expo_jbm329.gui.dialogs.analysis.regression_view import RegressionView
from expo_jbm329.gui.dialogs.analysis.statistics_config import StatisticsConfigWidget
from expo_jbm329.gui.dialogs.analysis.statistics_view import StatisticsView
from expo_jbm329.gui.dialogs.analysis.survival_view import SurvivalRegressionView
from expo_jbm329.gui.dialogs.analysis.timeseries_config import TimeSeriesConfigWidget
from expo_jbm329.gui.dialogs.analysis.timeseries_view import TimeSeriesView
from expo_jbm329.services.analysis.categories import AnalysisCategory, HypothesisTest
from expo_jbm329.services.analysis.clustering import ClusteringMethod
from expo_jbm329.services.analysis.correlation import (
    CorrelationError,
    CorrelationMethod,
    analyze_correlation_pair,
)
from expo_jbm329.services.analysis.correlation_export import (
    CorrelationExportComponent,
    CorrelationExportRequest,
)
from expo_jbm329.services.analysis.hypothesis_export import (
    HypothesisExportComponent as Component,
)
from expo_jbm329.services.analysis.hypothesis_export import HypothesisExportRequest
from expo_jbm329.services.analysis.multivariate_outliers import MultivariateOutlierError, MultivariateOutlierMethod
from expo_jbm329.services.analysis.outliers import OutlierMethod
from expo_jbm329.services.analysis.overview import (
    OverviewExportFormat,
    OverviewExportRequest,
    OverviewExportTable,
    analyze_dataset_overview,
)
from expo_jbm329.services.analysis.pca import PCAError
from expo_jbm329.services.analysis.regression import RegressionError
from expo_jbm329.services.analysis.regression_glm import RegressionModel
from expo_jbm329.services.analysis.statistics import (
    StatisticsExportFormat,
    StatisticsExportRequest,
    StatisticsExportTable,
    analyze_descriptive_statistics,
)
from expo_jbm329.services.analysis.timeseries import DecompositionModel
from expo_jbm329.services.file_writer import FileWriter
from expo_jbm329.utils.dataset_ref import DatasetRef
from expo_jbm329.workbench.controllers.analysis import analysis_controller as analysis_controller_module
from expo_jbm329.workbench.controllers.analysis.analysis_controller import AnalysisController
from tests.analysis.test_hypothesis_export import chi_snapshot, group_snapshot, paired_snapshot

_CONTROLLER_MODULE = "expo_jbm329.workbench.controllers.analysis.analysis_controller"


def _overview_export_context():
    result = analyze_dataset_overview(pd.DataFrame({"value": [1, 2]}))
    dialog = AnalysisDialog(
        parent=None,
        datasets=[DatasetRef("chosen", "Chosen", 2, 1), DatasetRef("other", "Other", 2, 1)],
        active_tab_id="chosen",
    )
    dialog.set_content_widget(OverviewView(result))
    dialog.set_exportable_overview(result)
    results = Mock()
    results.get_df_by_tab_id.side_effect = AssertionError("Export must not look up a dataset")
    results.current_df.side_effect = AssertionError("Export must not use the active tab")
    controller = AnalysisController(results=results, async_ops=DummyAsyncOps())
    return controller, dialog, result, Mock()


@pytest.mark.parametrize("format_choice", list(OverviewExportFormat))
def test_overview_export_routes_explicit_frames_and_never_reads_active_data(format_choice):
    controller, dialog, result, exporter = _overview_export_context()
    request = OverviewExportRequest(result, (OverviewExportTable.SAMPLE,), format_choice)
    controller._export_overview(dialog, exporter, request)
    expected = pd.DataFrame({"value": [1, 2]})
    method = {
        OverviewExportFormat.CSV: exporter.export_csv,
        OverviewExportFormat.EXCEL: exporter.export_excel,
        OverviewExportFormat.BINARY: exporter.export_data,
    }[format_choice]
    assert method.call_count == 1
    kwargs = method.call_args.kwargs
    assert kwargs["parent_widget"] is dialog
    assert kwargs["operation_target"] is dialog.content_panel()
    assert kwargs["show_success_dialog"] is True
    if format_choice is OverviewExportFormat.EXCEL:
        assert set(kwargs) == {"sheets", "parent_widget", "operation_target", "show_success_dialog"}
        assert list(kwargs["sheets"]) == ["Sample"]
        actual = kwargs["sheets"]["Sample"]
    else:
        assert set(kwargs) == {"df", "parent_widget", "operation_target", "show_success_dialog"}
        actual = kwargs["df"]
    pd.testing.assert_frame_equal(actual, expected)
    assert controller._results.mock_calls == []


def test_overview_excel_routes_all_three_named_tables_in_one_operation():
    controller, dialog, result, exporter = _overview_export_context()
    request = OverviewExportRequest(
        result,
        (OverviewExportTable.COLUMNS, OverviewExportTable.SAMPLE, OverviewExportTable.SUMMARY),
        OverviewExportFormat.EXCEL,
    )
    controller._export_overview(dialog, exporter, request)
    assert len(exporter.mock_calls) == 1
    sheets = exporter.export_excel.call_args.kwargs["sheets"]
    assert exporter.export_excel.call_args.kwargs["show_success_dialog"] is True
    assert list(sheets) == ["Columns", "Sample", "Summary"]
    assert list(sheets["Columns"].columns) == [
        "Column",
        "Type",
        "Storage type",
        "Missing",
        "Missing fraction",
        "Unique",
    ]
    assert sheets["Columns"].iloc[0]["Missing fraction"] == 0.0
    assert sheets["Sample"].shape == (2, 1)
    assert list(sheets["Sample"].columns) == ["value"]
    assert sheets["Summary"].iloc[0]["Metric"] == "row_count"


def test_overview_excel_uses_translated_names_and_headers_without_translating_values(monkeypatch):
    controller, dialog, _, exporter = _overview_export_context()
    translations = {
        "Columns": "Kolumner",
        "Summary": "Sammanfattning",
        "Sample": "Urval",
        "Column": "Kolumn",
        "Type": "Typ",
        "Storage type": "Lagringstyp",
        "Missing": "Saknas",
        "Missing fraction": "Andel saknade värden",
        "Unique": "Unika",
        "Section": "Avsnitt",
        "Metric": "Mått",
        "Count": "Antal",
        "Fraction": "Andel",
    }
    monkeypatch.setattr(analysis_controller_module, "tr", lambda _context, text: translations.get(text, text))
    request = OverviewExportRequest(
        dialog.exportable_overview(),
        (OverviewExportTable.COLUMNS, OverviewExportTable.SAMPLE, OverviewExportTable.SUMMARY),
        OverviewExportFormat.EXCEL,
    )

    controller._export_overview(dialog, exporter, request)
    sheets = exporter.export_excel.call_args.kwargs["sheets"]

    assert list(sheets) == ["Kolumner", "Urval", "Sammanfattning"]
    assert list(sheets["Kolumner"].columns) == [
        "Kolumn",
        "Typ",
        "Lagringstyp",
        "Saknas",
        "Andel saknade värden",
        "Unika",
    ]
    assert list(sheets["Sammanfattning"].columns) == ["Avsnitt", "Mått", "Kolumn", "Antal", "Andel"]
    assert sheets["Sammanfattning"].iloc[0]["Mått"] == "row_count"
    assert list(sheets["Urval"].columns) == ["value"]
    assert sheets["Urval"]["value"].dtype == "int64"


@pytest.mark.parametrize("format_choice", [OverviewExportFormat.CSV, OverviewExportFormat.BINARY])
def test_overview_summary_routes_translated_headers_for_csv_and_binary(format_choice, monkeypatch):
    controller, dialog, result, exporter = _overview_export_context()
    translations = {
        "Section": "Avsnitt",
        "Metric": "Mått",
        "Column": "Kolumn",
        "Count": "Antal",
        "Fraction": "Andel",
    }
    monkeypatch.setattr(analysis_controller_module, "tr", lambda _context, text: translations.get(text, text))
    request = OverviewExportRequest(result, (OverviewExportTable.SUMMARY,), format_choice)
    controller._export_overview(dialog, exporter, request)

    method = exporter.export_csv if format_choice is OverviewExportFormat.CSV else exporter.export_data
    assert method.call_count == 1
    assert list(method.call_args.kwargs["df"].columns) == ["Avsnitt", "Mått", "Kolumn", "Antal", "Andel"]
    assert method.call_args.kwargs["df"].iloc[0]["Mått"] == "row_count"
    assert method.call_args.kwargs["parent_widget"] is dialog
    assert method.call_args.kwargs["operation_target"] is dialog.content_panel()
    assert method.call_args.kwargs["show_success_dialog"] is True


@pytest.mark.parametrize("format_choice", [OverviewExportFormat.CSV, OverviewExportFormat.BINARY])
def test_overview_multi_table_single_file_request_is_rejected(format_choice):
    controller, dialog, result, exporter = _overview_export_context()
    request = OverviewExportRequest(result, (OverviewExportTable.COLUMNS, OverviewExportTable.SAMPLE), format_choice)
    controller._export_overview(dialog, exporter, request)
    assert exporter.mock_calls == []


@pytest.mark.parametrize("invalidate", ["category", "dataset", "placeholder", "replacement", "identity"])
def test_overview_export_rejects_stale_or_unrelated_snapshot(invalidate):
    controller, dialog, result, exporter = _overview_export_context()
    request = OverviewExportRequest(result, (OverviewExportTable.SAMPLE,), OverviewExportFormat.CSV)
    match invalidate:
        case "category":
            dialog.select_category(AnalysisCategory.STATISTICS)
        case "dataset":
            dialog.select_dataset("other")
        case "placeholder":
            dialog.show_placeholder("Analysis failed")
        case "replacement":
            dialog.set_content_widget(QWidget())
        case "identity":
            dialog.set_exportable_overview(analyze_dataset_overview(pd.DataFrame({"value": [9]})))
    controller._export_overview(dialog, exporter, request)
    assert exporter.mock_calls == []


def _statistics_export_context():
    result = analyze_descriptive_statistics(
        pd.DataFrame({
            "First numeric": [1.0, 2.0, 3.0, 4.0, 5.0],
            "Second numeric": [5.0, 4.0, 3.0, 2.0, 1.0],
            "Group": ["A", "A", "B", "B", None],
        })
    )
    dialog = AnalysisDialog(
        parent=None,
        datasets=[DatasetRef("chosen", "Chosen", 5, 3), DatasetRef("other", "Other", 5, 3)],
        active_tab_id="chosen",
    )
    dialog.select_category(AnalysisCategory.STATISTICS)
    dialog.set_content_widget(StatisticsView(result))
    dialog.set_exportable_statistics(result)
    results = Mock()
    results.get_df_by_tab_id.side_effect = AssertionError("Statistics export must use its displayed snapshot")
    results.current_df.side_effect = AssertionError("Statistics export must not read the active dataset")
    controller = AnalysisController(results=results, async_ops=DummyAsyncOps())
    return controller, dialog, result, Mock()


def _hypothesis_export_context(snapshot):
    dialog = AnalysisDialog(
        parent=None,
        datasets=[DatasetRef("chosen", "Chosen", 20, 4), DatasetRef("other", "Other", 20, 4)],
        active_tab_id="chosen",
    )
    dialog.select_category(AnalysisCategory.HYPOTHESIS_TESTS)
    content = AnalysisController._render_hypothesis_test_content(snapshot.test, snapshot.result)
    dialog.set_content_widget(content)
    snapshot = dataclasses.replace(
        snapshot,
        notes=tuple(note.replace("<b>", "").replace("</b>", "") for note in content.export_notes(snapshot.result)),
    )
    dialog.set_exportable_hypothesis(snapshot)
    results = Mock()
    results.get_df_by_tab_id.side_effect = AssertionError("Export must never fetch an active or pending dataset")
    results.current_df.side_effect = AssertionError("Export must never fetch active data")
    controller = AnalysisController(results=results, async_ops=DummyAsyncOps())
    return controller, dialog, snapshot, Mock()


@pytest.mark.parametrize("factory", [group_snapshot, chi_snapshot, paired_snapshot])
@pytest.mark.parametrize("format_choice", list(StatisticsExportFormat))
def test_hypothesis_exports_route_owned_numeric_frames_and_per_call_context(factory, format_choice):
    controller, dialog, snapshot, exporter = _hypothesis_export_context(factory())
    request = HypothesisExportRequest(snapshot, (Component.TEST_RESULTS,), format_choice)
    controller._export_hypothesis(dialog, exporter, request)
    method = {
        StatisticsExportFormat.EXCEL: exporter.export_excel,
        StatisticsExportFormat.CSV: exporter.export_csv,
        StatisticsExportFormat.BINARY: exporter.export_data,
    }[format_choice]
    assert method.call_count == 1
    kwargs = method.call_args.kwargs
    assert kwargs["parent_widget"] is dialog
    assert kwargs["operation_target"] is dialog.content_panel()
    assert kwargs["show_success_dialog"]
    frame = next(iter(kwargs["sheets"].values())) if format_choice is StatisticsExportFormat.EXCEL else kwargs["df"]
    assert frame["Test"].notna().any()
    assert pd.api.types.is_numeric_dtype(frame["Test statistic"])
    assert frame["Notes"].dropna().tolist() == list(snapshot.notes)
    if format_choice is StatisticsExportFormat.EXCEL:
        assert kwargs["chart_factory"] is None


@pytest.mark.parametrize("factory", [group_snapshot, chi_snapshot, paired_snapshot])
@pytest.mark.parametrize("locale", ["en", "sv"])
@pytest.mark.parametrize("charts_only", [False, True])
def test_hypothesis_localized_workbooks_read_back_all_types(tmp_path, factory, locale, charts_only):
    translator = QTranslator()
    path = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales" / f"app_{locale}.qm"
    assert translator.load(str(path))
    app = QApplication.instance()
    app.installTranslator(translator)
    try:
        controller, dialog, snapshot, exporter = _hypothesis_export_context(factory())
        components = (Component.CHARTS,) if charts_only else tuple(Component)
        request = HypothesisExportRequest(snapshot, components, StatisticsExportFormat.EXCEL)
        controller._export_hypothesis(dialog, exporter, request)
        kwargs = exporter.export_excel.call_args.kwargs
        chart_name = "Charts" if locale == "en" else "Diagram"
        assert kwargs["chart_sheet_name"] == chart_name
        images = kwargs["chart_factory"](None, None)
        assert len(images) == 1
        output = tmp_path / f"{locale}.xlsx"
        FileWriter().save_excel_sheets(
            kwargs["sheets"],
            output,
            chart_images=images,
            chart_sheet_name=chart_name,
            streaming=False,
        )
        workbook = load_workbook(output)
        try:
            assert workbook.sheetnames == [*kwargs["sheets"], chart_name]
            assert len(workbook[chart_name]._images) == 1
            assert workbook[chart_name]["A1"].value == " / ".join(snapshot.columns)
            if not charts_only:
                result_name = "Test results" if locale == "en" else "Testresultat"
                frame = kwargs["sheets"][result_name]
                notes_header = "Notes" if locale == "en" else "Noteringar"
                assert frame[notes_header].dropna().tolist() == list(snapshot.notes)
                assert isinstance(workbook[result_name]["B2"].value, (int, float))
                if locale == "sv":
                    assert "Teststatistik" in frame.columns
                    assert any("→" in note or "mät" in note or "Grupp" in note for note in snapshot.notes)
        finally:
            workbook.close()
    finally:
        app.removeTranslator(translator)


@pytest.mark.parametrize("change", ["identity", "dataset", "category", "placeholder", "replacement"])
def test_hypothesis_export_rejects_stale_dialog_requests(change):
    controller, dialog, snapshot, exporter = _hypothesis_export_context(group_snapshot())
    request = HypothesisExportRequest(snapshot, (Component.SUMMARY,), StatisticsExportFormat.EXCEL)
    match change:
        case "identity":
            dialog.set_exportable_hypothesis(group_snapshot())
        case "dataset":
            dialog._dataset_combo.setCurrentIndex(1)
        case "category":
            dialog.select_category(AnalysisCategory.STATISTICS)
        case "placeholder":
            dialog.show_placeholder("prompt")
        case "replacement":
            dialog.set_content_widget(QWidget())
    controller._export_hypothesis(dialog, exporter, request)
    controller._export_hypothesis(dialog, exporter, object())
    assert not exporter.mock_calls


def test_statistics_excel_exports_localized_tables_and_every_column_chart(monkeypatch):
    controller, dialog, result, exporter = _statistics_export_context()
    translations = {
        "Continuous": "Kontinuerliga",
        "Categorical": "Kategoriska",
        "Charts": "Diagram",
        "Column": "Kolumn",
        "Count": "Antal",
        "Category": "Kategori",
        "Fraction (non-missing)": "Andel (icke-saknade)",
        "Histogram": "Histogram",
        "Boxplot": "Lådagram",
        "No data": "Ingen data",
    }
    monkeypatch.setattr(analysis_controller_module, "tr", lambda _context, text: translations.get(text, text))
    request = StatisticsExportRequest(
        result,
        (
            StatisticsExportTable.CONTINUOUS,
            StatisticsExportTable.CATEGORICAL,
            StatisticsExportTable.CHARTS,
        ),
        StatisticsExportFormat.EXCEL,
    )

    controller._export_statistics(dialog, exporter, request)

    kwargs = exporter.export_excel.call_args.kwargs
    assert list(kwargs["sheets"]) == ["Kontinuerliga", "Kategoriska"]
    assert list(kwargs["sheets"]["Kontinuerliga"].columns)[0] == "Kolumn"
    assert list(kwargs["sheets"]["Kategoriska"].columns)[:4] == [
        "Kolumn",
        "Kategori",
        "Antal",
        "Andel (icke-saknade)",
    ]
    assert kwargs["chart_sheet_name"] == "Diagram"
    images = kwargs["chart_factory"](None, None)
    assert [image.heading for image in images] == ["First numeric", "Second numeric"]
    assert kwargs["parent_widget"] is dialog
    assert kwargs["operation_target"] is dialog.content_panel()
    assert kwargs["show_success_dialog"] is True
    assert controller._results.mock_calls == []


@pytest.mark.parametrize(
    ("format_choice", "method_name"),
    [
        (StatisticsExportFormat.CSV, "export_csv"),
        (StatisticsExportFormat.BINARY, "export_data"),
    ],
)
def test_statistics_single_table_routes_to_non_excel_exporters(format_choice, method_name, monkeypatch):
    controller, dialog, result, exporter = _statistics_export_context()
    monkeypatch.setattr(analysis_controller_module, "tr", lambda _context, text: text)
    request = StatisticsExportRequest(result, (StatisticsExportTable.CATEGORICAL,), format_choice)

    controller._export_statistics(dialog, exporter, request)

    kwargs = getattr(exporter, method_name).call_args.kwargs
    assert list(kwargs["df"].columns)[0] == "Column"
    assert kwargs["show_success_dialog"] is True
    assert exporter.export_excel.call_count == 0


def test_statistics_export_rejects_stale_snapshot_without_lookup():
    controller, dialog, result, exporter = _statistics_export_context()
    request = StatisticsExportRequest(result, (StatisticsExportTable.CONTINUOUS,), StatisticsExportFormat.CSV)
    dialog.invalidate_overview_export()

    controller._export_statistics(dialog, exporter, request)

    assert exporter.mock_calls == []


def test_overview_export_dialog_cancellation_does_not_emit(monkeypatch):
    from expo_jbm329.gui.dialogs.analysis.analysis_export_dialog import AnalysisExportDialog

    _, dialog, _, _ = _overview_export_context()
    requests = []
    dialog.overview_export_requested.connect(requests.append)
    monkeypatch.setattr(AnalysisExportDialog, "exec", lambda _self: QDialog.DialogCode.Rejected)
    dialog._show_export_selection()
    assert requests == []


def test_overview_export_dialog_acceptance_emits_typed_snapshot_request(monkeypatch):
    from expo_jbm329.gui.dialogs.analysis.analysis_export_dialog import AnalysisExportDialog

    _, dialog, result, _ = _overview_export_context()
    requests = []
    dialog.overview_export_requested.connect(requests.append)
    monkeypatch.setattr(AnalysisExportDialog, "exec", lambda _self: QDialog.DialogCode.Accepted)
    dialog._show_export_selection()
    assert requests == [OverviewExportRequest(result, (OverviewExportTable.COLUMNS,), OverviewExportFormat.EXCEL)]


@pytest.mark.parametrize("invalidate", ["category", "dataset", "placeholder", "identity"])
def test_export_selection_rejects_snapshot_changed_while_open(monkeypatch, invalidate):
    from expo_jbm329.gui.dialogs.analysis.analysis_export_dialog import AnalysisExportDialog

    _, dialog, _, _ = _overview_export_context()
    requests = []
    dialog.overview_export_requested.connect(requests.append)

    def accept_after_change(_self):
        match invalidate:
            case "category":
                dialog.select_category(AnalysisCategory.REGRESSION)
            case "dataset":
                dialog.select_dataset("other")
            case "placeholder":
                dialog.show_placeholder("Analysis failed")
            case "identity":
                dialog.set_exportable_overview(analyze_dataset_overview(pd.DataFrame({"value": [9]})))
        return QDialog.DialogCode.Accepted

    monkeypatch.setattr(AnalysisExportDialog, "exec", accept_after_change)
    dialog._show_export_selection()
    assert requests == []


@pytest.mark.parametrize(
    "category", [category for category in AnalysisCategory if category is not AnalysisCategory.OVERVIEW]
)
def test_other_analysis_export_preview_never_emits_or_exports(monkeypatch, category):
    from expo_jbm329.gui.dialogs.analysis.analysis_export_dialog import AnalysisExportDialog

    _, dialog, _, exporter = _overview_export_context()
    requests = []
    dialog.overview_export_requested.connect(requests.append)
    dialog.select_category(category)
    monkeypatch.setattr(AnalysisExportDialog, "exec", lambda _self: QDialog.DialogCode.Accepted)
    dialog._show_export_selection()
    assert requests == []
    assert exporter.mock_calls == []


def test_overview_job_is_stale_after_navigating_away_and_back():
    controller, dialog, result, _ = _overview_export_context()
    controller._run_analysis(
        dialog,
        AnalysisCategory.OVERVIEW,
        controller._category_handlers[AnalysisCategory.OVERVIEW],
        "chosen",
        pd.DataFrame({"value": [1, 2]}),
    )
    call = controller._async_ops.last_call
    dialog.select_category(AnalysisCategory.STATISTICS)
    dialog.select_category(AnalysisCategory.OVERVIEW)
    assert call["stale_check"]()
    call["on_result"](result)
    assert dialog.exportable_overview() is None


def test_open_analysis_dialog_reuses_injected_exporter_and_connects_request(dialog_factory):
    results = DummyResults(datasets=[DatasetRef("chosen", "Chosen", 2, 1)], active_tab_id="chosen")
    exporter = Mock()
    controller = AnalysisController(results=results, async_ops=DummyAsyncOps(), export_controller=exporter)
    _open_and_flush(controller, QWidget())
    dialog = dialog_factory[0]
    assert exporter.mock_calls == []
    result = analyze_dataset_overview(pd.DataFrame({"a": [1]}))
    dialog.set_exportable_overview(result)
    dialog.overview_export_requested.emit(
        OverviewExportRequest(result, (OverviewExportTable.SAMPLE,), OverviewExportFormat.CSV)
    )
    assert len(exporter.mock_calls) == 1
    assert exporter.export_csv.call_count == 1
    kwargs = exporter.export_csv.call_args.kwargs
    assert kwargs["parent_widget"] is dialog
    assert kwargs["operation_target"] is dialog.content_panel()
    assert kwargs["show_success_dialog"] is True
    pd.testing.assert_frame_equal(kwargs["df"], pd.DataFrame({"a": [1]}))


class DummyResults:
    def __init__(self, datasets=None, active_tab_id=None, dfs=None):
        self._datasets = datasets or []
        self._active_tab_id = active_tab_id
        self._dfs = dfs or {}

    def list_ready_datasets(self):
        return self._datasets

    def active_tab_id(self):
        return self._active_tab_id

    def get_df_by_tab_id(self, tab_id):
        return self._dfs[tab_id]


class DummyAsyncOps:
    """Records run_target_overlay_operation calls without executing anything.

    Tests simulate the background job explicitly via the captured kwargs
    (work/on_result/on_error/stale_check), mirroring how the real
    AsyncOperationController would eventually invoke them.
    """

    def __init__(self):
        self.calls: list[dict] = []

    def run_target_overlay_operation(self, **kwargs):
        self.calls.append(kwargs)
        return object()

    @property
    def last_call(self) -> dict:
        return self.calls[-1]


class DummySignal:
    def __init__(self):
        self._slot = None

    def connect(self, slot):
        self._slot = slot

    def emit(self, value):
        if self._slot is not None:
            self._slot(value)


class DummyAnalysisDialog:
    def __init__(self, **kwargs):
        self.kwargs = kwargs
        self.category_changed = DummySignal()
        self.dataset_changed = DummySignal()
        self.overview_export_requested = DummySignal()
        self.statistics_export_requested = DummySignal()
        self.hypothesis_export_requested = DummySignal()
        self.correlation_export_requested = DummySignal()
        self.placeholder_calls: list[str] = []
        self.content_widgets: list[QWidget] = []
        self.config_widgets: list[QWidget | None] = []
        self.exec_called = False
        self._panel = QWidget()
        self._selected_category: AnalysisCategory | None = None
        self._selected_dataset_tab_id: str | None = None
        self._current_content: QWidget | None = None
        self._current_config: QWidget | None = None
        self._export_overview = None
        self._export_statistics = None
        self._export_hypothesis = None
        self._export_correlation = None
        self._correlation_export_revision = 0
        self._analysis_revision = 0

    def invalidate_overview_export(self):
        self._export_overview = None
        self._export_statistics = None
        self._export_hypothesis = None
        self.invalidate_correlation_export()
        self._analysis_revision += 1

    def invalidate_correlation_export(self):
        self._export_correlation = None
        self._correlation_export_revision += 1

    def correlation_export_revision(self):
        return self._correlation_export_revision

    def analysis_revision(self):
        return self._analysis_revision

    def set_exportable_overview(self, result):
        self._export_overview = result

    def exportable_overview(self):
        return self._export_overview

    def set_exportable_statistics(self, result):
        self._export_statistics = result

    def exportable_statistics(self):
        return self._export_statistics

    def set_exportable_hypothesis(self, snapshot):
        self._export_hypothesis = snapshot

    def exportable_hypothesis(self):
        return self._export_hypothesis

    def set_exportable_correlation(self, snapshot):
        self._export_correlation = snapshot

    def exportable_correlation(self):
        if self._selected_category is AnalysisCategory.CORRELATION:
            return self._export_correlation
        return None

    def show_placeholder(self, text):
        self.invalidate_overview_export()
        self.placeholder_calls.append(text)
        self._current_content = None

    def set_content_widget(self, widget):
        self.invalidate_overview_export()
        self.content_widgets.append(widget)
        self._current_content = widget

    def content_widget(self):
        return self._current_content

    def set_config_widget(self, widget):
        self.config_widgets.append(widget)
        self._current_config = widget

    def config_widget(self):
        return self._current_config

    def selected_category(self):
        return self._selected_category

    def selected_dataset_tab_id(self):
        return self._selected_dataset_tab_id

    def content_panel(self):
        return self._panel

    def exec(self):
        self.exec_called = True
        return QDialog.DialogCode.Accepted


@pytest.fixture
def dialog_factory(monkeypatch):
    dialogs: list[DummyAnalysisDialog] = []

    def _factory(**kwargs):
        dlg = DummyAnalysisDialog(**kwargs)
        dialogs.append(dlg)
        return dlg

    monkeypatch.setattr(f"{_CONTROLLER_MODULE}.AnalysisDialog", _factory)
    yield dialogs

    # Matplotlib canvases (StatisticsView) schedule a deferred draw_idle();
    # flush it here, while `dialogs` still keeps every created widget alive,
    # so the queued paint never fires later against an already-garbage-
    # collected canvas in an unrelated test.
    QApplication.processEvents()


def _open_and_flush(ctrl: AnalysisController, parent: QWidget) -> None:
    """Open the dialog and let the deferred initial refresh (singleShot) run."""
    ctrl.open_dialog(parent)
    QApplication.processEvents()


def _simulate_success(call_kwargs: dict) -> None:
    """Simulate a background job completing successfully.

    Mirrors AsyncOperationController's own stale_check gating: a stale
    result must never reach on_result.
    """
    stale_check = call_kwargs.get("stale_check")
    if stale_check is not None and stale_check():
        return
    result = call_kwargs["work"](cancel_cb=lambda: False)
    call_kwargs["on_result"](result)


def _assert_apply_prompt(ctrl: AnalysisController, dlg: DummyAnalysisDialog) -> None:
    content = dlg.content_widget()
    assert isinstance(content, QLabel)
    assert content.text() == ctrl._tr(ctrl.TR_APPLY_PROMPT)  # noqa: SLF001


def test_open_dialog_does_nothing_when_there_are_no_datasets(dialog_factory):
    ctrl = AnalysisController(results=DummyResults(datasets=[]), async_ops=DummyAsyncOps())
    _open_and_flush(ctrl, QWidget())

    assert dialog_factory == []


def test_reports_selected_before_initial_refresh_do_not_launch_analysis(monkeypatch):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    dialogs: list[AnalysisDialog] = []

    def open_reports(dialog):
        dialogs.append(dialog)
        dialog.select_reports()
        return QDialog.DialogCode.Rejected

    monkeypatch.setattr(AnalysisDialog, "exec", open_reports)
    results = DummyResults(datasets=[dataset], active_tab_id="t1")
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(results=results, async_ops=async_ops)

    parent = QWidget()
    _open_and_flush(ctrl, parent)

    assert len(dialogs) == 1
    dialog = dialogs[0]
    assert dialog.selected_category() is None
    assert async_ops.calls == []
    # Even a programmatic dataset change must not launch a report-page analysis.
    dialog.dataset_changed.emit("t1")
    assert async_ops.calls == []
    assert dialog.content_widget() is None


@pytest.mark.parametrize("recompute", [False, True])
def test_report_navigation_discards_pending_results_and_errors(recompute):
    frame = pd.DataFrame({"a": [1, 2, 3]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    results = DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": frame})
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(results=results, async_ops=async_ops)
    dialog = AnalysisDialog(parent=None, datasets=[dataset], active_tab_id="t1")
    content = QLabel("Existing result")
    dialog.set_content_widget(content)
    if recompute:
        ctrl._recompute_content(
            dialog,
            category=AnalysisCategory.OVERVIEW,
            scope_suffix="preview",
            compute=lambda df, callbacks: "New result",
            apply_result=lambda result: dialog.show_placeholder(str(result)),
            is_stale=lambda: False,
        )
    else:
        ctrl._refresh_content(dialog)
    assert len(async_ops.calls) == 1
    call = async_ops.last_call

    dialog.select_reports()
    page = dialog.report_page()
    assert call["stale_check"]()
    _simulate_success(call)
    call["on_error"]("Late error")
    ctrl._refresh_content(dialog)

    assert dialog.content_widget() is content
    assert dialog.report_page() is page
    assert async_ops.calls == [call]


@pytest.mark.parametrize("category", [AnalysisCategory.OVERVIEW, AnalysisCategory.REGRESSION])
def test_returning_from_reports_preserves_automatic_and_apply_first_behavior(monkeypatch, category):
    frame = pd.DataFrame({"a": range(8), "b": range(1, 9)})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=8, column_count=2)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": frame}),
        async_ops=async_ops,
    )
    dialogs: list[AnalysisDialog] = []

    def navigate(dialog):
        dialogs.append(dialog)
        dialog.select_reports()
        QApplication.processEvents()
        assert async_ops.calls == []
        dialog.select_category(category)
        return QDialog.DialogCode.Rejected

    monkeypatch.setattr(AnalysisDialog, "exec", navigate)

    parent = QWidget()
    _open_and_flush(ctrl, parent)

    dialog = dialogs[0]
    assert dialog.selected_category() is category
    if category is AnalysisCategory.OVERVIEW:
        assert len(async_ops.calls) == 1
        _simulate_success(async_ops.last_call)
        assert isinstance(dialog.content_widget(), OverviewView)
    else:
        assert async_ops.calls == []
        assert isinstance(dialog.config_widget(), RegressionConfigWidget)
        assert isinstance(dialog.content_widget(), RegressionView)
        dialog.config_widget().model_requested.emit()
        assert len(async_ops.calls) == 1


def test_open_dialog_builds_dialog_with_datasets_and_active_tab(dialog_factory):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=10, column_count=3)

    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1"),
        async_ops=DummyAsyncOps(),
    )
    _open_and_flush(ctrl, QWidget())

    assert len(dialog_factory) == 1
    dlg = dialog_factory[0]
    assert dlg.kwargs["datasets"] == [dataset]
    assert dlg.kwargs["active_tab_id"] == "t1"
    assert dlg.exec_called


def test_open_dialog_schedules_exactly_one_deferred_initial_refresh(monkeypatch, dialog_factory):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=10, column_count=3)
    calls: list[object] = []
    monkeypatch.setattr(AnalysisController, "_refresh_content", lambda self, dialog: calls.append(dialog))

    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1"),
        async_ops=DummyAsyncOps(),
    )
    _open_and_flush(ctrl, QWidget())

    assert calls == [dialog_factory[0]]


def test_deferred_initial_refresh_is_skipped_after_category_interaction(monkeypatch):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    dialogs: list[DummyAnalysisDialog] = []
    calls: list[object] = []

    class InteractingDialog(DummyAnalysisDialog):
        def exec(self):
            self.exec_called = True
            self._selected_category = AnalysisCategory.STATISTICS
            self._selected_dataset_tab_id = "t1"
            self.category_changed.emit(AnalysisCategory.STATISTICS.value)
            return QDialog.DialogCode.Accepted

    def _factory(**kwargs):
        dialog = InteractingDialog(**kwargs)
        dialogs.append(dialog)
        return dialog

    monkeypatch.setattr(f"{_CONTROLLER_MODULE}.AnalysisDialog", _factory)
    monkeypatch.setattr(AnalysisController, "_refresh_content", lambda self, dialog: calls.append(dialog))
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1"),
        async_ops=DummyAsyncOps(),
    )

    _open_and_flush(ctrl, QWidget())

    assert calls == [dialogs[0]]


def test_initial_refresh_shows_overview_content_when_overview_is_preselected(monkeypatch):
    """AnalysisDialog always pre-selects Overview and an active dataset at
    construction time; the deferred initial refresh must turn that default
    selection into real content without any further signal.
    """
    df = pd.DataFrame({"a": [1, 2, 3]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    dialogs: list[DummyAnalysisDialog] = []

    def _factory(**kwargs):
        dlg = DummyAnalysisDialog(**kwargs)
        dlg._selected_category = AnalysisCategory.OVERVIEW
        dlg._selected_dataset_tab_id = "t1"
        dialogs.append(dlg)
        return dlg

    monkeypatch.setattr(f"{_CONTROLLER_MODULE}.AnalysisDialog", _factory)

    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialogs[0]
    assert len(async_ops.calls) == 1
    _simulate_success(async_ops.last_call)

    assert len(dlg.content_widgets) == 1
    assert isinstance(dlg.content_widgets[0], OverviewView)


def test_category_without_a_registered_handler_shows_not_implemented_placeholder(dialog_factory):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=10, column_count=3)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1"),
        async_ops=async_ops,
    )
    ctrl._category_handlers.pop(AnalysisCategory.TIME_SERIES)  # noqa: SLF001
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.TIME_SERIES
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.TIME_SERIES.value)

    assert dlg.placeholder_calls == [ctrl._tr(ctrl.TR_NOT_IMPLEMENTED)]  # noqa: SLF001
    assert dlg.content_widgets == []
    assert async_ops.calls == []  # never dispatches a background job


def test_overview_category_runs_as_a_background_job_with_a_busy_overlay(dialog_factory):
    df = pd.DataFrame({"a": [1, 2, 3], "b": ["x", "y", "z"]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=2)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)

    assert len(async_ops.calls) == 1
    call = async_ops.last_call
    assert call["target"] is dlg.content_panel()
    assert call["runner"] == "pool"
    assert call["scope"] == "analysis:overview"
    # Nothing is shown until the job "completes" - the busy overlay is what
    # covers the result pane in the meantime.
    assert dlg.content_widgets == []

    _simulate_success(call)

    assert len(dlg.content_widgets) == 1
    assert isinstance(dlg.content_widgets[0], OverviewView)


def test_statistics_category_runs_as_a_background_job_with_a_busy_overlay(dialog_factory):
    df = pd.DataFrame({"a": [1.0, 2.0, 3.0], "b": ["x", "y", "z"]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=2)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.STATISTICS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.STATISTICS.value)

    assert len(async_ops.calls) == 1
    call = async_ops.last_call
    assert call["target"] is dlg.content_panel()
    assert call["runner"] == "pool"
    assert call["scope"] == "analysis:statistics"
    assert dlg.content_widgets == []

    _simulate_success(call)

    assert len(dlg.content_widgets) == 1
    assert isinstance(dlg.content_widgets[0], StatisticsView)
    assert dlg.placeholder_calls == []


def test_statistics_category_also_builds_a_column_picker_config_widget(dialog_factory):
    df = pd.DataFrame({"a": [1.0, 2.0, 3.0], "b": [4.0, 5.0, 6.0]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=2)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.STATISTICS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.STATISTICS.value)
    _simulate_success(async_ops.last_call)

    assert len(dlg.config_widgets) == 2  # proactive None, then the real config widget
    config = dlg.config_widgets[-1]
    assert isinstance(config, StatisticsConfigWidget)
    assert config.selected_column() == "a"


def test_statistics_category_shows_categorical_summaries_without_numeric_config(dialog_factory):
    df = pd.DataFrame({"sex": ["F", "M", "F", None]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=4, column_count=1)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.STATISTICS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.STATISTICS.value)
    _simulate_success(async_ops.last_call)

    content = dlg.content_widgets[-1]
    assert isinstance(content, StatisticsView)
    assert content.findChild(QTableWidget) is not None
    assert dlg.config_widgets[-1] is None


def test_changing_the_statistics_config_column_updates_the_content_view_directly(dialog_factory):
    """Switching columns is a pure GUI-thread operation - it must not
    dispatch a new background job (all columns' data is already computed).
    """
    df = pd.DataFrame({"a": [1.0, 2.0, 3.0], "b": [4.0, 5.0, 6.0]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=2)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.STATISTICS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.STATISTICS.value)
    _simulate_success(async_ops.last_call)

    jobs_before = len(async_ops.calls)
    config = dlg.config_widgets[-1]
    content = dlg.content_widgets[-1]
    assert isinstance(content, StatisticsView)
    config._column_combo.setCurrentIndex(1)  # noqa: SLF001

    assert len(async_ops.calls) == jobs_before  # no new background job
    assert config.selected_column() == "b"
    numeric_table = next(table for table in content.findChildren(QTableWidget) if table.columnCount() == 17)
    assert numeric_table.item(1, 15).text().startswith("5")
    assert numeric_table.item(1, 16).text().startswith("5")
    assert "b: Shapiro-Wilk suggests" in content._recommendation_label.text()  # noqa: SLF001


def test_clicking_a_statistics_table_row_updates_the_config_without_a_new_job(dialog_factory):
    df = pd.DataFrame({"a": [1.0, 2.0, 3.0], "b": [4.0, 5.0, 6.0]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=2)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.STATISTICS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.STATISTICS.value)
    _simulate_success(async_ops.last_call)

    jobs_before = len(async_ops.calls)
    content = dlg.content_widgets[-1]
    config = dlg.config_widgets[-1]
    assert isinstance(content, StatisticsView)
    assert isinstance(config, StatisticsConfigWidget)
    table = next(table for table in content.findChildren(QTableWidget) if table.columnCount() == 17)

    table.cellClicked.emit(1, 0)

    assert len(async_ops.calls) == jobs_before
    assert config.selected_column() == "b"
    assert "b: Shapiro-Wilk suggests" in content._recommendation_label.text()  # noqa: SLF001


def test_switching_to_a_category_hides_a_stale_config_widget_while_loading(dialog_factory):
    df = pd.DataFrame({"a": [1.0, 2.0, 3.0]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.STATISTICS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.STATISTICS.value)
    _simulate_success(async_ops.last_call)
    assert dlg.config_widgets[-1] is not None

    # Switch to Overview: the stale Statistics config widget must be hidden
    # immediately, even before Overview's own job completes.
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)

    assert dlg.config_widgets[-1] is None


def test_overview_result_is_discarded_when_the_dataset_changes_before_it_completes(dialog_factory):
    df1 = pd.DataFrame({"a": [1, 2, 3]})
    df2 = pd.DataFrame({"a": [1, 2, 3, 4, 5]})
    datasets = [
        DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1),
        DatasetRef(tab_id="t2", title="Sheet2", row_count=5, column_count=1),
    ]
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=datasets, active_tab_id="t1", dfs={"t1": df1, "t2": df2}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)
    first_call = async_ops.last_call

    # User switches dataset before the first job "completes".
    dlg._selected_dataset_tab_id = "t2"
    dlg.dataset_changed.emit("t2")
    second_call = async_ops.last_call

    # The stale (first) job finally completes: its result must be dropped.
    _simulate_success(first_call)
    assert dlg.content_widgets == []

    # The fresh (second) job completes normally.
    _simulate_success(second_call)
    assert len(dlg.content_widgets) == 1
    assert dlg.content_widgets[0]._result.row_count == 5  # noqa: SLF001


def test_overview_category_shows_error_placeholder_when_dataset_lookup_fails(dialog_factory):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)

    assert async_ops.calls == []  # dataset lookup fails before any job starts
    assert dlg.content_widgets == []
    assert dlg.placeholder_calls == [ctrl._tr(ctrl.TR_ANALYSIS_ERROR)]  # noqa: SLF001


def test_overview_category_shows_error_placeholder_when_the_job_itself_fails(dialog_factory):
    df = pd.DataFrame({"a": [1, 2, 3]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)

    async_ops.last_call["on_error"]("boom")

    assert dlg.content_widgets == []
    assert dlg.placeholder_calls == [ctrl._tr(ctrl.TR_ANALYSIS_ERROR)]  # noqa: SLF001


def test_job_error_is_ignored_once_stale(dialog_factory):
    df1 = pd.DataFrame({"a": [1, 2, 3]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df1}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)
    call = async_ops.last_call

    # User navigates away to an unimplemented category before the job fails.
    ctrl._category_handlers.pop(AnalysisCategory.TIME_SERIES)  # noqa: SLF001
    dlg._selected_category = AnalysisCategory.TIME_SERIES
    dlg.category_changed.emit(AnalysisCategory.TIME_SERIES.value)
    assert dlg.placeholder_calls == [ctrl._tr(ctrl.TR_NOT_IMPLEMENTED)]  # noqa: SLF001

    call["on_error"]("boom")

    # The stale error must not clobber the (already correct) placeholder.
    assert dlg.placeholder_calls == [ctrl._tr(ctrl.TR_NOT_IMPLEMENTED)]  # noqa: SLF001


def test_cancelled_work_returns_none_and_is_ignored(dialog_factory):
    df = pd.DataFrame({"a": [1, 2, 3]})
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1", dfs={"t1": df}),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)
    call = async_ops.last_call

    result = call["work"](cancel_cb=lambda: True)
    assert result is None

    call["on_result"](result)
    assert dlg.content_widgets == []


def test_no_category_selected_does_not_touch_the_dialog(dialog_factory):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=3, column_count=1)
    async_ops = DummyAsyncOps()
    ctrl = AnalysisController(
        results=DummyResults(datasets=[dataset], active_tab_id="t1"),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg.dataset_changed.emit("t1")

    assert dlg.placeholder_calls == []
    assert dlg.content_widgets == []
    assert async_ops.calls == []


# ----------------------------------------------------------------------
# Hypothesis Tests: Apply-first recompute
# ----------------------------------------------------------------------


def _hypothesis_tests_df() -> pd.DataFrame:
    # 24 rows with fully distinct "value"/"other" values (> MAX_GROUPS) so
    # only "grp" and "color" are eligible categorical columns - Group
    # Comparison defaults to ("value", "grp") with two numeric choices to
    # switch between, and chi-square defaults to ("grp", "color").
    return pd.DataFrame({
        "value": [float(i) for i in range(24)],
        "other": [float(i) * 10 for i in range(24)],
        "grp": ["A", "B"] * 12,
        "color": ["r", "g", "b"] * 8,
    })


def _gc_numeric_combo(config: HypothesisTestsConfigWidget) -> QComboBox:
    gc_config = config._group_comparison_config  # noqa: SLF001
    assert gc_config is not None
    return gc_config._numeric_combo  # noqa: SLF001


def _open_hypothesis_tests(async_ops: DummyAsyncOps, dialog_factory, df: pd.DataFrame | None = None):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=24, column_count=4)
    ctrl = AnalysisController(
        results=DummyResults(
            datasets=[dataset],
            active_tab_id="t1",
            dfs={"t1": _hypothesis_tests_df() if df is None else df},
        ),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.HYPOTHESIS_TESTS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.HYPOTHESIS_TESTS.value)
    return ctrl, dlg


def _hypothesis_tests_config(dlg: DummyAnalysisDialog) -> HypothesisTestsConfigWidget:
    config = dlg.config_widgets[-1]
    assert isinstance(config, HypothesisTestsConfigWidget)
    return config


def _apply_hypothesis_test(async_ops: DummyAsyncOps, dlg: DummyAnalysisDialog) -> dict:
    """Click Apply and return the test job's call, without completing it."""
    _hypothesis_tests_config(dlg)._apply_button.click()  # noqa: SLF001
    return async_ops.last_call


def _select_test(config: HypothesisTestsConfigWidget, test: HypothesisTest) -> None:
    combo = config._test_combo  # noqa: SLF001
    combo.setCurrentIndex(combo.findData(test.value))


def test_hypothesis_pending_edits_retain_displayed_applied_export_until_new_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dialog = _open_hypothesis_tests(async_ops, dialog_factory)
    assert dialog.exportable_hypothesis() is None
    call = _apply_hypothesis_test(async_ops, dialog)
    assert dialog.exportable_hypothesis() is None
    _simulate_success(call)
    snapshot = dialog.exportable_hypothesis()
    assert snapshot is not None
    assert snapshot.columns == ("value", "grp")
    config = _hypothesis_tests_config(dialog)
    _gc_numeric_combo(config).setCurrentIndex(1)
    assert config.current_configuration() != config.applied_configuration()
    assert dialog.exportable_hypothesis() is snapshot
    selection_dialog = AnalysisExportDialog(hypothesis=snapshot)
    text = "\n".join(label.text() for label in selection_dialog.findChildren(QLabel))
    assert "value / grp" in text
    assert "other / grp" not in text
    second = _apply_hypothesis_test(async_ops, dialog)
    assert dialog.exportable_hypothesis() is None
    _simulate_success(second)
    assert dialog.exportable_hypothesis().columns == ("other", "grp")


@pytest.mark.parametrize("change", ["test-return", "dataset", "category-return", "prompt", "error", "apply"])
def test_hypothesis_late_callbacks_cannot_restore_invalidated_exports(dialog_factory, change):
    async_ops = DummyAsyncOps()
    _, dialog = _open_hypothesis_tests(async_ops, dialog_factory)
    config = _hypothesis_tests_config(dialog)
    call = _apply_hypothesis_test(async_ops, dialog)
    value = call["work"]()
    match change:
        case "test-return":
            _select_test(config, HypothesisTest.CHI_SQUARE)
            _select_test(config, HypothesisTest.GROUP_COMPARISON)
        case "dataset":
            dialog._selected_dataset_tab_id = "changed"
        case "category-return":
            dialog.invalidate_overview_export()
        case "prompt":
            dialog.show_placeholder("prompt")
        case "error":
            call["on_error"]("failure")
        case "apply":
            _apply_hypothesis_test(async_ops, dialog)
    assert call["stale_check"]()
    call["on_result"](value)
    assert dialog.exportable_hypothesis() is None


def test_hypothesis_failed_and_cancelled_jobs_never_register_exports(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dialog = _open_hypothesis_tests(
        async_ops,
        dialog_factory,
        pd.DataFrame({"value": [1, 2, 3], "other": [1, 2, 3]}),
    )
    _select_test(_hypothesis_tests_config(dialog), HypothesisTest.PAIRED_COMPARISON)
    _simulate_success(_apply_hypothesis_test(async_ops, dialog))
    assert dialog.exportable_hypothesis() is None
    call = _apply_hypothesis_test(async_ops, dialog)
    assert call["work"](cancel_cb=lambda: True) is None
    call["on_result"](None)
    assert dialog.exportable_hypothesis() is None


def test_hypothesis_tests_initialize_without_a_job_and_prompt_for_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_hypothesis_tests(async_ops, dialog_factory)

    assert async_ops.calls == []
    _assert_apply_prompt(ctrl, dlg)
    config = _hypothesis_tests_config(dlg)
    assert config.current_configuration() == (HypothesisTest.GROUP_COMPARISON, ("value", "grp"))
    assert config.chi_square_selection() == ("grp", "color")
    assert config.applied_configuration() is None


def test_hypothesis_tests_with_no_eligible_columns_show_the_error_and_the_test_selector(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory, pd.DataFrame({"a": ["x", "y", "z"]}))

    assert async_ops.calls == []
    assert isinstance(dlg.content_widget(), GroupComparisonView)
    config = _hypothesis_tests_config(dlg)
    assert config.group_comparison_selection() is None
    assert config.chi_square_selection() is None
    assert not config._apply_button.isEnabled()  # noqa: SLF001


def test_hypothesis_test_column_changes_start_no_job_until_applied(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)

    _gc_numeric_combo(_hypothesis_tests_config(dlg)).setCurrentIndex(1)  # "value" -> "other"

    assert async_ops.calls == []


def test_applying_a_hypothesis_test_runs_a_background_job_and_keeps_the_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    config = _hypothesis_tests_config(dlg)
    config_widgets_before = len(dlg.config_widgets)
    _gc_numeric_combo(config).setCurrentIndex(1)  # "value" -> "other"

    call = _apply_hypothesis_test(async_ops, dlg)

    assert len(async_ops.calls) == 1
    assert call["scope"] == "analysis:hypothesis_tests:group_comparison:other:grp"
    assert call["target"] is dlg.content_panel()
    assert call["cancelable"] is False
    assert call["indeterminate"] is True
    assert call["timeout_ms"] == 60_000

    _simulate_success(call)

    assert isinstance(dlg.content_widget(), GroupComparisonView)
    assert len(dlg.config_widgets) == config_widgets_before


def test_column_change_after_apply_keeps_the_result_until_applied_again(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    _simulate_success(_apply_hypothesis_test(async_ops, dlg))
    view = dlg.content_widget()
    jobs_before = len(async_ops.calls)

    _gc_numeric_combo(_hypothesis_tests_config(dlg)).setCurrentIndex(1)

    assert len(async_ops.calls) == jobs_before
    assert dlg.content_widget() is view


def test_stale_hypothesis_test_result_is_discarded_when_applied_again(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    config = _hypothesis_tests_config(dlg)

    _gc_numeric_combo(config).setCurrentIndex(1)  # "value" -> "other"
    first = _apply_hypothesis_test(async_ops, dlg)
    _gc_numeric_combo(config).setCurrentIndex(0)  # "other" -> "value" again
    second = _apply_hypothesis_test(async_ops, dlg)

    content_widgets_before = len(dlg.content_widgets)
    _simulate_success(first)
    assert len(dlg.content_widgets) == content_widgets_before  # stale result dropped

    _simulate_success(second)
    assert len(dlg.content_widgets) == content_widgets_before + 1


def test_stale_hypothesis_test_result_is_discarded_when_category_changes(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    call = _apply_hypothesis_test(async_ops, dlg)

    # User navigates to Overview before the test completes.
    dlg._selected_category = AnalysisCategory.OVERVIEW
    dlg.category_changed.emit(AnalysisCategory.OVERVIEW.value)
    _simulate_success(async_ops.last_call)  # Overview's own job

    content_widgets_before = len(dlg.content_widgets)
    _simulate_success(call)
    assert len(dlg.content_widgets) == content_widgets_before  # stale, dropped


def test_hypothesis_test_error_shows_the_error_placeholder_without_touching_config(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    config_widgets_before = len(dlg.config_widgets)

    _apply_hypothesis_test(async_ops, dlg)["on_error"]("boom")

    assert dlg.placeholder_calls[-1] == ctrl._tr(ctrl.TR_ANALYSIS_ERROR)  # noqa: SLF001
    assert len(dlg.config_widgets) == config_widgets_before  # config untouched


# ----------------------------------------------------------------------
# Hypothesis Tests: switching between tests
# ----------------------------------------------------------------------


def test_switching_test_replaces_the_result_with_the_prompt_without_a_job(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    _simulate_success(_apply_hypothesis_test(async_ops, dlg))
    config_widgets_before = len(dlg.config_widgets)
    jobs_before = len(async_ops.calls)

    _select_test(_hypothesis_tests_config(dlg), HypothesisTest.CHI_SQUARE)

    assert len(async_ops.calls) == jobs_before
    _assert_apply_prompt(ctrl, dlg)
    assert len(dlg.config_widgets) == config_widgets_before


def test_applying_chi_square_runs_a_chi_square_background_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    _select_test(_hypothesis_tests_config(dlg), HypothesisTest.CHI_SQUARE)

    call = _apply_hypothesis_test(async_ops, dlg)

    assert call["scope"] == "analysis:hypothesis_tests:chi_square:grp:color"
    _simulate_success(call)
    assert isinstance(dlg.content_widget(), ChiSquareView)


def test_applying_paired_comparison_runs_and_renders_a_background_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    config = _hypothesis_tests_config(dlg)
    _select_test(config, HypothesisTest.PAIRED_COMPARISON)

    call = _apply_hypothesis_test(async_ops, dlg)

    assert call["scope"] == "analysis:hypothesis_tests:paired_comparison:value:other"
    _simulate_success(call)
    assert isinstance(dlg.content_widget(), PairedComparisonView)
    assert len(dlg.config_widgets) == 1


def test_paired_comparison_excludes_incomplete_subjects_and_renders_result(dialog_factory):
    df = pd.DataFrame({
        "before": [1.0, 2.0, 3.0, None],
        "after": [2.0, 3.0, 5.0, 9.0],
        "group": ["A", "B", "A", "B"],
    })
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory, df)
    _select_test(_hypothesis_tests_config(dlg), HypothesisTest.PAIRED_COMPARISON)

    _simulate_success(_apply_hypothesis_test(async_ops, dlg))

    view = dlg.content_widget()
    assert isinstance(view, PairedComparisonView)
    labels = view.findChildren(QLabel)
    assert any("Complete subjects: 3 of 4" in label.text() for label in labels)
    table = view.findChild(QTableWidget)
    assert table is not None
    assert table.rowCount() == 2
    assert table.item(0, 1).text() == "3"
    canvas = view.findChild(FigureCanvasQTAgg)
    assert canvas is not None
    assert len(canvas.figure.axes[1].lines) == 3
    assert list(canvas.figure.axes[1].lines[-1].get_ydata()) == [3.0, 5.0]


def test_switching_back_to_group_comparison_keeps_its_pending_selection(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    config = _hypothesis_tests_config(dlg)
    _gc_numeric_combo(config).setCurrentIndex(1)  # "value" -> "other"

    _select_test(config, HypothesisTest.CHI_SQUARE)
    _select_test(config, HypothesisTest.GROUP_COMPARISON)

    _assert_apply_prompt(ctrl, dlg)
    assert _apply_hypothesis_test(async_ops, dlg)["scope"] == "analysis:hypothesis_tests:group_comparison:other:grp"


def test_changing_the_chi_square_column_is_applied_with_the_new_pair(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    config = _hypothesis_tests_config(dlg)
    _select_test(config, HypothesisTest.CHI_SQUARE)

    chi_config = config._chi_square_config  # noqa: SLF001
    assert chi_config is not None
    chi_config._row_combo.setCurrentIndex(chi_config._row_combo.findText("color"))  # noqa: SLF001

    assert async_ops.calls == []
    call = _apply_hypothesis_test(async_ops, dlg)
    assert call["scope"] == "analysis:hypothesis_tests:chi_square:color:grp"
    _simulate_success(call)
    assert isinstance(dlg.content_widget(), ChiSquareView)


def test_stale_group_comparison_result_is_discarded_after_switching_test(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    group_comparison_call = _apply_hypothesis_test(async_ops, dlg)

    _select_test(_hypothesis_tests_config(dlg), HypothesisTest.CHI_SQUARE)
    content_widgets_before = len(dlg.content_widgets)
    _simulate_success(group_comparison_call)

    assert len(dlg.content_widgets) == content_widgets_before  # stale, dropped
    _assert_apply_prompt(ctrl, dlg)


def test_stale_paired_result_is_discarded_after_switching_test(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_hypothesis_tests(async_ops, dialog_factory)
    config = _hypothesis_tests_config(dlg)
    _select_test(config, HypothesisTest.PAIRED_COMPARISON)
    paired_call = _apply_hypothesis_test(async_ops, dlg)

    _select_test(config, HypothesisTest.CHI_SQUARE)
    content_widgets_before = len(dlg.content_widgets)
    _simulate_success(paired_call)

    assert len(dlg.content_widgets) == content_widgets_before
    _assert_apply_prompt(ctrl, dlg)


def test_switching_to_an_unavailable_test_shows_its_error_and_disables_apply(dialog_factory):
    # Only "grp" is categorical: Group Comparison works, chi-square has no column pair.
    df = pd.DataFrame({"value": [float(i) for i in range(24)], "grp": ["A", "B"] * 12})
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory, df)
    config = _hypothesis_tests_config(dlg)

    _select_test(config, HypothesisTest.CHI_SQUARE)

    assert async_ops.calls == []
    assert isinstance(dlg.content_widget(), ChiSquareView)
    assert not config._apply_button.isEnabled()  # noqa: SLF001


# ----------------------------------------------------------------------
# Correlation Explorer
# ----------------------------------------------------------------------


def _correlation_df() -> pd.DataFrame:
    return pd.DataFrame({
        "a": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0],
        "b": [2.0, 4.1, 5.9, 8.2, 9.8, 12.1],
        "c": [5.0, 3.0, 4.0, 1.0, 2.0, 0.5],
        "text": ["x", "y", "z", "x", "y", "z"],
    })


def _correlation_export_context():
    frame = _correlation_df()
    frame["constant"] = 1.0
    outcome = AnalysisController._compute_correlation(
        frame,
        analysis_controller_module._JobCallbacks(None, None),  # noqa: SLF001
        columns=("a", "b", "c", "constant"),
    )
    assert outcome is not None and outcome.export_snapshot is not None
    dialog = AnalysisDialog(
        parent=None,
        datasets=[DatasetRef("chosen", "Chosen", len(frame), len(frame.columns))],
        active_tab_id="chosen",
    )
    dialog.select_category(AnalysisCategory.CORRELATION)
    dialog.set_content_widget(CorrelationView(outcome.matrix, outcome.pair_detail))
    dialog.set_exportable_correlation(outcome.export_snapshot)
    results = Mock()
    results.get_df_by_tab_id.side_effect = AssertionError("Export must not look up a dataset")
    results.current_df.side_effect = AssertionError("Export must not use the active tab")
    controller = AnalysisController(results=results, async_ops=DummyAsyncOps())
    return controller, dialog, outcome.export_snapshot, Mock()


@pytest.mark.parametrize("format_choice", list(StatisticsExportFormat))
def test_correlation_export_routes_only_localized_ranked_table(
    format_choice: StatisticsExportFormat,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, dialog, snapshot, exporter = _correlation_export_context()
    translations = {
        "Strongest correlations": "Starkaste korrelationer",
        "Variable 1": "Variabel 1",
        "Variable 2": "Variabel 2",
        "Coefficient": "Koefficient",
        "95% CI lower": "95 % KI nedre",
        "95% CI upper": "95 % KI övre",
        "p": "p",
        "Holm p": "Holm p",
        "n": "n",
        "Strength": "Styrka",
        "Significant after Holm": "Signifikant efter Holm",
        "Correlation method": "Korrelationsmetod",
        "Ranking and Holm explanation": "Rangordning och Holm-förklaring",
        "Pearson": "Pearson",
        "Negligible": "Försumbar",
        "Weak": "Svag",
        "Moderate": "Måttlig",
        "Strong": "Stark",
        "N/A": "Ej tillämpligt",
    }
    monkeypatch.setattr(analysis_controller_module, "tr", lambda _context, text: translations.get(text, text))
    request = CorrelationExportRequest(
        snapshot,
        (CorrelationExportComponent.STRONGEST_CORRELATIONS,),
        format_choice,
    )

    controller._export_correlation(dialog, exporter, request)

    method = {
        StatisticsExportFormat.EXCEL: exporter.export_excel,
        StatisticsExportFormat.CSV: exporter.export_csv,
        StatisticsExportFormat.BINARY: exporter.export_data,
    }[format_choice]
    assert method.call_count == 1
    kwargs = method.call_args.kwargs
    assert kwargs["parent_widget"] is dialog
    assert kwargs["operation_target"] is dialog.content_panel()
    assert kwargs["show_success_dialog"] is True
    if format_choice is StatisticsExportFormat.EXCEL:
        assert set(kwargs) == {
            "sheets",
            "chart_factory",
            "chart_sheet_name",
            "parent_widget",
            "operation_target",
            "show_success_dialog",
        }
        assert list(kwargs["sheets"]) == ["Starkaste korrelationer"]
        assert kwargs["chart_factory"] is None
        assert kwargs["chart_sheet_name"] is None
        frame = kwargs["sheets"]["Starkaste korrelationer"]
    else:
        assert set(kwargs) == {"df", "parent_widget", "operation_target", "show_success_dialog"}
        frame = kwargs["df"]
    assert frame.columns.tolist() == [
        "Variabel 1",
        "Variabel 2",
        "Koefficient",
        "95 % KI nedre",
        "95 % KI övre",
        "p",
        "Holm p",
        "n",
        "Styrka",
        "Signifikant efter Holm",
        "Korrelationsmetod",
        "Rangordning och Holm-förklaring",
    ]
    assert list(zip(frame["Variabel 1"], frame["Variabel 2"], strict=True)) == [
        (pair.x_column, pair.y_column) for pair in snapshot.matrix.pairs
    ]
    assert frame["Koefficient"].iloc[0] == snapshot.matrix.pairs[0].coefficient
    assert pd.api.types.is_numeric_dtype(frame["Koefficient"])
    assert pd.api.types.is_numeric_dtype(frame["95 % KI nedre"])
    assert frame["Korrelationsmetod"].eq("Pearson").all()
    assert "Holm adjustment" in frame["Rangordning och Holm-förklaring"].iloc[0]
    assert "*" not in frame["Rangordning och Holm-förklaring"].iloc[0]
    assert "Ej tillämpligt" in frame["Styrka"].tolist()
    assert frame["Koefficient"].isna().any()
    assert frame["95 % KI nedre"].isna().any()
    assert controller._results.mock_calls == []


def test_correlation_export_uses_compiled_swedish_labels() -> None:
    controller, dialog, snapshot, exporter = _correlation_export_context()
    translator = QTranslator()
    catalog = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales" / "app_sv.qm"
    assert translator.load(str(catalog))
    app = QApplication.instance()
    app.installTranslator(translator)
    try:
        controller._export_correlation(
            dialog,
            exporter,
            CorrelationExportRequest(
                snapshot,
                (CorrelationExportComponent.STRONGEST_CORRELATIONS,),
                StatisticsExportFormat.EXCEL,
            ),
        )
    finally:
        app.removeTranslator(translator)

    sheets = exporter.export_excel.call_args.kwargs["sheets"]
    assert list(sheets) == ["Starkaste korrelationerna"]
    frame = sheets["Starkaste korrelationerna"]
    assert frame.columns.tolist() == [
        "Variabel 1",
        "Variabel 2",
        "Koefficient",
        "95 % KI nedre",
        "95 % KI övre",
        "p",
        "Holm p",
        "n",
        "Styrka",
        "Signifikant efter Holm",
        "Korrelationsmetod",
        "Förklaring av rangordning och Holm-korrigering",
    ]
    assert "Stark" in frame["Styrka"].tolist()
    assert "Ej tillämpligt" in frame["Styrka"].tolist()
    assert "Holm-korrigering" in frame["Förklaring av rangordning och Holm-korrigering"].iloc[0]


def test_correlation_export_rejects_stale_navigation_without_dataset_lookup() -> None:
    controller, dialog, snapshot, exporter = _correlation_export_context()
    request = CorrelationExportRequest(
        snapshot,
        (CorrelationExportComponent.STRONGEST_CORRELATIONS,),
        StatisticsExportFormat.CSV,
    )
    dialog.select_category(AnalysisCategory.STATISTICS)

    controller._export_correlation(dialog, exporter, request)

    assert exporter.mock_calls == []
    assert controller._results.mock_calls == []


@pytest.mark.parametrize(
    "components",
    [
        (CorrelationExportComponent.MATRIX_PLOT,),
        (CorrelationExportComponent.SCATTERPLOTS,),
        tuple(CorrelationExportComponent),
    ],
)
def test_correlation_excel_routes_independent_and_mixed_charts_from_owned_snapshot(components):
    controller, dialog, snapshot, exporter = _correlation_export_context()
    request = CorrelationExportRequest(snapshot, components, StatisticsExportFormat.EXCEL)
    controller._export_correlation(dialog, exporter, request)
    kwargs = exporter.export_excel.call_args.kwargs
    assert kwargs["parent_widget"] is dialog
    assert kwargs["operation_target"] is dialog.content_panel()
    assert kwargs["show_success_dialog"]
    assert bool(kwargs["sheets"]) is (CorrelationExportComponent.STRONGEST_CORRELATIONS in components)
    assert kwargs["chart_sheet_name"] == "Charts"
    progress = []
    images = kwargs["chart_factory"](progress.append, None)
    expected = int(CorrelationExportComponent.MATRIX_PLOT in components)
    if CorrelationExportComponent.SCATTERPLOTS in components:
        expected += len(snapshot.matrix.pairs)
    assert len(images) == expected
    assert progress[-1] == 100
    assert all(image.image_data.startswith(b"\x89PNG") for image in images)
    assert controller._results.mock_calls == []


def test_invalid_correlation_matrix_selection_returns_structured_error_without_indexing() -> None:
    result = AnalysisController._compute_correlation(
        _correlation_df(),
        analysis_controller_module._JobCallbacks(None, None),  # noqa: SLF001
        columns=("missing", "a"),
    )

    assert result is not None
    assert result.matrix.error is CorrelationError.INVALID_COLUMN
    assert result.export_snapshot is None


def _open_correlation(async_ops: DummyAsyncOps, dialog_factory, df: pd.DataFrame | None = None):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=6, column_count=4)
    ctrl = AnalysisController(
        results=DummyResults(
            datasets=[dataset],
            active_tab_id="t1",
            dfs={"t1": _correlation_df() if df is None else df},
        ),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.CORRELATION
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.CORRELATION.value)
    return ctrl, dlg


def _correlation_config(dlg: DummyAnalysisDialog) -> CorrelationConfigWidget:
    config = dlg.config_widgets[-1]
    assert isinstance(config, CorrelationConfigWidget)
    return config


def _correlation_view(dlg: DummyAnalysisDialog) -> CorrelationView:
    view = dlg.content_widget()
    assert isinstance(view, CorrelationView)
    return view


def _apply_correlation(async_ops: DummyAsyncOps, dlg: DummyAnalysisDialog) -> dict:
    """Click Apply and return the matrix job's call, without completing it."""
    _correlation_config(dlg)._apply_button.click()  # noqa: SLF001
    return async_ops.last_call


def _open_applied_correlation(async_ops: DummyAsyncOps, dialog_factory):
    ctrl, dlg = _open_correlation(async_ops, dialog_factory)
    _simulate_success(_apply_correlation(async_ops, dlg))
    return ctrl, dlg


def test_correlation_initializes_without_a_job_and_prompts_for_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_correlation(async_ops, dialog_factory)

    assert async_ops.calls == []
    _assert_apply_prompt(ctrl, dlg)
    config = _correlation_config(dlg)
    assert config.matrix_configuration() == (CorrelationMethod.PEARSON, ("a", "b", "c"))
    assert config.checked_columns() == ("a", "b", "c")
    assert config.is_pair_selection_enabled() is False


def test_pair_change_before_the_first_apply_starts_no_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory)

    _correlation_config(dlg).set_pair("c", "a")

    assert async_ops.calls == []


def test_method_change_without_apply_starts_no_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    jobs_before = len(async_ops.calls)

    config._method_combo.setCurrentIndex(config._method_combo.findData(CorrelationMethod.SPEARMAN))  # noqa: SLF001

    assert len(async_ops.calls) == jobs_before
    assert config.matrix_configuration() == (CorrelationMethod.PEARSON, ("a", "b", "c"))
    assert _correlation_view(dlg).method() is CorrelationMethod.PEARSON


def test_pending_matrix_controls_and_pair_only_recompute_keep_export_snapshot(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    snapshot = dlg.exportable_correlation()
    assert snapshot is not None
    config = _correlation_config(dlg)
    config._method_combo.setCurrentIndex(config._method_combo.findData(CorrelationMethod.SPEARMAN))  # noqa: SLF001
    config._column_list.item(2).setCheckState(Qt.CheckState.Unchecked)  # noqa: SLF001

    assert dlg.exportable_correlation() is snapshot
    config.set_pair("c", "a")
    pair_call = async_ops.last_call
    assert pair_call["scope"] == "analysis:correlation:pair:pearson:c:a"
    assert dlg.exportable_correlation() is snapshot
    _simulate_success(pair_call)
    assert dlg.exportable_correlation() is snapshot

    matrix_call = _apply_correlation(async_ops, dlg)
    assert dlg.exportable_correlation() is None
    _simulate_success(matrix_call)
    assert dlg.exportable_correlation() is not None
    assert dlg.exportable_correlation().matrix.method is CorrelationMethod.SPEARMAN


def test_correlation_apply_runs_as_a_cancelable_background_job_with_progress(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory)

    call = _apply_correlation(async_ops, dlg)

    assert call["scope"] == "analysis:correlation:matrix:pearson"
    assert call["target"] is dlg.content_panel()
    assert call["cancelable"] is True
    assert call["indeterminate"] is False
    assert call["timeout_ms"] == 600_000


def test_correlation_job_reports_progress_through_the_injected_callback(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory)
    call = _apply_correlation(async_ops, dlg)
    progress: list[int] = []

    result = call["work"](progress_cb=progress.append, cancel_cb=lambda: False)

    assert result is not None
    assert progress[-1] == 100


def test_matrix_worker_owns_selected_columns_before_analysis_and_pair_detail(dialog_factory):
    async_ops = DummyAsyncOps()
    source = _correlation_df()
    source_before = source.copy(deep=True)
    _, dlg = _open_correlation(async_ops, dialog_factory, source)
    call = _apply_correlation(async_ops, dlg)
    outcome = call["work"](cancel_cb=lambda: False)

    assert outcome.export_snapshot is not None
    assert outcome.export_snapshot.matrix_data.columns.tolist() == ["a", "b", "c"]
    pd.testing.assert_frame_equal(outcome.export_snapshot.matrix_data, source_before[["a", "b", "c"]])
    source.loc[:, ["a", "b", "c"]] = -999.0
    pd.testing.assert_frame_equal(outcome.export_snapshot.matrix_data, source_before[["a", "b", "c"]])
    assert outcome.pair_detail is not None
    expected_detail = analyze_correlation_pair(
        source_before,
        outcome.pair_detail.pair.x_column,
        outcome.pair_detail.pair.y_column,
        outcome.matrix.method,
    )
    assert outcome.pair_detail.pair.x_column == expected_detail.pair.x_column
    assert outcome.pair_detail.pair.y_column == expected_detail.pair.y_column
    assert outcome.pair_detail.pair.coefficient == expected_detail.pair.coefficient
    assert outcome.pair_detail.pair.n == expected_detail.pair.n
    assert np.isnan(outcome.pair_detail.pair.adjusted_p_value)

    call["on_result"](outcome)
    assert dlg.exportable_correlation() is outcome.export_snapshot


def test_first_apply_details_the_strongest_pair_and_enables_pair_selection(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    config.set_pair("c", "a")  # disabled pickers only hold a placeholder pair

    _simulate_success(_apply_correlation(async_ops, dlg))

    detail = _correlation_view(dlg).pair_detail()
    assert detail is not None
    assert (detail.pair.x_column, detail.pair.y_column) == ("a", "b")
    assert config.current_pair() == ("a", "b")
    assert config.is_pair_selection_enabled() is True
    assert async_ops.last_call["scope"] == "analysis:correlation:matrix:pearson"  # no extra pair job


def test_correlation_without_enough_numeric_columns_shows_error_without_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory, pd.DataFrame({"a": [1.0, 2.0, 3.0]}))

    assert async_ops.calls == []
    view = _correlation_view(dlg)
    assert view.table() is None
    assert dlg.config_widgets[-1] is None


def test_cancel_during_the_matrix_computation_returns_none(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory)
    call = _apply_correlation(async_ops, dlg)
    polls: list[None] = []

    def _cancel_after_first_poll() -> bool:
        polls.append(None)
        return len(polls) > 1

    assert call["work"](cancel_cb=_cancel_after_first_poll) is None


def test_applied_method_change_recomputes_the_matrix_and_keeps_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    first_view = _correlation_view(dlg)
    configs_before = len(dlg.config_widgets)

    config._method_combo.setCurrentIndex(config._method_combo.findData(CorrelationMethod.SPEARMAN))  # noqa: SLF001
    call = _apply_correlation(async_ops, dlg)

    assert call["scope"] == "analysis:correlation:matrix:spearman"
    _simulate_success(call)

    view = _correlation_view(dlg)
    assert view is not first_view
    assert view.method() is CorrelationMethod.SPEARMAN
    detail = view.pair_detail()
    assert detail is not None
    assert detail.method is CorrelationMethod.SPEARMAN
    assert len(dlg.config_widgets) == configs_before  # config is kept


def test_apply_recomputes_the_matrix_with_the_checked_columns(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)

    item = config._column_list.item(2)  # noqa: SLF001 - "c"
    item.setCheckState(item.checkState().Unchecked)
    _simulate_success(_apply_correlation(async_ops, dlg))

    view = _correlation_view(dlg)
    table = view.table()
    assert table is not None
    assert table.rowCount() == 1
    assert view.columns() == ("a", "b")
    assert config._x_combo.count() == 2  # noqa: SLF001
    assert config.current_pair() == ("a", "b")


def test_replacement_matrix_drops_excluded_pair_without_extra_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    config.set_pair("c", "a")
    _simulate_success(async_ops.last_call)
    item = config._column_list.item(2)  # noqa: SLF001
    item.setCheckState(item.checkState().Unchecked)

    call = _apply_correlation(async_ops, dlg)
    jobs_before = len(async_ops.calls)
    outcome = call["work"]()
    assert outcome.pair_detail is not None
    assert (outcome.pair_detail.pair.x_column, outcome.pair_detail.pair.y_column) == ("a", "b")
    call["on_result"](outcome)

    assert config.current_pair() == ("a", "b")
    assert len(async_ops.calls) == jobs_before
    detail = _correlation_view(dlg).pair_detail()
    assert detail is not None
    assert (detail.pair.x_column, detail.pair.y_column) == ("a", "b")


def test_pair_controller_rejects_columns_outside_displayed_matrix(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    item = config._column_list.item(2)  # noqa: SLF001
    item.setCheckState(item.checkState().Unchecked)
    _simulate_success(_apply_correlation(async_ops, dlg))
    jobs_before = len(async_ops.calls)

    config.set_pair("c", "a")
    ctrl._recompute_correlation_pair(dlg, config, "c", "a")  # noqa: SLF001
    ctrl._recompute_correlation_pair(dlg, config, "a", "a")  # noqa: SLF001

    assert len(async_ops.calls) == jobs_before


def test_replacement_matrix_preserves_a_valid_pair_orientation(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    config.set_pair("c", "a")
    _simulate_success(async_ops.last_call)
    item = config._column_list.item(1)  # noqa: SLF001 - "b"
    item.setCheckState(item.checkState().Unchecked)

    call = _apply_correlation(async_ops, dlg)
    jobs_before = len(async_ops.calls)
    _simulate_success(call)

    assert config.current_pair() == ("c", "a")
    assert len(async_ops.calls) == jobs_before
    detail = _correlation_view(dlg).pair_detail()
    assert detail is not None
    assert (detail.pair.x_column, detail.pair.y_column) == ("c", "a")


@pytest.mark.parametrize("failed", [True, False])
def test_unsuccessful_matrix_replacement_does_not_publish_requested_pair_columns(dialog_factory, failed):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    item = config._column_list.item(2)  # noqa: SLF001
    item.setCheckState(item.checkState().Unchecked)
    call = _apply_correlation(async_ops, dlg)

    if failed:
        call["on_error"]("boom")
    else:
        call["on_result"](None)

    assert config._x_combo.count() == 3  # noqa: SLF001
    jobs_before = len(async_ops.calls)
    config.set_pair("c", "a")
    assert len(async_ops.calls) == jobs_before


def test_cancelled_matrix_recompute_shows_placeholder_and_keeps_config(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)

    call = _apply_correlation(async_ops, dlg)
    call["on_result"](call["work"](cancel_cb=lambda: True))

    assert dlg.placeholder_calls[-1] == ctrl._tr(ctrl.TR_ANALYSIS_CANCELLED)  # noqa: SLF001
    assert dlg.config_widgets[-1] is config
    assert config.is_pair_selection_enabled() is False


def test_stale_matrix_recompute_is_discarded_when_applied_again(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    combo = config._method_combo  # noqa: SLF001

    combo.setCurrentIndex(combo.findData(CorrelationMethod.SPEARMAN))
    first = _apply_correlation(async_ops, dlg)
    assert dlg.exportable_correlation() is None
    combo.setCurrentIndex(combo.findData(CorrelationMethod.KENDALL))
    second = _apply_correlation(async_ops, dlg)

    contents_before = len(dlg.content_widgets)
    _simulate_success(first)
    assert len(dlg.content_widgets) == contents_before

    _simulate_success(second)
    assert _correlation_view(dlg).method() is CorrelationMethod.KENDALL


def test_repeated_apply_with_same_configuration_rejects_the_older_matrix_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory)
    first = _apply_correlation(async_ops, dlg)
    second = _apply_correlation(async_ops, dlg)

    first_result = first["work"](cancel_cb=lambda: False)
    first["on_result"](first_result)
    assert dlg.exportable_correlation() is None
    second_result = second["work"](cancel_cb=lambda: False)
    second["on_result"](second_result)
    assert dlg.exportable_correlation() is not None


def test_pair_change_recomputes_only_the_pair_detail_over_the_pair_panel(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    view = _correlation_view(dlg)
    contents_before = len(dlg.content_widgets)

    config.set_pair("c", "a")

    call = async_ops.last_call
    assert call["scope"] == "analysis:correlation:pair:pearson:c:a"
    assert call["target"] is view.pair_panel()
    assert call["cancelable"] is False

    _simulate_success(call)

    assert len(dlg.content_widgets) == contents_before  # view updated in place
    detail = view.pair_detail()
    assert detail is not None
    assert (detail.pair.x_column, detail.pair.y_column) == ("c", "a")


def test_clicking_a_table_row_selects_the_pair_and_recomputes_it(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    view = _correlation_view(dlg)
    table = view.table()
    assert table is not None

    table.cellClicked.emit(2, 0)

    last_pair = view._result.pairs[2]  # noqa: SLF001
    assert config.current_pair() == (last_pair.x_column, last_pair.y_column)
    assert async_ops.last_call["scope"].startswith("analysis:correlation:pair:")


def test_stale_pair_recompute_is_discarded_when_the_pair_changes_again(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    view = _correlation_view(dlg)
    initial_detail = view.pair_detail()

    config.set_pair("c", "a")
    first = async_ops.last_call
    config.set_pair("b", "c")
    second = async_ops.last_call

    _simulate_success(first)
    assert view.pair_detail() is initial_detail

    _simulate_success(second)
    detail = view.pair_detail()
    assert detail is not None
    assert (detail.pair.x_column, detail.pair.y_column) == ("b", "c")


def test_pair_change_while_the_matrix_is_computing_is_caught_up_afterwards(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)

    matrix_call = _apply_correlation(async_ops, dlg)
    config.set_pair("c", "a")
    pair_call = async_ops.last_call
    assert pair_call["scope"] == "analysis:correlation:pair:pearson:c:a"

    _simulate_success(matrix_call)
    catchup_call = async_ops.last_call
    assert catchup_call["scope"] == "analysis:correlation:pair:pearson:c:a"
    _simulate_success(pair_call)
    _simulate_success(catchup_call)

    detail = _correlation_view(dlg).pair_detail()
    assert detail is not None
    assert (detail.pair.x_column, detail.pair.y_column) == ("c", "a")


def test_pair_recompute_uses_the_displayed_matrix_method(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)

    # Kendall is applied, but its matrix job hasn't finished: the view is still Pearson.
    config._method_combo.setCurrentIndex(config._method_combo.findData(CorrelationMethod.KENDALL))  # noqa: SLF001
    _apply_correlation(async_ops, dlg)
    config.set_pair("c", "a")

    assert async_ops.last_call["scope"] == "analysis:correlation:pair:pearson:c:a"


def test_recompute_shows_error_placeholder_when_the_dataset_cannot_be_reloaded(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_correlation(async_ops, dialog_factory)
    ctrl._results._dfs.clear()  # noqa: SLF001 - the dataset disappeared

    _correlation_config(dlg)._apply_button.click()  # noqa: SLF001

    assert async_ops.calls == []
    assert dlg.placeholder_calls[-1] == ctrl._tr(ctrl.TR_ANALYSIS_ERROR)  # noqa: SLF001


def test_recompute_without_a_selected_dataset_does_nothing(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_correlation(async_ops, dialog_factory)
    dlg._selected_dataset_tab_id = None

    _correlation_config(dlg)._apply_button.click()  # noqa: SLF001

    assert async_ops.calls == []


def test_failed_recompute_shows_error_placeholder_unless_stale(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_correlation(async_ops, dialog_factory)
    config = _correlation_config(dlg)
    combo = config._method_combo  # noqa: SLF001

    combo.setCurrentIndex(combo.findData(CorrelationMethod.SPEARMAN))
    first = _apply_correlation(async_ops, dlg)
    combo.setCurrentIndex(combo.findData(CorrelationMethod.KENDALL))
    second = _apply_correlation(async_ops, dlg)

    first["on_error"]("boom")
    assert dlg.placeholder_calls == []

    second["on_error"]("boom")
    assert dlg.placeholder_calls == [ctrl._tr(ctrl.TR_ANALYSIS_ERROR)]  # noqa: SLF001


def test_non_correlation_jobs_keep_an_indeterminate_non_cancelable_overlay(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_hypothesis_tests(async_ops, dialog_factory)

    call = _apply_hypothesis_test(async_ops, dlg)
    assert call["cancelable"] is False
    assert call["indeterminate"] is True
    assert call["timeout_ms"] == 60_000


# ----------------------------------------------------------------------
# Linear Regression
# ----------------------------------------------------------------------


def _regression_df() -> pd.DataFrame:
    return pd.DataFrame({
        "y": [1.0, 2.5, 2.0, 4.5, 5.0, 6.5, 6.0, 8.5],
        "x": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0],
        "z": [2.0, 1.0, 4.0, 3.0, 6.0, 5.0, 8.0, 7.0],
        "g": ["a", "b", "a", "b", "a", "b", "a", "a"],
    })


def _open_regression(async_ops: DummyAsyncOps, dialog_factory, df: pd.DataFrame | None = None):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=8, column_count=4)
    ctrl = AnalysisController(
        results=DummyResults(
            datasets=[dataset],
            active_tab_id="t1",
            dfs={"t1": _regression_df() if df is None else df},
        ),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.REGRESSION
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.REGRESSION.value)
    async_ops.calls.clear()
    return ctrl, dlg


def _regression_config(dlg: DummyAnalysisDialog) -> RegressionConfigWidget:
    config = dlg.config_widgets[-1]
    assert isinstance(config, RegressionConfigWidget)
    return config


def _regression_view(dlg: DummyAnalysisDialog) -> RegressionView:
    view = dlg.content_widgets[-1]
    assert isinstance(view, RegressionView)
    return view


def _check_predictors(config: RegressionConfigWidget, *columns: str) -> None:
    predictor_list = config._predictor_list  # noqa: SLF001
    for index in range(predictor_list.count()):
        item = predictor_list.item(index)
        assert item is not None
        if item.flags() & Qt.ItemFlag.ItemIsUserCheckable and item.flags() & Qt.ItemFlag.ItemIsEnabled:
            wanted = item.data(Qt.ItemDataRole.UserRole) in columns
            item.setCheckState(Qt.CheckState.Checked if wanted else Qt.CheckState.Unchecked)


def test_regression_initializes_without_a_background_job_or_busy_overlay(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory)

    assert async_ops.calls == []
    view = _regression_view(dlg)
    assert view.result().error is RegressionError.NO_PREDICTORS_SELECTED
    assert view.result().target == "y"
    config = _regression_config(dlg)
    assert config.model_configuration() == (RegressionModel.LINEAR, "y", ())


def test_regression_without_numeric_columns_can_still_configure_binary_models(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory, df=pd.DataFrame({"g": ["a", "b", "a"]}))

    assert async_ops.calls == []
    assert _regression_view(dlg).result().error is RegressionError.NO_NUMERIC_COLUMN
    config = _regression_config(dlg)
    _select_model = config._model_combo  # noqa: SLF001
    _select_model.setCurrentIndex(_select_model.findData(RegressionModel.LOGISTIC.value))
    assert config.current_target() == "g"


def test_failing_initializer_shows_error_placeholder_without_a_job(dialog_factory, monkeypatch):
    def _raise(_df: pd.DataFrame) -> object:
        raise ValueError

    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_regression(async_ops, dialog_factory)
    handler = ctrl._category_handlers[AnalysisCategory.REGRESSION]  # noqa: SLF001
    monkeypatch.setitem(
        ctrl._category_handlers,  # noqa: SLF001
        AnalysisCategory.REGRESSION,
        dataclasses.replace(handler, initialize=_raise),
    )

    dlg.category_changed.emit(AnalysisCategory.REGRESSION.value)

    assert async_ops.calls == []
    assert dlg.placeholder_calls[-1] == ctrl._tr(ctrl.TR_ANALYSIS_ERROR)  # noqa: SLF001
    assert dlg.config_widgets[-1] is None


def test_applying_predictors_refits_the_model_in_a_background_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory)
    config = _regression_config(dlg)

    _check_predictors(config, "x", "g")
    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:regression:y:x:g"
    assert call["cancelable"] is False
    _simulate_success(call)

    view = _regression_view(dlg)
    assert view.result().error is None
    assert view.result().predictors == ("x", "g")
    assert dlg.config_widgets[-1] is config  # the config is never rebuilt


def test_changing_the_target_refits_without_it_as_a_predictor(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory)
    config = _regression_config(dlg)
    _check_predictors(config, "x", "z")
    config._apply_button.click()  # noqa: SLF001
    _simulate_success(async_ops.last_call)

    combo = config._target_combo  # noqa: SLF001
    combo.setCurrentIndex(combo.findText("x"))

    call = async_ops.last_call
    assert call["scope"] == "analysis:regression:x:z"
    _simulate_success(call)
    assert _regression_view(dlg).result().target == "x"
    assert _regression_view(dlg).result().predictors == ("z",)


def test_stale_regression_result_is_discarded(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory)
    config = _regression_config(dlg)

    _check_predictors(config, "x")
    config._apply_button.click()  # noqa: SLF001
    first = async_ops.last_call
    _check_predictors(config, "z")
    config._apply_button.click()  # noqa: SLF001
    second = async_ops.last_call
    views_before = len(dlg.content_widgets)

    _simulate_success(first)
    assert len(dlg.content_widgets) == views_before

    _simulate_success(second)
    assert _regression_view(dlg).result().predictors == ("z",)


def test_applying_logistic_regression_uses_glm_view_and_keeps_linear_defaults(dialog_factory):
    rng = np.random.default_rng(21)
    size = 120
    x = rng.normal(size=size)
    probability = 1 / (1 + np.exp(-(-0.2 + 0.7 * x)))
    df = pd.DataFrame({
        "binary": rng.binomial(1, probability, size=size),
        "x": x,
        "z": rng.normal(size=size),
    })
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory, df)
    config = _regression_config(dlg)
    config._model_combo.setCurrentIndex(  # noqa: SLF001
        config._model_combo.findData(RegressionModel.LOGISTIC.value)  # noqa: SLF001
    )
    assert config.current_model() is RegressionModel.LOGISTIC
    assert config.current_target() == "binary"
    _check_predictors(config, "x")
    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:regression:logistic:binary:x"
    _simulate_success(call)

    view = dlg.content_widget()
    assert isinstance(view, GeneralizedRegressionView)
    assert view.result().error is None
    assert view.result().model is RegressionModel.LOGISTIC
    assert view.result().target == "binary"
    assert dlg.config_widgets[-1] is config


@pytest.mark.parametrize("model", [RegressionModel.POISSON, RegressionModel.NEGATIVE_BINOMIAL])
def test_applying_count_regression_keeps_selected_model(dialog_factory, model):
    rng = np.random.default_rng(22)
    size = 120
    x = rng.normal(size=size)
    mean = np.exp(0.3 + 0.4 * x)
    counts = rng.poisson(mean) if model is RegressionModel.POISSON else rng.negative_binomial(2, 2 / (2 + mean))
    df = pd.DataFrame({"count": counts, "x": x})
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory, df)
    config = _regression_config(dlg)
    config._model_combo.setCurrentIndex(  # noqa: SLF001
        config._model_combo.findData(model.value)  # noqa: SLF001
    )
    _check_predictors(config, "x")
    config._apply_button.click()  # noqa: SLF001
    _simulate_success(async_ops.last_call)

    view = dlg.content_widget()
    assert isinstance(view, GeneralizedRegressionView)
    assert view.result().error is None
    assert view.result().model is model
    assert view.result().plot_data is not None
    canvas = view.findChild(FigureCanvasQTAgg)
    assert canvas is not None
    assert len(canvas.figure.axes) == 2
    assert len(canvas.figure.axes[0].collections[0].get_offsets()) == size
    assert dlg.config_widgets[-1] is config


def test_stale_glm_result_is_discarded_after_switching_models(dialog_factory):
    rng = np.random.default_rng(23)
    size = 120
    x = rng.normal(size=size)
    probability = 1 / (1 + np.exp(-(-0.2 + 0.7 * x)))
    df = pd.DataFrame({"binary": rng.binomial(1, probability, size=size), "x": x, "z": rng.normal(size=size)})
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_regression(async_ops, dialog_factory, df)
    config = _regression_config(dlg)
    config._model_combo.setCurrentIndex(  # noqa: SLF001
        config._model_combo.findData(RegressionModel.LOGISTIC.value)  # noqa: SLF001
    )
    _check_predictors(config, "x")
    config._apply_button.click()  # noqa: SLF001
    stale_call = async_ops.last_call

    config._model_combo.setCurrentIndex(  # noqa: SLF001
        config._model_combo.findData(RegressionModel.LINEAR.value)  # noqa: SLF001
    )
    _simulate_success(stale_call)

    _assert_apply_prompt(ctrl, dlg)


def test_cox_regression_uses_apply_first_background_path_and_survival_view(dialog_factory):
    rng = np.random.default_rng(27)
    size = 160
    x = rng.normal(size=size)
    event_time = rng.exponential(scale=np.exp(-0.35 * x), size=size)
    censor_time = rng.exponential(scale=1.6, size=size)
    event = event_time <= censor_time
    df = pd.DataFrame({
        "duration": np.minimum(event_time, censor_time),
        "event": event.astype(int),
        "x": x,
    })
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory, df)
    config = _regression_config(dlg)
    config._model_combo.setCurrentIndex(config._model_combo.findData(RegressionModel.COX.value))  # noqa: SLF001

    assert async_ops.calls == []
    assert config.current_duration() == "duration"
    assert config.current_event() == "event"
    _check_predictors(config, "x")
    assert async_ops.calls == []
    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:regression:cox:duration:event:x"
    _simulate_success(call)

    view = dlg.content_widget()
    assert isinstance(view, SurvivalRegressionView)
    assert view.result().error is None
    assert dlg.config_widgets[-1] is config
    assert view.result().plot_data is not None
    assert view.result().plot_data.at_risk[0] == size
    canvas = view.findChild(FigureCanvasQTAgg)
    assert canvas is not None
    assert len(canvas.figure.axes) == 3
    assert canvas.figure.axes[0].get_xscale() == "log"
    np.testing.assert_array_equal(
        canvas.figure.axes[1].lines[0].get_xdata(),
        view.result().plot_data.times,
    )


def test_cox_job_is_stale_after_switching_models_away_and_back(dialog_factory):
    rng = np.random.default_rng(29)
    size = 150
    x = rng.normal(size=size)
    event_time = rng.exponential(scale=np.exp(-0.3 * x), size=size)
    censor_time = rng.exponential(scale=1.5, size=size)
    event = event_time <= censor_time
    df = pd.DataFrame({
        "duration": np.minimum(event_time, censor_time),
        "event": event.astype(int),
        "binary": rng.binomial(1, 0.5, size=size),
        "x": x,
    })
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory, df)
    config = _regression_config(dlg)
    config._model_combo.setCurrentIndex(config._model_combo.findData(RegressionModel.COX.value))  # noqa: SLF001
    _check_predictors(config, "x")
    config._apply_button.click()  # noqa: SLF001
    call = async_ops.last_call

    config._model_combo.setCurrentIndex(config._model_combo.findData(RegressionModel.LOGISTIC.value))  # noqa: SLF001
    config._model_combo.setCurrentIndex(config._model_combo.findData(RegressionModel.COX.value))  # noqa: SLF001
    views_before = len(dlg.content_widgets)
    _simulate_success(call)

    assert len(dlg.content_widgets) == views_before
    assert not isinstance(dlg.content_widget(), SurvivalRegressionView)


def test_cox_job_is_stale_after_duration_change_and_category_round_trip(dialog_factory):
    rng = np.random.default_rng(31)
    size = 150
    x = rng.normal(size=size)
    event_time = rng.exponential(scale=np.exp(-0.3 * x), size=size)
    censor_time = rng.exponential(scale=1.5, size=size)
    event = event_time <= censor_time
    df = pd.DataFrame({
        "duration": np.minimum(event_time, censor_time),
        "duration2": np.minimum(event_time, censor_time) + 0.25,
        "event": event.astype(int),
        "binary": rng.binomial(1, 0.5, size=size),
        "x": x,
    })
    async_ops = DummyAsyncOps()
    _, dlg = _open_regression(async_ops, dialog_factory, df)
    config = _regression_config(dlg)
    config._model_combo.setCurrentIndex(config._model_combo.findData(RegressionModel.COX.value))  # noqa: SLF001
    _check_predictors(config, "x")
    config._apply_button.click()  # noqa: SLF001
    stale_call = async_ops.last_call

    config._duration_combo.setCurrentIndex(config._duration_combo.findText("duration2"))  # noqa: SLF001
    dlg.category_changed.emit(AnalysisCategory.STATISTICS.value)
    dlg.category_changed.emit(AnalysisCategory.REGRESSION.value)
    assert _regression_config(dlg) is not config
    views_before = len(dlg.content_widgets)
    _simulate_success(stale_call)

    assert len(dlg.content_widgets) == views_before
    assert not isinstance(dlg.content_widget(), SurvivalRegressionView)


# ----------------------------------------------------------------------
# Outlier Explorer
# ----------------------------------------------------------------------


def _outliers_df() -> pd.DataFrame:
    return pd.DataFrame({
        "a": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0, 9.0, 100.0],
        "b": [float(i) for i in range(10)],
        "text": list("abcdefghij"),
    })


def _open_outliers(async_ops: DummyAsyncOps, dialog_factory, df: pd.DataFrame | None = None):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=10, column_count=3)
    ctrl = AnalysisController(
        results=DummyResults(
            datasets=[dataset],
            active_tab_id="t1",
            dfs={"t1": _outliers_df() if df is None else df},
        ),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.OUTLIERS
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.OUTLIERS.value)
    return ctrl, dlg


def _outliers_config(dlg: DummyAnalysisDialog) -> OutliersConfigWidget:
    config = dlg.config_widgets[-1]
    assert isinstance(config, OutliersConfigWidget)
    return config


def _outliers_view(dlg: DummyAnalysisDialog) -> OutliersView:
    view = dlg.content_widget()
    assert isinstance(view, OutliersView)
    return view


def _multivariate_outliers_config(dlg: DummyAnalysisDialog) -> MultivariateOutliersConfigWidget:
    config = dlg.config_widget()
    assert isinstance(config, MultivariateOutliersConfigWidget)
    return config


def _multivariate_outliers_view(dlg: DummyAnalysisDialog) -> MultivariateOutliersView:
    view = dlg.content_widget()
    assert isinstance(view, MultivariateOutliersView)
    return view


def _detail_column(view: OutliersView) -> str:
    detail = view.column_detail()
    assert detail is not None
    return detail.summary.column


def _select_outlier_method(config: OutliersConfigWidget, method: OutlierMethod) -> None:
    combo = config._method_combo  # noqa: SLF001
    combo.setCurrentIndex(combo.findData(method))


def _apply_outliers(async_ops: DummyAsyncOps, dlg: DummyAnalysisDialog) -> dict:
    """Click the univariate Apply and return the summary job's call, without completing it."""
    _outliers_config(dlg)._apply_button.click()  # noqa: SLF001
    return async_ops.last_call


def _open_applied_outliers(async_ops: DummyAsyncOps, dialog_factory):
    ctrl, dlg = _open_outliers(async_ops, dialog_factory)
    _simulate_success(_apply_outliers(async_ops, dlg))
    return ctrl, dlg


def _switch_to_multivariate(dlg: DummyAnalysisDialog) -> MultivariateOutliersConfigWidget:
    _outliers_config(dlg)._mode_combo.setCurrentIndex(1)  # noqa: SLF001
    return _multivariate_outliers_config(dlg)


def test_outliers_initialize_without_a_job_and_prompt_for_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_outliers(async_ops, dialog_factory)

    assert async_ops.calls == []
    _assert_apply_prompt(ctrl, dlg)
    config = _outliers_config(dlg)
    assert config.summary_configuration() == (OutlierMethod.IQR, 1.5)
    assert config.is_column_selection_enabled() is False


def test_outlier_column_change_before_the_first_apply_starts_no_job(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_outliers(async_ops, dialog_factory)

    _outliers_config(dlg).set_column("b")

    assert async_ops.calls == []


def test_outlier_apply_runs_as_a_background_job_with_a_standard_overlay(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_outliers(async_ops, dialog_factory)

    call = _apply_outliers(async_ops, dlg)

    assert call["scope"] == "analysis:outliers:summary:iqr:1.5"
    assert call["target"] is dlg.content_panel()
    assert call["cancelable"] is False
    assert call["indeterminate"] is True


def test_first_outlier_apply_details_the_top_ranked_column_and_enables_column_selection(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)
    config.set_column("b")  # the disabled picker only holds a placeholder column

    _simulate_success(_apply_outliers(async_ops, dlg))

    assert _detail_column(_outliers_view(dlg)) == "a"
    assert config.current_column() == "a"
    assert config.is_column_selection_enabled() is True
    assert async_ops.last_call["scope"] == "analysis:outliers:summary:iqr:1.5"  # no extra column job


def test_outliers_without_numeric_columns_show_error_without_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_outliers(async_ops, dialog_factory, pd.DataFrame({"t": ["x", "y", "z"]}))

    assert async_ops.calls == []
    view = _outliers_view(dlg)
    assert view.table() is None
    assert view.column_detail() is None
    assert dlg.config_widgets[-1] is None


def test_outlier_method_and_threshold_changes_start_no_job_until_applied(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)
    jobs_before = len(async_ops.calls)

    _select_outlier_method(config, OutlierMethod.Z_SCORE)
    config._threshold_spin.setValue(2.5)  # noqa: SLF001

    assert len(async_ops.calls) == jobs_before
    assert config.summary_configuration() == (OutlierMethod.IQR, 1.5)
    assert _outliers_view(dlg).configuration() == (OutlierMethod.IQR, 1.5)


def test_switching_outliers_to_multivariate_prompts_without_fitting(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_applied_outliers(async_ops, dialog_factory)
    jobs_before = len(async_ops.calls)

    multivariate_config = _switch_to_multivariate(dlg)

    assert len(async_ops.calls) == jobs_before
    _assert_apply_prompt(ctrl, dlg)
    assert multivariate_config.analysis_configuration() == (
        ("a", "b"),
        MultivariateOutlierMethod.ISOLATION_FOREST,
        True,
        0.05,
        20,
    )


def test_switching_to_multivariate_with_too_few_numeric_columns_explains_why(dialog_factory):
    async_ops = DummyAsyncOps()
    df = pd.DataFrame({"a": [1.0, 2.0, 3.0], "text": ["x", "y", "z"]})
    _, dlg = _open_outliers(async_ops, dialog_factory, df)

    multivariate_config = _switch_to_multivariate(dlg)

    assert async_ops.calls == []
    view = _multivariate_outliers_view(dlg)
    assert view._result.error is MultivariateOutlierError.NOT_ENOUGH_NUMERIC_COLUMNS  # noqa: SLF001
    assert dlg.config_widget() is multivariate_config


def test_switching_to_multivariate_shows_error_when_the_dataset_cannot_be_reloaded(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_outliers(async_ops, dialog_factory)
    ctrl._results._dfs.clear()  # noqa: SLF001 - the dataset disappeared

    _outliers_config(dlg)._mode_combo.setCurrentIndex(1)  # noqa: SLF001

    assert async_ops.calls == []
    assert dlg.placeholder_calls[-1] == ctrl._tr(ctrl.TR_ANALYSIS_ERROR)  # noqa: SLF001
    assert dlg.config_widget() is None


def test_first_multivariate_apply_fits_the_default_configuration(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_outliers(async_ops, dialog_factory)
    config = _switch_to_multivariate(dlg)

    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:outliers:multivariate:isolation_forest:standardized:contamination:0.05:a:b"
    assert call["target"] is dlg.content_panel()
    assert call["cancelable"] is False
    _simulate_success(call)
    assert _multivariate_outliers_view(dlg).configuration() == config.analysis_configuration()


def test_applying_multivariate_configuration_refits_and_retains_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_outliers(async_ops, dialog_factory)
    config = _switch_to_multivariate(dlg)

    config._method_combo.setCurrentIndex(config._method_combo.findData(MultivariateOutlierMethod.LOCAL_OUTLIER_FACTOR))  # noqa: SLF001
    config._contamination_spin.setValue(10.0)  # noqa: SLF001
    config._neighbors_spin.setValue(3)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert (
        call["scope"]
        == "analysis:outliers:multivariate:local_outlier_factor:standardized:contamination:0.1:neighbors:3:a:b"
    )
    _simulate_success(call)

    assert dlg.config_widget() is config
    assert _multivariate_outliers_view(dlg).configuration() == (
        ("a", "b"),
        MultivariateOutlierMethod.LOCAL_OUTLIER_FACTOR,
        True,
        0.1,
        3,
    )


def test_stale_multivariate_fit_is_discarded(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_outliers(async_ops, dialog_factory)
    config = _switch_to_multivariate(dlg)

    config._contamination_spin.setValue(10.0)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001
    first = async_ops.last_call
    config._contamination_spin.setValue(20.0)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001
    second = async_ops.last_call
    views_before = len(dlg.content_widgets)

    _simulate_success(first)
    assert len(dlg.content_widgets) == views_before

    _simulate_success(second)
    assert _multivariate_outliers_view(dlg).configuration()[3] == 0.2


def test_switching_multivariate_mode_back_restores_the_univariate_prompt(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _switch_to_multivariate(dlg)
    jobs_before = len(async_ops.calls)

    config._mode_combo.setCurrentIndex(0)  # noqa: SLF001

    assert len(async_ops.calls) == jobs_before
    _assert_apply_prompt(ctrl, dlg)
    assert isinstance(dlg.config_widget(), OutliersConfigWidget)


def test_applied_outlier_method_recomputes_the_summary_and_keeps_the_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)
    first_view = _outliers_view(dlg)
    configs_before = len(dlg.config_widgets)
    config.set_column("b")
    _simulate_success(async_ops.last_call)

    _select_outlier_method(config, OutlierMethod.Z_SCORE)
    call = _apply_outliers(async_ops, dlg)

    assert call["scope"] == "analysis:outliers:summary:z_score:3.0"
    _simulate_success(call)

    view = _outliers_view(dlg)
    assert view is not first_view
    assert view.configuration() == (OutlierMethod.Z_SCORE, 3.0)
    detail = view.column_detail()
    assert detail is not None
    assert detail.method is OutlierMethod.Z_SCORE
    assert detail.summary.column == "b"  # the selected column is kept
    assert len(dlg.config_widgets) == configs_before


def test_stale_outlier_summary_is_discarded_when_applied_again(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)
    spin = config._threshold_spin  # noqa: SLF001

    spin.setValue(2.0)
    first = _apply_outliers(async_ops, dlg)
    spin.setValue(3.0)
    second = _apply_outliers(async_ops, dlg)

    contents_before = len(dlg.content_widgets)
    _simulate_success(first)
    assert len(dlg.content_widgets) == contents_before

    _simulate_success(second)
    assert _outliers_view(dlg).configuration() == (OutlierMethod.IQR, 3.0)


def test_outlier_column_change_recomputes_only_the_detail_over_the_detail_panel(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)
    view = _outliers_view(dlg)
    contents_before = len(dlg.content_widgets)

    config.set_column("b")

    call = async_ops.last_call
    assert call["scope"] == "analysis:outliers:column:iqr:1.5:b"
    assert call["target"] is view.detail_panel()
    assert call["cancelable"] is False

    _simulate_success(call)

    assert len(dlg.content_widgets) == contents_before  # view updated in place
    assert _detail_column(view) == "b"


def test_clicking_an_outlier_summary_row_selects_the_column(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)
    table = _outliers_view(dlg).table()
    assert table is not None

    table.cellClicked.emit(1, 0)

    assert config.current_column() == "b"
    assert async_ops.last_call["scope"] == "analysis:outliers:column:iqr:1.5:b"


def test_stale_outlier_column_recompute_is_discarded_when_the_column_changes_again(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)
    view = _outliers_view(dlg)
    initial_detail = view.column_detail()

    config.set_column("b")
    first = async_ops.last_call
    config.set_column("a")
    second = async_ops.last_call

    _simulate_success(first)
    assert view.column_detail() is initial_detail

    _simulate_success(second)
    assert _detail_column(view) == "a"


def test_outlier_column_change_while_the_summary_is_computing_is_caught_up_afterwards(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)

    _select_outlier_method(config, OutlierMethod.MODIFIED_Z_SCORE)
    summary_call = _apply_outliers(async_ops, dlg)
    dlg.show_placeholder("computing")  # no view is displayed while the summary job runs
    jobs_before = len(async_ops.calls)

    config.set_column("b")
    assert len(async_ops.calls) == jobs_before  # nothing to update yet

    _simulate_success(summary_call)
    assert _detail_column(_outliers_view(dlg)) == "a"

    column_call = async_ops.last_call
    assert column_call["scope"] == "analysis:outliers:column:modified_z_score:3.5:b"
    _simulate_success(column_call)
    assert _detail_column(_outliers_view(dlg)) == "b"


def test_outlier_column_recompute_uses_the_displayed_summary_configuration(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_applied_outliers(async_ops, dialog_factory)
    config = _outliers_config(dlg)

    # Z-score is applied, but its summary job hasn't finished: the view still shows IQR.
    _select_outlier_method(config, OutlierMethod.Z_SCORE)
    _apply_outliers(async_ops, dlg)
    config.set_column("b")

    assert async_ops.last_call["scope"] == "analysis:outliers:column:iqr:1.5:b"


# ----------------------------------------------------------------------
# Principal Component Analysis
# ----------------------------------------------------------------------


def _pca_df() -> pd.DataFrame:
    return pd.DataFrame({
        "a": [1.0, 2.0, 3.0, 4.0, 5.0],
        "b": [2.0, 4.0, 7.0, 8.0, 11.0],
        "c": [9.0, 2.0, 5.0, 3.0, 7.0],
        "text": list("abcde"),
    })


def _open_pca(async_ops: DummyAsyncOps, dialog_factory, df: pd.DataFrame | None = None):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=5, column_count=4)
    ctrl = AnalysisController(
        results=DummyResults(
            datasets=[dataset],
            active_tab_id="t1",
            dfs={"t1": _pca_df() if df is None else df},
        ),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.PCA
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.PCA.value)
    return ctrl, dlg


def _pca_config(dlg: DummyAnalysisDialog) -> PCAConfigWidget:
    config = dlg.config_widgets[-1]
    assert isinstance(config, PCAConfigWidget)
    return config


def _pca_view(dlg: DummyAnalysisDialog) -> PCAView:
    view = dlg.content_widget()
    assert isinstance(view, PCAView)
    return view


def _set_pca_checked(config: PCAConfigWidget, column: str, checked: bool) -> None:
    items = config._column_list.findItems(column, Qt.MatchFlag.MatchExactly)  # noqa: SLF001
    assert len(items) == 1
    items[0].setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)


def test_pca_initializes_without_a_job_and_prompts_for_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_pca(async_ops, dialog_factory)

    assert async_ops.calls == []
    _assert_apply_prompt(ctrl, dlg)
    config = _pca_config(dlg)
    assert config.analysis_configuration() == (("a", "b", "c"), True)
    assert config.checked_columns() == ("a", "b", "c")


def test_first_pca_apply_fits_the_default_configuration(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_pca(async_ops, dialog_factory)

    _pca_config(dlg)._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:pca:fit:standardized:a:b:c"
    _simulate_success(call)
    assert _pca_view(dlg).configuration() == (("a", "b", "c"), True)


def test_pca_without_enough_numeric_columns_shows_error_without_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_pca(async_ops, dialog_factory, pd.DataFrame({"a": [1.0, 2.0], "text": ["x", "y"]}))

    assert async_ops.calls == []
    view = _pca_view(dlg)
    assert view.loadings_table() is None
    assert PCAError.NOT_ENOUGH_NUMERIC_COLUMNS.value in str(view._result.error)  # noqa: SLF001
    assert dlg.config_widgets[-1] is None


def test_apply_recomputes_pca_and_keeps_the_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_pca(async_ops, dialog_factory)
    config = _pca_config(dlg)
    prompt = dlg.content_widget()
    configs_before = len(dlg.config_widgets)

    _set_pca_checked(config, "c", False)
    config._standardize_checkbox.setChecked(False)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:pca:fit:raw:a:b"
    assert call["target"] is dlg.content_panel()
    assert call["cancelable"] is False

    _simulate_success(call)

    view = _pca_view(dlg)
    assert view is not prompt
    assert view.configuration() == (("a", "b"), False)
    assert len(dlg.config_widgets) == configs_before


def test_stale_pca_recompute_is_discarded_when_apply_changes_again(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_pca(async_ops, dialog_factory)
    config = _pca_config(dlg)

    _set_pca_checked(config, "c", False)
    config._apply_button.click()  # noqa: SLF001
    first = async_ops.last_call

    config._standardize_checkbox.setChecked(False)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001
    second = async_ops.last_call

    contents_before = len(dlg.content_widgets)
    _simulate_success(first)
    assert len(dlg.content_widgets) == contents_before

    _simulate_success(second)
    assert _pca_view(dlg).configuration() == (("a", "b"), False)


# ----------------------------------------------------------------------
# Clustering
# ----------------------------------------------------------------------


def _clustering_df() -> pd.DataFrame:
    return pd.DataFrame({
        "x": [0.0, 0.1, -0.1, 10.0, 10.1, 9.9],
        "y": [0.0, -0.1, 0.1, 10.0, 10.1, 9.9],
        "z": [1.0, 1.1, 0.9, 5.0, 5.1, 4.9],
        "text": list("abcdef"),
    })


def _open_clustering(async_ops: DummyAsyncOps, dialog_factory, df: pd.DataFrame | None = None):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=6, column_count=4)
    ctrl = AnalysisController(
        results=DummyResults(
            datasets=[dataset],
            active_tab_id="t1",
            dfs={"t1": _clustering_df() if df is None else df},
        ),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())

    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.CLUSTERING
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.CLUSTERING.value)
    return ctrl, dlg


def _clustering_config(dlg: DummyAnalysisDialog) -> ClusteringConfigWidget:
    config = dlg.config_widgets[-1]
    assert isinstance(config, ClusteringConfigWidget)
    return config


def _clustering_view(dlg: DummyAnalysisDialog) -> ClusteringView:
    view = dlg.content_widget()
    assert isinstance(view, ClusteringView)
    return view


def _set_clustering_checked(config: ClusteringConfigWidget, column: str, checked: bool) -> None:
    items = config._column_list.findItems(column, Qt.MatchFlag.MatchExactly)  # noqa: SLF001
    assert len(items) == 1
    items[0].setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)


def test_clustering_initializes_without_a_job_and_prompts_for_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_clustering(async_ops, dialog_factory)

    assert async_ops.calls == []
    _assert_apply_prompt(ctrl, dlg)
    config = _clustering_config(dlg)
    assert config.analysis_configuration() == (("x", "y", "z"), ClusteringMethod.K_MEANS, True, 3, 0.5, 5)
    assert config.checked_columns() == ("x", "y", "z")


def test_first_clustering_apply_fits_the_default_configuration(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_clustering(async_ops, dialog_factory)

    _clustering_config(dlg)._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:clustering:fit:k_means:standardized:clusters:3:x:y:z"
    _simulate_success(call)
    assert _clustering_view(dlg).configuration() == (("x", "y", "z"), ClusteringMethod.K_MEANS, True, 3, 0.5, 5)


def test_clustering_without_numeric_columns_shows_error_without_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_clustering(async_ops, dialog_factory, pd.DataFrame({"text": ["x", "y"]}))

    assert async_ops.calls == []
    assert _clustering_view(dlg).cluster_table() is None
    assert dlg.config_widgets[-1] is None


def test_clustering_apply_recomputes_dbscan_and_keeps_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_clustering(async_ops, dialog_factory)
    config = _clustering_config(dlg)
    prompt = dlg.content_widget()
    configs_before = len(dlg.config_widgets)

    _set_clustering_checked(config, "z", False)
    config._standardize_checkbox.setChecked(False)  # noqa: SLF001
    combo = config._method_combo  # noqa: SLF001
    combo.setCurrentIndex(combo.findData(ClusteringMethod.DBSCAN))
    config._epsilon_spin.setValue(0.75)  # noqa: SLF001
    config._min_samples_spin.setValue(2)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:clustering:fit:dbscan:raw:eps:0.75:min_samples:2:x:y"
    assert call["target"] is dlg.content_panel()
    assert call["cancelable"] is False

    _simulate_success(call)

    view = _clustering_view(dlg)
    assert view is not prompt
    assert view.configuration() == (("x", "y"), ClusteringMethod.DBSCAN, False, 3, 0.75, 2)
    assert len(dlg.config_widgets) == configs_before


def test_stale_clustering_recompute_is_discarded_after_another_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_clustering(async_ops, dialog_factory)
    config = _clustering_config(dlg)

    config._cluster_count_spin.setValue(2)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001
    first = async_ops.last_call
    config._cluster_count_spin.setValue(4)  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001
    second = async_ops.last_call

    contents_before = len(dlg.content_widgets)
    _simulate_success(first)
    assert len(dlg.content_widgets) == contents_before

    _simulate_success(second)
    assert _clustering_view(dlg).configuration()[3] == 4


# ----------------------------------------------------------------------
# Time Series Explorer
# ----------------------------------------------------------------------


def _time_series_df() -> pd.DataFrame:
    return pd.DataFrame({
        "when": pd.date_range("2025-01-01", periods=21, freq="D"),
        "value": range(21),
        "other": range(100, 121),
    })


def _open_time_series(async_ops: DummyAsyncOps, dialog_factory, df: pd.DataFrame | None = None):
    dataset = DatasetRef(tab_id="t1", title="Sheet1", row_count=21, column_count=3)
    ctrl = AnalysisController(
        results=DummyResults(
            datasets=[dataset],
            active_tab_id="t1",
            dfs={"t1": _time_series_df() if df is None else df},
        ),
        async_ops=async_ops,
    )
    _open_and_flush(ctrl, QWidget())
    dlg = dialog_factory[0]
    dlg._selected_category = AnalysisCategory.TIME_SERIES
    dlg._selected_dataset_tab_id = "t1"
    dlg.category_changed.emit(AnalysisCategory.TIME_SERIES.value)
    return ctrl, dlg


def _time_series_config(dlg: DummyAnalysisDialog) -> TimeSeriesConfigWidget:
    config = dlg.config_widgets[-1]
    assert isinstance(config, TimeSeriesConfigWidget)
    return config


def _time_series_view(dlg: DummyAnalysisDialog) -> TimeSeriesView:
    view = dlg.content_widget()
    assert isinstance(view, TimeSeriesView)
    return view


def test_time_series_initializes_without_a_job_and_prompts_for_apply(dialog_factory):
    async_ops = DummyAsyncOps()
    ctrl, dlg = _open_time_series(async_ops, dialog_factory)

    assert async_ops.calls == []
    _assert_apply_prompt(ctrl, dlg)
    config = _time_series_config(dlg)
    assert config.analysis_configuration() == ("when", "value", None, None, DecompositionModel.ADDITIVE)
    assert config.pending_configuration() == config.analysis_configuration()


def test_first_time_series_apply_analyzes_the_default_configuration(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_time_series(async_ops, dialog_factory)

    _time_series_config(dlg)._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:time_series:fit:when:value:original:auto:additive"
    _simulate_success(call)
    assert _time_series_view(dlg).configuration() == ("when", "value", None, 7, DecompositionModel.ADDITIVE)


def test_time_series_without_datetime_has_no_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_time_series(async_ops, dialog_factory, pd.DataFrame({"value": [1.0, 2.0, 3.0]}))

    assert async_ops.calls == []
    assert isinstance(dlg.content_widget(), TimeSeriesView)
    assert dlg.config_widgets[-1] is None


def test_time_series_apply_recomputes_and_keeps_config(dialog_factory):
    async_ops = DummyAsyncOps()
    _, dlg = _open_time_series(async_ops, dialog_factory)
    config = _time_series_config(dlg)
    configs_before = len(dlg.config_widgets)

    config._value_combo.setCurrentIndex(config._value_combo.findText("other"))  # noqa: SLF001
    config._frequency_combo.setCurrentIndex(config._frequency_combo.findData("W"))  # noqa: SLF001
    config._auto_period.setChecked(False)  # noqa: SLF001
    config._period_spin.setValue(3)  # noqa: SLF001
    config._model_combo.setCurrentIndex(config._model_combo.findData(DecompositionModel.MULTIPLICATIVE))  # noqa: SLF001
    config._apply_button.click()  # noqa: SLF001

    call = async_ops.last_call
    assert call["scope"] == "analysis:time_series:fit:when:other:W:3:multiplicative"
    _simulate_success(call)
    assert _time_series_view(dlg).configuration() == ("when", "other", "W", 3, DecompositionModel.MULTIPLICATIVE)
    assert len(dlg.config_widgets) == configs_before
