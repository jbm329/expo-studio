from __future__ import annotations

from pathlib import Path

import pytest
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg
from matplotlib.figure import Figure
from openpyxl import load_workbook
from PyQt6.QtCore import QCoreApplication, QTranslator
from PyQt6.QtWidgets import QApplication

from expo_jbm329.gui.dialogs.analysis.pca_view import PCAView
from expo_jbm329.services.analysis.categories import AnalysisCategory
from expo_jbm329.services.analysis.pca_charts import draw_score_plot, draw_scree_plot, render_pca_charts
from expo_jbm329.services.analysis.pca_export import PCAExportComponent, PCAExportSnapshot, pca_export_tables
from expo_jbm329.services.analysis.report import (
    ReportDocument,
    capture_report_labels,
    create_report_entry,
    prepare_report_content,
    render_prepared_report_html,
)
from expo_jbm329.services.excel_chart import ExcelChartCancelledError
from expo_jbm329.services.file_writer import FileWriter
from expo_jbm329.utils.dataset_ref import DatasetRef
from expo_jbm329.utils.format_utils import fmt_pct
from tests.analysis.test_pca_export import CHART_LABELS, pca_snapshot

_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"


@pytest.mark.parametrize("locale", ["en", "sv"])
@pytest.mark.parametrize("sampled", [False, True])
@pytest.mark.parametrize("standardize", [False, True])
def test_live_worker_and_html_share_captured_localized_semantics(locale, sampled, standardize, monkeypatch):
    result = pca_snapshot(sampled=sampled, standardize=standardize).result
    translator = QTranslator()
    catalog = Path(__file__).parents[2] / "src" / "expo_jbm329" / "i18n" / "locales" / f"app_{locale}.qm"
    assert translator.load(str(catalog))
    app = QApplication.instance()
    assert app is not None
    app.installTranslator(translator)
    try:
        view = PCAView(result)
        labels = view.chart_labels()
        report_labels = capture_report_labels(QCoreApplication.translate)
        expected_titles = ("Scree plot", "PC1 vs PC2") if locale == "en" else ("Scree-diagram", "PC1 mot PC2")
        assert (labels.scree_metadata.accessible_name, labels.scatter_metadata.accessible_name) == expected_titles
        assert labels.pc1_axis == f"PC1 ({fmt_pct(result.explained_variance_ratio[0])})"
        assert labels.pc2_axis == f"PC2 ({fmt_pct(result.explained_variance_ratio[1])})"
        assert labels.component_axis == ("Component" if locale == "en" else "Komponent")
        assert labels.variance_axis == ("Explained variance (%)" if locale == "en" else "Förklarad varians (%)")
        assert labels.explained == ("Explained variance" if locale == "en" else "Förklarad varians")
        assert labels.cumulative == ("Cumulative" if locale == "en" else "Kumulativ")
        caption = (
            "The PC1-vs-PC2 plot shows a deterministic sample of 7 rows."
            if locale == "en"
            else "PC1-mot-PC2-diagrammet visar ett deterministiskt urval av 7 rader."
        )
        assert labels.scatter_metadata.captions == ((caption,) if sampled else ())
    finally:
        app.removeTranslator(translator)

    def semantics(ax):
        legend = ax.get_legend()
        return (
            ax.get_title(),
            ax.get_xlabel(),
            ax.get_ylabel(),
            tuple(text.get_text() for text in legend.get_texts()) if legend is not None else (),
            tuple(ax.get_xticks()) if ax.get_title() == expected_titles[0] else (),
            tuple(tuple(line.get_ydata()) for line in ax.get_lines()),
        )

    live_axes = {
        canvas.figure.axes[0].get_title(): canvas.figure.axes[0] for canvas in view.findChildren(FigureCanvasQTAgg)
    }
    expected = [semantics(live_axes[title]) for title in expected_titles]
    rendered = []
    original_print_png = FigureCanvasAgg.print_png

    def inspect_png(canvas, output, *args, **kwargs):
        rendered.append(semantics(canvas.figure.axes[0]))
        original_print_png(canvas, output, *args, **kwargs)

    monkeypatch.setattr(FigureCanvasAgg, "print_png", inspect_png)
    snapshot = PCAExportSnapshot(result, view.chart_labels())
    assert snapshot.chart_labels is labels
    document = ReportDocument()
    entry = create_report_entry(AnalysisCategory.PCA, DatasetRef("pca", "PCA dataset", 8, 3), snapshot)
    document.upsert(entry.with_selection(("charts",), ""))
    prepared = prepare_report_content(document, labels=report_labels)
    assert rendered == expected
    charts = prepared.entries[0].components[0].charts
    assert [chart.heading for chart in charts] == ["", ""]
    assert tuple(chart.accessible_name for chart in charts) == expected_titles
    assert charts[0].captions == ()
    assert charts[1].captions == ((caption,) if sampled else ())
    html = render_prepared_report_html(prepared)
    assert html.index(f'alt="{expected_titles[0]}"') < html.index(f'alt="{expected_titles[1]}"')
    assert 'alt=""' not in html
    assert "<h3>" + ("PCA plots" if locale == "en" else "PCA-diagram") + "</h3>" in html
    assert html.count("<figcaption>") == int(sampled)
    if sampled:
        assert f"<figcaption>\n<p>{caption}</p>\n</figcaption>" in html
    view.close()


def test_scree_plot_draws_component_bars_and_the_cumulative_line():
    result = pca_snapshot().result
    ax = Figure().add_subplot(111)

    draw_scree_plot(ax, result, CHART_LABELS)

    assert [bar.get_height() for bar in ax.patches] == pytest.approx([
        ratio * 100 for ratio in result.explained_variance_ratio
    ])
    (line,) = ax.get_lines()
    assert list(line.get_ydata()) == pytest.approx([ratio * 100 for ratio in result.cumulative_variance_ratio])
    assert ax.get_title() == "Scree plot"
    assert (ax.get_xlabel(), ax.get_ylabel()) == ("Component", "Explained variance (%)")
    assert [text.get_text() for text in ax.get_legend().get_texts()] == ["Cumulative", "Explained variance"]


def test_score_plot_draws_the_sampled_scores():
    result = pca_snapshot().result
    ax = Figure().add_subplot(111)

    draw_score_plot(ax, result, CHART_LABELS)

    (scores,) = ax.collections
    offsets = scores.get_offsets()
    assert offsets[:, 0].tolist() == pytest.approx(list(result.sample_pc1))
    assert offsets[:, 1].tolist() == pytest.approx(list(result.sample_pc2))
    assert ax.get_title() == "PC1 vs PC2"
    assert (ax.get_xlabel(), ax.get_ylabel()) == ("PC1 (80%)", "PC2 (15%)")
    horizontal, vertical = ax.get_lines()
    assert list(horizontal.get_ydata()) == [0, 0]
    assert list(vertical.get_xdata()) == [0, 0]


def test_render_returns_both_plots_with_captions_and_progress():
    progress: list[int] = []

    scree, scores = render_pca_charts(pca_snapshot(sampled=True), progress_cb=progress.append, cancel_cb=lambda: False)

    assert (scree.heading, scores.heading) == ("", "")
    assert (scree.accessible_name, scores.accessible_name) == ("Scree plot", "PC1 vs PC2")
    assert scree.captions == ()
    assert scores.captions == ("The PC1-vs-PC2 plot shows a deterministic sample of 7 rows.",)
    for image in (scree, scores):
        assert image.image_data.startswith(_PNG_SIGNATURE)
        assert (image.width_px, image.height_px) == (960, 600)
    assert progress == [50, 100]


@pytest.mark.parametrize("cancel_after", [0, 1, 2])
def test_render_honors_cancellation_before_each_image_and_at_the_end(cancel_after):
    calls = 0

    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls > cancel_after

    with pytest.raises(ExcelChartCancelledError):
        render_pca_charts(pca_snapshot(), cancel_cb=cancel)


@pytest.mark.parametrize("streaming", [False, True])
@pytest.mark.parametrize("sampled", [False, True])
def test_workbook_keeps_typed_tables_and_both_plots(tmp_path, streaming, sampled):
    snapshot = pca_snapshot(sampled=sampled)
    frames = pca_export_tables(
        snapshot,
        (PCAExportComponent.FEATURE_LOADINGS, PCAExportComponent.EXPLAINED_VARIANCE, PCAExportComponent.MODEL_SUMMARY),
    )
    sheets = {key.value: frame for key, frame in frames.items()}
    path = tmp_path / "pca.xlsx"

    FileWriter().save_excel_sheets(
        sheets,
        path,
        chart_images=render_pca_charts(snapshot),
        chart_sheet_name="Charts",
        streaming=streaming,
    )

    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["feature_loadings", "explained_variance", "model_summary", "Charts"]
        assert len(workbook["Charts"]._images) == 2
        chart_sheet = workbook["Charts"]
        assert chart_sheet["A1"].value is None
        cells = [cell.value for row in chart_sheet.iter_rows() for cell in row if cell.value is not None]
        assert cells == list(snapshot.chart_labels.scatter_metadata.captions)
        assert chart_sheet._images[0].anchor._from.row < chart_sheet._images[1].anchor._from.row
        loadings = workbook["feature_loadings"]
        assert [cell.value for cell in loadings[1]] == ["feature", "PC1", "PC2", "PC3"]
        assert loadings["A2"].value == "x"
        assert isinstance(loadings["B2"].value, float)
        for row, values in enumerate(snapshot.result.loadings, start=2):
            assert [cell.value for cell in loadings[row]][1:] == pytest.approx(values)
        variance = workbook["explained_variance"]
        assert [cell.value for cell in variance[1]] == ["component", "variance", "share", "cumulative_share"]
        assert variance["A2"].value == "PC1"
        assert isinstance(variance["C2"].value, float)
        assert [row[2].value for row in list(variance.rows)[1:]] == pytest.approx(
            snapshot.result.explained_variance_ratio
        )
        summary = workbook["model_summary"]
        assert summary["A2"].value == "features"
        assert summary["B2"].value == pytest.approx(3.0)
    finally:
        workbook.close()
