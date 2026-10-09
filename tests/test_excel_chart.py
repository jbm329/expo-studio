from __future__ import annotations

from dataclasses import replace
from typing import TYPE_CHECKING

import pandas as pd
import pytest
from openpyxl import load_workbook

from expo_jbm329.services.analysis.statistics import (
    StatisticsExportTable,
    analyze_descriptive_statistics,
    statistics_export_tables,
)
from expo_jbm329.services.analysis.statistics_charts import StatisticsChartLabels, render_statistics_charts
from expo_jbm329.services.excel_chart import ExcelChartImage
from expo_jbm329.services.file_writer import EXCEL_MAX_ROWS, ExportCancelledError, FileWriter

if TYPE_CHECKING:
    from pathlib import Path


@pytest.fixture
def chart_images() -> tuple[ExcelChartImage, ...]:
    result = analyze_descriptive_statistics(pd.DataFrame({"First": [1.0, 2.0, 3.0], "Second": [4.0, 5.0, 6.0]}))
    return render_statistics_charts(result.columns, StatisticsChartLabels("Histogram", "Boxplot", "No data"))


@pytest.mark.parametrize("streaming", [True, False])
@pytest.mark.parametrize("tables_selected", [True, False])
def test_named_excel_embeds_one_image_per_numeric_column_in_single_charts_sheet(
    tmp_path: Path,
    chart_images: tuple[ExcelChartImage, ...],
    streaming: bool,
    tables_selected: bool,
) -> None:
    path = tmp_path / "statistics.xlsx"
    tables = (
        statistics_export_tables(
            analyze_descriptive_statistics(pd.DataFrame({"First": [1.0, 2.0], "Second": [3.0, 4.0]})),
            (StatisticsExportTable.CONTINUOUS,),
        )
        if tables_selected
        else {}
    )

    FileWriter().save_excel_sheets(
        tables,
        path,
        chart_images=chart_images,
        chart_sheet_name="Charts",
        streaming=streaming,
    )

    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == (["Continuous", "Charts"] if tables_selected else ["Charts"])
        chart_sheet = workbook["Charts"]
        assert [chart_sheet.cell(row, 1).value for row in (1, 24)] == ["First", "Second"]
        assert len(chart_sheet._images) == 2
        assert [image.anchor._from.row for image in chart_sheet._images] == [1, 24]
        if tables_selected:
            assert workbook["Continuous"].cell(1, 1).value == "column"
            assert workbook["Continuous"].cell(2, 1).value == "First"
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [True, False])
def test_chart_image_writing_cancellation_preserves_existing_destination(
    tmp_path: Path,
    chart_images: tuple[ExcelChartImage, ...],
    streaming: bool,
) -> None:
    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"original")
    calls = 0
    cancel_on_call = 3 if streaming else 2

    def cancel() -> bool:
        nonlocal calls
        calls += 1
        return calls == cancel_on_call

    with pytest.raises(ExportCancelledError):
        FileWriter().save_excel_sheets(
            {},
            path,
            chart_images=chart_images,
            chart_sheet_name="Charts",
            streaming=streaming,
            cancel_cb=cancel,
        )

    assert path.read_bytes() == b"original"
    assert not list(tmp_path.glob(".existing-*.xlsx"))


@pytest.mark.parametrize(
    ("sheets", "chart_images", "chart_sheet_name"),
    [
        ({"Charts": pd.DataFrame({"x": [1]})}, (), "Charts"),
        ({}, (), "Charts"),
        ({}, (), None),
    ],
)
def test_chart_workbook_validation_rejects_missing_or_conflicting_content(
    sheets: dict[str, pd.DataFrame],
    chart_images: tuple[ExcelChartImage, ...],
    chart_sheet_name: str | None,
) -> None:
    with pytest.raises(ValueError):
        FileWriter().save_excel_sheets(
            sheets,
            "unused.xlsx",
            chart_images=chart_images,
            chart_sheet_name=chart_sheet_name,
        )


def test_excel_chart_image_validates_bytes_and_dimensions() -> None:
    with pytest.raises(ValueError):
        ExcelChartImage("Column", b"", 10, 10)
    with pytest.raises(ValueError):
        ExcelChartImage("Column", b"png", 0, 10)


def test_charts_sheet_row_limit_fails_explicitly(tmp_path: Path) -> None:
    path = tmp_path / "row-limit.xlsx"
    oversized = ExcelChartImage("Column", b"png", 1, EXCEL_MAX_ROWS * 20)

    with pytest.raises(ValueError, match="Charts worksheet exceeds the Excel row limit"):
        FileWriter().save_excel_sheets(
            {},
            path,
            chart_images=(oversized,),
            chart_sheet_name="Charts",
            streaming=False,
        )

    assert not path.exists()


@pytest.mark.parametrize("streaming", [True, False])
def test_image_captions_are_literal_wrapped_and_do_not_overlap_following_image(
    tmp_path: Path,
    chart_images: tuple[ExcelChartImage, ...],
    streaming: bool,
) -> None:
    captioned = replace(
        chart_images[0],
        heading="=literal heading",
        captions=("=literal caption", "Long explanation " * 180),
    )
    path = tmp_path / "captions.xlsx"
    FileWriter().save_excel_sheets(
        {},
        path,
        chart_images=(captioned, chart_images[1]),
        chart_sheet_name="Charts",
        streaming=streaming,
    )
    workbook = load_workbook(path)
    try:
        sheet = workbook["Charts"]
        assert sheet["A1"].value == "=literal heading"
        assert sheet["A1"].data_type == "s"
        assert sheet["A23"].value == "=literal caption"
        assert sheet["A23"].data_type == "s"
        assert sheet["A23"].alignment.wrap_text
        assert sheet.row_dimensions[24].height == 300
        assert sheet.row_dimensions[25].height <= 300
        assert sheet["A27"].value == chart_images[1].heading
        assert sheet._images[1].anchor._from.row == 27
        assert sheet.column_dimensions["A"].width == 150
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [True, False])
def test_caption_rows_are_included_in_excel_row_limit(
    tmp_path: Path,
    chart_images: tuple[ExcelChartImage, ...],
    streaming: bool,
    monkeypatch,
) -> None:
    monkeypatch.setattr("expo_jbm329.services.file_writer.EXCEL_MAX_ROWS", 23)
    image = replace(chart_images[0], captions=("first note", "second note"))
    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"original")
    with pytest.raises(ValueError, match="row limit"):
        FileWriter().save_excel_sheets(
            {},
            path,
            chart_images=(image,),
            chart_sheet_name="Charts",
            streaming=streaming,
        )
    assert path.read_bytes() == b"original"
    assert not list(tmp_path.glob(".existing-*.xlsx"))


@pytest.mark.parametrize("streaming", [True, False])
def test_cancellation_during_caption_writing_preserves_original_file(
    tmp_path: Path,
    chart_images: tuple[ExcelChartImage, ...],
    streaming: bool,
    monkeypatch,
) -> None:
    image = replace(chart_images[0], captions=("cancel at this note",))
    path = tmp_path / "existing.xlsx"
    path.write_bytes(b"original")
    original = FileWriter._add_excel_chart_images

    def cancel_during_caption(worksheet, images, *, cancel_check):
        calls = 0

        def cancel():
            nonlocal calls
            calls += 1
            if calls == 2:
                message = "Cancelled while writing captions."
                raise ExportCancelledError(message, path=str(path))
            cancel_check()

        original(worksheet, images, cancel_check=cancel)

    monkeypatch.setattr(FileWriter, "_add_excel_chart_images", staticmethod(cancel_during_caption))
    with pytest.raises(ExportCancelledError):
        FileWriter().save_excel_sheets(
            {},
            path,
            chart_images=(image,),
            chart_sheet_name="Charts",
            streaming=streaming,
        )
    assert path.read_bytes() == b"original"
    assert not list(tmp_path.glob(".existing-*.xlsx"))
