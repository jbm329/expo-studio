from __future__ import annotations

from typing import TYPE_CHECKING

import pandas as pd
import pytest
from openpyxl import load_workbook

from expo_jbm329.services.file_writer import ExportCancelledError, FileWriter

if TYPE_CHECKING:
    from pathlib import Path


@pytest.fixture
def writer() -> FileWriter:
    return FileWriter()


def test_validate_excel_sheet_name_rejects_invalid_names(writer: FileWriter) -> None:
    with pytest.raises(ValueError):
        writer._validate_excel_sheet_name("")
    with pytest.raises(ValueError):
        writer._validate_excel_sheet_name("bad:name")


def test_save_csv_writes_file(tmp_path: Path, writer: FileWriter) -> None:
    path = tmp_path / "out.csv"
    df = pd.DataFrame({"X": [1, 2]})

    result = writer.save_csv(df, path)

    assert result == path
    assert pd.read_csv(path)["X"].tolist() == [1, 2]


def test_save_csv_can_be_cancelled_before_write(tmp_path: Path, writer: FileWriter) -> None:
    path = tmp_path / "out.csv"
    df = pd.DataFrame({"X": [1]})

    with pytest.raises(ExportCancelledError):
        writer.save_csv(df, path, cancel_cb=lambda: True)


@pytest.mark.parametrize("streaming", [True, False])
@pytest.mark.parametrize("index", [True, False])
def test_named_excel_roundtrip(tmp_path: Path, writer: FileWriter, streaming: bool, index: bool) -> None:
    frame = pd.DataFrame({"Number": [1, 2], "Text": pd.Series(["hello", pd.NA], dtype="string")})
    frame.index = pd.Index([10, 20], name="row")
    sheets = {
        " Exact Å ": frame,
        "Different": pd.DataFrame({"Flag": [True]}),
        "Headers only": pd.DataFrame(columns=["Empty"]),
    }
    progress: list[int] = []
    path = tmp_path / "named.xlsx"
    assert writer.save_excel_sheets(sheets, path, streaming=streaming, index=index, progress_cb=progress.append) == path
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == list(sheets)
        rows = list(workbook[" Exact Å "].values)
        assert rows[0] == (("row", "Number", "Text") if index else ("Number", "Text"))
        assert rows[1] == ((10, 1, "hello") if index else (1, "hello"))
        assert rows[2] == ((20, 2, None) if index else (2, None))
        assert list(workbook["Different"].values)[1] == ((0, True) if index else (True,))
        assert workbook["Headers only"].max_row == 1
    finally:
        workbook.close()
    assert progress == sorted(set(progress))
    assert progress[0] == 0
    assert progress[-1] == 100
    assert progress.count(100) == 1
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == [path]


@pytest.mark.parametrize("index", [False, True])
def test_streaming_splits_all_rows_in_order(tmp_path: Path, writer: FileWriter, index: bool) -> None:
    path = tmp_path / "split.xlsx"
    frame = pd.DataFrame({"value": range(5)}, index=pd.Index(range(10, 15), name="row"))
    writer.save_excel_sheets(
        {"First": frame, "Last": pd.DataFrame({"other": [9]})},
        path,
        max_rows_per_sheet=3,
        chunk_size_rows=4,
        index=index,
    )
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["First", "First_2", "First_3", "Last"]
        actual = [row for name in workbook.sheetnames[:-1] for row in list(workbook[name].values)[1:]]
        expected = [(10 + value, value) if index else (value,) for value in range(5)]
        assert actual == expected
        for name in workbook.sheetnames[:-1]:
            assert list(workbook[name].values)[0] == (("row", "value") if index else ("value",))
    finally:
        workbook.close()


@pytest.mark.parametrize(
    "sheets",
    [
        {},
        {"": pd.DataFrame({"a": [1]})},
        {"bad:name": pd.DataFrame({"a": [1]})},
        {"'bad": pd.DataFrame({"a": [1]})},
        {"bad'": pd.DataFrame({"a": [1]})},
        {"bad\x01": pd.DataFrame({"a": [1]})},
        {"A" * 32: pd.DataFrame({"a": [1]})},
        {"Data": pd.DataFrame({"a": [1]}), "DATA": pd.DataFrame({"a": [2]})},
        {"Data": pd.DataFrame({"a": [1, 2, 3]}), "data_2": pd.DataFrame({"a": [1]})},
        {"data_2": pd.DataFrame({"a": [1]}), "Data": pd.DataFrame({"a": [1, 2, 3]})},
        {"A" * 30: pd.DataFrame({"a": [1, 2, 3]})},
        {"No columns": pd.DataFrame(index=[1])},
        {"Empty": pd.DataFrame()},
        {"Not frame": object()},
        {1: pd.DataFrame({"a": [1]})},
        [],
    ],
)
def test_invalid_named_requests_preflight_without_io(tmp_path: Path, writer: FileWriter, sheets) -> None:
    path = tmp_path / "new_directory" / "out.xlsx"
    with pytest.raises((ValueError, TypeError)):
        writer.save_excel_sheets(sheets, path, max_rows_per_sheet=3)
    assert not path.parent.exists()


@pytest.mark.parametrize(
    "options",
    [
        {"max_rows_per_sheet": 0},
        {"max_rows_per_sheet": 1},
        {"max_rows_per_sheet": 1_048_577},
        {"max_rows_per_sheet": True},
        {"max_rows_per_sheet": 2.5},
        {"chunk_size_rows": 0},
        {"chunk_size_rows": -1},
        {"chunk_size_rows": True},
        {"chunk_size_rows": "2"},
        {"streaming": "false"},
        {"streaming": False, "max_rows_per_sheet": 3},
    ],
)
def test_invalid_named_settings(tmp_path: Path, writer: FileWriter, options) -> None:
    path = tmp_path / "out.xlsx"
    with pytest.raises((ValueError, TypeError)):
        writer.save_excel_sheets({"Data": pd.DataFrame({"a": [1, 2, 3]})}, path, **options)
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == []


@pytest.mark.parametrize("streaming", [True, False])
@pytest.mark.parametrize("index", [True, False])
def test_named_column_limits_include_index(tmp_path: Path, writer: FileWriter, streaming: bool, index: bool) -> None:
    frame = pd.DataFrame(columns=range(16_384 if index else 16_385))
    with pytest.raises(ValueError, match="column limit"):
        writer.save_excel_sheets({"Data": frame}, tmp_path / "out.xlsx", streaming=streaming, index=index)
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == []


@pytest.mark.parametrize("streaming", [True, False])
def test_header_only_named_workbook(tmp_path: Path, writer: FileWriter, streaming: bool) -> None:
    path = tmp_path / "headers.xlsx"
    progress: list[int] = []
    writer.save_excel_sheets(
        {"One": pd.DataFrame(columns=["A"]), "Two": pd.DataFrame(columns=["B", "C"])},
        path,
        streaming=streaming,
        progress_cb=progress.append,
    )
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["One", "Two"]
        assert list(workbook["One"].values) == [("A",)]
        assert list(workbook["Two"].values) == [("B", "C")]
    finally:
        workbook.close()
    assert progress == sorted(set(progress))
    assert progress[-1] == 100


@pytest.mark.parametrize("streaming", [True, False])
@pytest.mark.parametrize("existing", [True, False])
@pytest.mark.parametrize("phase", ["before", "middle", "between", "commit"])
def test_named_cancellation_is_atomic(
    tmp_path: Path, writer: FileWriter, streaming: bool, existing: bool, phase: str, monkeypatch
) -> None:
    path = tmp_path / "out.xlsx"
    original = b"original destination"
    if existing:
        path.write_bytes(original)
    progress: list[int] = []
    saved = False
    save = FileWriter._save_excel_workbook

    def record_save(book, filename) -> None:
        nonlocal saved
        save(book, filename)
        saved = True

    monkeypatch.setattr(FileWriter, "_save_excel_workbook", staticmethod(record_save))
    sheets = {"One": pd.DataFrame({"a": range(4)}), "Two": pd.DataFrame({"b": range(4)})}

    def cancel() -> bool:
        if phase == "before":
            return True
        if phase == "commit":
            return saved
        threshold = 25 if phase == "middle" else 50
        return bool(progress and progress[-1] >= threshold)

    with pytest.raises(ExportCancelledError) as caught:
        writer.save_excel_sheets(
            sheets,
            path,
            streaming=streaming,
            chunk_size_rows=2,
            progress_cb=progress.append,
            cancel_cb=cancel,
            corr_id="atomic",
        )
    assert caught.value.corr_id == "atomic"
    assert caught.value.path == path.as_posix()
    assert caught.value.rows_written is not None
    assert caught.value.sheets_written is not None
    assert path.read_bytes() == original if existing else not path.exists()
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == ([path] if existing else [])
    assert 100 not in progress
    assert progress == sorted(set(progress))


@pytest.mark.parametrize("streaming", [True, False])
@pytest.mark.parametrize("existing", [True, False])
@pytest.mark.parametrize("phase", ["write", "replace"])
def test_named_failure_preserves_destination(
    tmp_path: Path, writer: FileWriter, streaming: bool, existing: bool, phase: str, monkeypatch
) -> None:
    from pathlib import Path

    path = tmp_path / "out.xlsx"
    if existing:
        path.write_bytes(b"original")
    progress: list[int] = []

    def fail(*args, **kwargs) -> None:
        message = "injected failure"
        raise OSError(message)

    monkeypatch.setattr(
        FileWriter if phase == "write" else Path, "_save_excel_workbook" if phase == "write" else "replace", fail
    )
    with pytest.raises(OSError, match="injected failure"):
        writer.save_excel_sheets(
            {"One": pd.DataFrame({"a": [1]})}, path, streaming=streaming, progress_cb=progress.append
        )
    assert path.read_bytes() == b"original" if existing else not path.exists()
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == ([path] if existing else [])
    assert 100 not in progress


def test_streaming_normalizes_cells_and_index(tmp_path: Path, writer: FileWriter) -> None:
    import numpy as np

    path = tmp_path / "clean.xlsx"
    frame = pd.DataFrame({
        "value": [
            pd.Timestamp("2026-01-01", tz="UTC"),
            b"hello",
            np.int64(3),
            np.float64("inf"),
            pd.Period("2026-01", freq="M"),
            pd.Interval(1, 2),
            pd.NA,
        ],
    })
    writer.save_excel_sheets({"Cells": frame}, path, index=True, na_rep="missing")
    workbook = load_workbook(path)
    try:
        values = [row[1] for row in list(workbook["Cells"].values)[1:]]
        assert values[0] == pd.Timestamp("2026-01-01").to_pydatetime()
        assert values[1:] == ["hello", 3, "missing", "2026-01", "(1, 2]", "missing"]
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [True, False])
def test_legacy_single_excel_roundtrip(tmp_path: Path, writer: FileWriter, streaming: bool) -> None:
    path = tmp_path / "legacy.xlsx"
    writer.save_excel(pd.DataFrame({"a": [1, 2]}), path, streaming=streaming)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Data"]
        assert list(workbook["Data"].values) == [("a",), (1,), (2,)]
    finally:
        workbook.close()


def test_legacy_streaming_cancellation_keeps_partial_workbook(tmp_path: Path, writer: FileWriter) -> None:
    path = tmp_path / "legacy.xlsx"
    progress: list[int] = []
    with pytest.raises(ExportCancelledError) as caught:
        writer.save_excel(
            pd.DataFrame({"a": range(4)}),
            path,
            progress_cb=progress.append,
            cancel_cb=lambda: bool(progress and progress[-1] >= 25),
        )
    assert caught.value.rows_written == 1
    workbook = load_workbook(path)
    try:
        assert list(workbook["Data"].values) == [("a",), (0,)]
    finally:
        workbook.close()


def test_named_validation_freezes_membership_not_contents(writer: FileWriter) -> None:
    frame = pd.DataFrame({"a": [1]})
    sheets = {"One": frame}
    captured = writer.validate_excel_sheets(sheets)
    sheets["Other"] = pd.DataFrame({"b": [2]})
    frame.loc[0, "a"] = 3
    assert list(captured) == ["One"]
    assert captured["One"] is frame
    assert captured["One"].loc[0, "a"] == 3


@pytest.mark.parametrize("streaming", [True, False])
def test_named_archive_failure_closes_handles(tmp_path: Path, writer: FileWriter, streaming: bool, monkeypatch) -> None:
    from openpyxl.writer.excel import ExcelWriter

    path = tmp_path / "out.xlsx"
    path.write_bytes(b"original")

    def fail(*args, **kwargs) -> None:
        message = "archive failed"
        raise OSError(message)

    monkeypatch.setattr(ExcelWriter, "write_data", fail)
    with pytest.raises(OSError, match="archive failed"):
        writer.save_excel_sheets({"One": pd.DataFrame({"a": [1]})}, path, streaming=streaming)
    assert path.read_bytes() == b"original"
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == [path]


@pytest.mark.parametrize("streaming", [True, False])
def test_named_bad_cell_failure_is_atomic(tmp_path: Path, writer: FileWriter, streaming: bool) -> None:
    path = tmp_path / "out.xlsx"
    path.write_bytes(b"original")
    frame = (
        pd.DataFrame({"value": [pd.Timestamp("2026-01-01", tz="UTC")]})
        if not streaming
        else pd.DataFrame({"value": [1, {"unsupported": 2}]})
    )
    progress: list[int] = []
    with pytest.raises(ValueError):
        writer.save_excel_sheets({"One": frame}, path, streaming=streaming, progress_cb=progress.append)
    assert path.read_bytes() == b"original"
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == [path]
    assert 100 not in progress


@pytest.mark.parametrize("streaming", [True, False])
def test_named_progress_failure_is_atomic(tmp_path: Path, writer: FileWriter, streaming: bool) -> None:
    path = tmp_path / "out.xlsx"
    path.write_bytes(b"original")

    def progress(value: int) -> None:
        if value > 0:
            message = "progress failed"
            raise RuntimeError(message)

    with pytest.raises(RuntimeError, match="progress failed"):
        writer.save_excel_sheets(
            {"One": pd.DataFrame({"a": [1, 2]}), "Two": pd.DataFrame({"b": [3]})},
            path,
            streaming=streaming,
            progress_cb=progress,
        )
    assert path.read_bytes() == b"original"
    assert [entry for entry in tmp_path.iterdir() if entry.is_file()] == [path]


@pytest.mark.parametrize("setting", ["max_rows_per_sheet", "chunk_size_rows"])
def test_named_preflight_validates_reloaded_settings(tmp_path: Path, writer: FileWriter, setting: str) -> None:
    writer.reload_settings({"excel": {setting: 0}})
    with pytest.raises(ValueError):
        writer.save_excel_sheets({"One": pd.DataFrame({"a": [1]})}, tmp_path / "out.xlsx")


@pytest.mark.parametrize("index", [True, False])
def test_named_pandas_row_limit_includes_header(tmp_path: Path, writer: FileWriter, index: bool) -> None:
    path = tmp_path / "out.xlsx"
    writer.save_excel_sheets(
        {"One": pd.DataFrame({"a": [1, 2]})}, path, streaming=False, index=index, max_rows_per_sheet=3
    )
    workbook = load_workbook(path)
    try:
        assert workbook["One"].max_row == 3
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [True, False])
def test_named_multiindex_index_roundtrip(tmp_path: Path, writer: FileWriter, streaming: bool) -> None:
    frame = pd.DataFrame(
        {"value": [1, 2]}, index=pd.MultiIndex.from_tuples([("a", 10), ("b", 20)], names=["group", "row"])
    )
    path = tmp_path / "out.xlsx"
    writer.save_excel_sheets({"One": frame}, path, streaming=streaming, index=True)
    workbook = load_workbook(path)
    try:
        assert list(workbook["One"].values) == [("group", "row", "value"), ("a", 10, 1), ("b", 20, 2)]
    finally:
        workbook.close()


@pytest.mark.parametrize("index", [True, False])
def test_named_multiindex_column_preflight(tmp_path: Path, writer: FileWriter, index: bool) -> None:
    frame = pd.DataFrame([[1]], columns=pd.MultiIndex.from_tuples([("group", "value")]))
    with pytest.raises(ValueError, match="row limit" if index else "requires index"):
        writer.save_excel_sheets(
            {"One": frame}, tmp_path / "out.xlsx", streaming=False, index=index, max_rows_per_sheet=3
        )


def test_named_commit_cancellation_reports_split_counters(tmp_path: Path, writer: FileWriter, monkeypatch) -> None:
    saved = False
    original_save = FileWriter._save_excel_workbook

    def save(book, filename) -> None:
        nonlocal saved
        original_save(book, filename)
        saved = True

    monkeypatch.setattr(FileWriter, "_save_excel_workbook", staticmethod(save))
    with pytest.raises(ExportCancelledError) as caught:
        writer.save_excel_sheets(
            {"One": pd.DataFrame({"a": range(5)}), "Two": pd.DataFrame(columns=["b"])},
            tmp_path / "out.xlsx",
            max_rows_per_sheet=3,
            cancel_cb=lambda: saved,
        )
    assert caught.value.rows_written == 5
    assert caught.value.sheets_written == 4


def test_named_notification_failure_after_commit_does_not_report_export_failure(
    tmp_path: Path, writer: FileWriter, caplog
) -> None:
    path = tmp_path / "out.xlsx"
    path.write_bytes(b"original")

    def progress(value: int) -> None:
        if value == 100:
            message = "notification failed"
            raise RuntimeError(message)

    assert writer.save_excel_sheets({"One": pd.DataFrame({"a": [1]})}, path, progress_cb=progress) == path
    workbook = load_workbook(path)
    try:
        assert list(workbook["One"].values) == [("a",), (1,)]
    finally:
        workbook.close()
    assert "progress notification failed after Excel commit" in caplog.text


@pytest.mark.parametrize("streaming", [True, False])
def test_legacy_excel_header_only(tmp_path: Path, writer: FileWriter, streaming: bool) -> None:
    path = tmp_path / "out.xlsx"
    writer.save_excel(pd.DataFrame(columns=["a"]), path, streaming=streaming)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Data"]
        assert list(workbook["Data"].values) == [("a",)]
    finally:
        workbook.close()


@pytest.mark.parametrize("streaming", [True, False])
def test_legacy_excel_cancellation_before_data(tmp_path: Path, writer: FileWriter, streaming: bool) -> None:
    path = tmp_path / "out.xlsx"
    with pytest.raises(ExportCancelledError) as caught:
        writer.save_excel(pd.DataFrame({"a": [1]}), path, streaming=streaming, cancel_cb=lambda: True, corr_id="legacy")
    assert caught.value.rows_written == 0
    assert caught.value.corr_id == "legacy"
    if streaming:
        workbook = load_workbook(path)
        try:
            assert list(workbook["Data"].values) == [("a",)]
        finally:
            workbook.close()
    else:
        assert not path.exists()


@pytest.mark.parametrize("index", [True, False])
def test_legacy_excel_splitting_keeps_all_rows(tmp_path: Path, writer: FileWriter, index: bool) -> None:
    path = tmp_path / "out.xlsx"
    writer.save_excel(pd.DataFrame({"a": range(5)}), path, max_rows_per_sheet=3, index=index)
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Data", "Data_2", "Data_3"]
        values = [row for sheet in workbook for row in list(sheet.values)[1:]]
        assert values == [(value, value) if index else (value,) for value in range(5)]
    finally:
        workbook.close()
